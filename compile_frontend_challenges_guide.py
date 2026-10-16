#!/usr/bin/env python3
"""
compile_frontend_challenges_guide.py — Generate high-typography Frontend Challenges Guide via Word COM
"""

import sys
import json
from pathlib import Path
from datetime import datetime

# UTF-8 console output
if sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr.encoding.lower() != 'utf-8':
    sys.stderr.reconfigure(encoding='utf-8')

import win32com.client
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_OUT = Path(r"e:\anti\reports\FRONTEND_CHALLENGES_SENIOR_GUIDE.docx")
PDF_OUT = Path(r"e:\anti\reports\FRONTEND_CHALLENGES_SENIOR_GUIDE.pdf")
DATA_PATH = Path(r"e:\anti\data\frontend_challenges.json")

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def main():
    print("[+] Compiling Frontend Challenges Senior Guide...")
    with open(DATA_PATH, "r", encoding="utf-8") as f:
        challenges = json.load(f)

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Title Banner
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("GLOBAL FRONT-END CHALLENGES ARSENAL")
    r.font.name = "Arial"
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = st.add_run("Practical Engineering Tests, Take-Home Codebases & Tech Stack Intelligence")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Executive Summary Card
    card_table = doc.add_table(rows=1, cols=1)
    card_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    card_cell = card_table.cell(0, 0)
    set_cell_background(card_cell, "F1F5F9")
    set_cell_margins(card_cell, top=140, bottom=140, left=180, right=180)

    p_card = card_cell.paragraphs[0]
    r_c1 = p_card.add_run("EXECUTIVE INTELLIGENCE BRIEFING\n")
    r_c1.font.bold = True
    r_c1.font.size = Pt(11)
    r_c1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    summary_text = (
        f"This publication synthesizes {len(challenges)} open-source frontend hiring challenges collected from real tech "
        "companies worldwide (curated by Felipe Fialho). Instead of synthetic LeetCode trivia, these assessments reflect actual "
        "production expectations: component state isolation, REST/GraphQL API integration, responsive CSS architectures, "
        "and client-side performance engineering.\n\n"
        "Key Technologies: React, TypeScript, Vue.js, Angular, Next.js, Redux, Sass, GraphQL, and automated testing frameworks."
    )
    r_c2 = p_card.add_run(summary_text)
    r_c2.font.size = Pt(9.5)
    r_c2.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Featured Challenges Table
    h2 = doc.add_paragraph()
    r_h2 = h2.add_run("Featured Real-World Frontend Take-Home Assessments")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    headers = ["Company", "Stack & Technologies", "Lang", "Challenge Repository"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "0F172A")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Populate table with challenges
    for row_idx, item in enumerate(challenges[:35]):
        row_cells = table.add_row().cells
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        vals = [
            item.get("company_name", "")[:24],
            item.get("stack", "")[:36],
            item.get("lang", "").upper(),
            item.get("challenge_url", "")[:42]
        ]
        for col_idx, val in enumerate(vals):
            row_cells[col_idx].text = val
            set_cell_background(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], top=80, bottom=80, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Save DOCX
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DOCX_OUT))
    print(f"[+] DOCX compiled: {DOCX_OUT}")

    # Compile PDF via Word COM
    print("[+] Compiling PDF via Word COM...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        wdoc = word.Documents.Open(str(DOCX_OUT.resolve()))
        wdoc.ExportAsFixedFormat(str(PDF_OUT.resolve()), 17)
        wdoc.Close(False)
        print(f"[+] PDF compiled successfully: {PDF_OUT}")
    except Exception as e:
        print(f"[-] Word COM error: {e}")
    finally:
        word.Quit()

if __name__ == "__main__":
    main()
