#!/usr/bin/env python3
"""
compile_new_grad_2027_guide.py — Generate high-typography New Grad 2027 Master Guide via Word COM.
Produces reports/NEW_GRAD_2027_MASTER_GUIDE.docx and reports/NEW_GRAD_2027_MASTER_GUIDE.pdf
"""

import sys
import json
import sqlite3
from pathlib import Path
from datetime import datetime

# UTF-8 console output
if sys.stdout.encoding.lower() != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')
if sys.stderr.encoding.lower() != 'utf-8':
    sys.stderr.reconfigure(encoding='utf-8')

import win32com.client
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_OUT = Path(r"e:\anti\reports\NEW_GRAD_2027_MASTER_GUIDE.docx")
PDF_OUT = Path(r"e:\anti\reports\NEW_GRAD_2027_MASTER_GUIDE.pdf")
DATA_PATH = Path(r"e:\anti\data\new_grad_2027_active.json")

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def main():
    print("[+] Compiling New Grad 2027 Master Publication...")
    with open(DATA_PATH, "r", encoding="utf-8") as f:
        listings = json.load(f)

    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.7)
        s.bottom_margin = Inches(0.7)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    # Title Banner
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run("GLOBAL NEW GRAD 2027 ROLES COMPENDIUM")
    r.font.name = "Arial"
    r.font.size = Pt(22)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = st.add_run("Targeted Opportunities, Role Analysis & Strategic Application Intelligence (vanshb03/New-Grad-2027)")
    r_sub.font.name = "Arial"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Executive Summary Card
    card_table = doc.add_table(rows=1, cols=1)
    card_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    card_cell = card_table.rows[0].cells[0]
    set_cell_background(card_cell, "F1F5F9")
    set_cell_margins(card_cell, top=140, bottom=140, left=180, right=180)

    p_card = card_cell.paragraphs[0]
    p_card.paragraph_format.line_spacing = 1.2
    run_card_title = p_card.add_run("EXECUTIVE BRIEFING & PIPELINE TELEMETRY\n")
    run_card_title.font.bold = True
    run_card_title.font.size = Pt(11)
    run_card_title.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Analysis metrics
    total_active = len(listings)
    companies = len(set(x.get("company_name", "") for x in listings))
    sponsorship_offers = len([x for x in listings if x.get("sponsorship") == "Offers Sponsorship"])
    sponsorship_no = len([x for x in listings if x.get("sponsorship") == "Does Not Offer Sponsorship"])
    us_cit_req = len([x for x in listings if "Citizenship" in str(x.get("sponsorship", ""))])
    other_spons = total_active - (sponsorship_offers + sponsorship_no + us_cit_req)

    summary_text = (
        f"• Total Active Roles Ingested: {total_active:,d} verified full-time positions across {companies:,d} top-tier employers.\n"
        f"• Sponsorship Profile: {sponsorship_offers:,d} Offer Sponsorship | {sponsorship_no:,d} Do Not Offer | {us_cit_req:,d} US Citizen Only | {other_spons:,d} Flexible/Other.\n"
        f"• Target Graduate Profile: Aditya Mehra (DSU '26, GPA 8.78/10.0, Systems & Autonomous AI Agent Architect).\n"
        f"• Sourcing Provenance: Synchronized and deduplicated from upstream repository vanshb03/New-Grad-2027.\n"
        f"• Strategic Directive: Early-application velocity, direct referral alignment, and deep systems engineering matching."
    )
    run_body = p_card.add_run(summary_text)
    run_body.font.size = Pt(9.5)
    run_body.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section Heading: Featured Roles Table
    h2 = doc.add_paragraph()
    h2_run = h2.add_run("PRIMARY STRATEGIC OPPORTUNITIES INDEX (SAMPLE 50 POSITIONS)")
    h2_run.font.bold = True
    h2_run.font.size = Pt(13)
    h2_run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
    h2.paragraph_format.space_after = Pt(6)

    # Table
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    col_widths = [Inches(1.8), Inches(2.6), Inches(1.3), Inches(1.3)]
    for i, w in enumerate(col_widths):
        table.columns[i].width = w

    headers = ["Company", "Role Title", "Location", "Sponsorship"]
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=100, bottom=100, left=120, right=120)
        p = hdr_cells[i].paragraphs[0]
        p.runs[0].font.bold = True
        p.runs[0].font.size = Pt(9)
        p.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    # Select top 50 positions
    selected = listings[:50]
    for row_idx, item in enumerate(selected):
        row_cells = table.add_row().cells
        bg_col = "FFFFFF" if row_idx % 2 == 0 else "F8FAFC"
        
        locs = ", ".join(item.get("locations", []))[:28] if isinstance(item.get("locations"), list) else str(item.get("locations", ""))[:28]
        vals = [
            item.get("company_name", "")[:24],
            item.get("title", "")[:35],
            locs,
            item.get("sponsorship", "Other")[:20]
        ]
        for col_idx, val in enumerate(vals):
            row_cells[col_idx].text = val
            set_cell_background(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], top=70, bottom=70, left=100, right=100)
            p = row_cells[col_idx].paragraphs[0]
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

    # Save DOCX
    DOCX_OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DOCX_OUT))
    print(f"[+] DOCX compiled: {DOCX_OUT}")

    # Compile PDF via Word COM
    print("[+] Compiling PDF via Word COM...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        wdoc = word.Documents.Open(str(DOCX_OUT.resolve()))
        wdoc.ExportAsFixedFormat(str(PDF_OUT.resolve()), 17) # 17 = wdExportFormatPDF
        wdoc.Close(False)
        print(f"[+] PDF compiled successfully: {PDF_OUT}")
    except Exception as e:
        print(f"[-] Word COM error: {e}")
    finally:
        word.Quit()

if __name__ == "__main__":
    main()
