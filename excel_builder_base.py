# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

print("Initializing workbook builder...")
wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Styling Definitions
font_title = Font(name="Calibri", size=16, bold=True, color="1B365D")
font_subtitle = Font(name="Calibri", size=11, italic=True, color="4B5563")
font_section = Font(name="Calibri", size=13, bold=True, color="1B365D")
font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
font_sub_header = Font(name="Calibri", size=11, bold=True, color="1B365D")
font_bold = Font(name="Calibri", size=11, bold=True, color="000000")
font_normal = Font(name="Calibri", size=10, color="000000")
font_small = Font(name="Calibri", size=9, color="4B5563")

fill_navy = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
fill_ice = PatternFill(start_color="E8EEF5", end_color="E8EEF5", fill_type="solid")
fill_zebra = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
fill_p1 = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid") # light green
fill_p2 = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid") # light yellow
fill_p3 = PatternFill(start_color="F3F4F6", end_color="F3F4F6", fill_type="solid") # light gray
fill_p4 = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # light red

thin_border_side = Side(border_style="thin", color="D1D5DB")
border_thin = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=thin_border_side)
border_header = Border(left=thin_border_side, right=thin_border_side, top=thin_border_side, bottom=Side(border_style="medium", color="1B365D"))

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center", wrap_text=True)
align_header = Alignment(horizontal="center", vertical="center", wrap_text=True)

def apply_table_styles(ws, start_row, end_row, start_col, end_col, is_header=True):
    for row in range(start_row, end_row + 1):
        for col in range(start_col, end_col + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = border_thin
            if row == start_row and is_header:
                cell.font = font_header
                cell.fill = fill_navy
                cell.alignment = align_header
            else:
                if row % 2 == 1:
                    if cell.fill.fill_type is None:
                        cell.fill = fill_zebra

def auto_fit_columns(ws, max_cols=30):
    ws.views.sheetView[0].showGridLines = True
    for col in range(1, min(ws.max_column + 1, max_cols + 1)):
        col_letter = get_column_letter(col)
        max_len = 0
        for row in range(1, min(ws.max_row + 1, 100)):
            val = ws.cell(row=row, column=col).value
            if val is not None:
                val_str = str(val)
                if len(val_str) > max_len:
                    max_len = len(val_str)
        width = max(max_len + 3, 12)
        if width > 55:
            width = 55
        ws.column_dimensions[col_letter].width = width

print("Base setup ready.")
