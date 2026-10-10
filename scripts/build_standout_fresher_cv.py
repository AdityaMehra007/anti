import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client

BASE_DIR = Path(r"e:\anti")

# Palette: Clean, polished, high-contrast, modern executive
C_NAVY = RGBColor(15, 23, 42)       # #0f172a Deep Slate Navy
C_SLATE = RGBColor(71, 85, 105)     # #475569 Slate
C_BODY = RGBColor(30, 41, 59)       # #1e293b Body text
C_MUTED = RGBColor(100, 116, 139)   # #64748b Metadata
C_LINK = "1E3A8A"                   # #1e3a8a Deep Classic Blue

def add_clean_divider(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')       # 0.5 pt
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CBD5E1')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_hyperlink_run(paragraph, url, display_text, font_size_pt=8.2, color_hex="1E3A8A", bold=False, underline=True):
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'<w:color w:val="{color_hex}"/>'
        f'{"<w:b/>" if bold else ""}'
        f'{"<w:u w:val=\"single\"/>" if underline else ""}'
        f'<w:sz w:val="{int(font_size_pt * 2)}"/>'
        f'</w:rPr>'
        f'<w:t>{display_text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    paragraph._p.append(hyperlink)

def create_standout_cv(filename="ADITYA_MEHRA_STANDOUT_FRESHER_CV.docx"):
    doc = Document()

    # Exact page margins calibrated for complete 1-page fit
    for section in doc.sections:
        section.page_width = Inches(8.27)    # A4
        section.page_height = Inches(11.69)  # A4
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.50)
        section.right_margin = Inches(0.50)

    # 1. HEADER - NAME
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    r_name = p_name.add_run("ADITYA MEHRA")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(19)
    r_name.font.bold = True
    r_name.font.color.rgb = C_NAVY

    # SUBTITLE: Polished for MNCs & Startups
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2)
    r_sub = p_sub.add_run("BBA Graduate  |  Business Operations, AI-Driven Builder & Associate")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.2)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_SLATE

    # CONTACT ROW 1: Location | Phone | Email
    p_con1 = doc.add_paragraph()
    p_con1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con1.paragraph_format.space_before = Pt(0)
    p_con1.paragraph_format.space_after = Pt(1)
    
    r_loc = p_con1.add_run("Bengaluru, Karnataka, India  |  ")
    r_loc.font.name = "Calibri"; r_loc.font.size = Pt(8.2); r_loc.font.color.rgb = C_MUTED
    
    r_ph = p_con1.add_run("Phone: ")
    r_ph.font.name = "Calibri"; r_ph.font.size = Pt(8.2); r_ph.font.bold = True; r_ph.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "tel:+917003456624", "+91 7003456624", font_size_pt=8.2, color_hex=C_LINK, underline=False)
    
    r_sep1 = p_con1.add_run("  |  Email: ")
    r_sep1.font.name = "Calibri"; r_sep1.font.size = Pt(8.2); r_sep1.font.bold = True; r_sep1.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "mailto:adityamehra799@gmail.com", "adityamehra799@gmail.com", font_size_pt=8.2, color_hex=C_LINK, underline=False)

    # CONTACT ROW 2: Website / Proof of Work | LinkedIn | GitHub
    p_con2 = doc.add_paragraph()
    p_con2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con2.paragraph_format.space_before = Pt(0)
    p_con2.paragraph_format.space_after = Pt(4)

    r_web_lbl = p_con2.add_run("Proof of Work & Portfolio: ")
    r_web_lbl.font.name = "Calibri"; r_web_lbl.font.size = Pt(8.2); r_web_lbl.font.bold = True; r_web_lbl.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://adityamehra007.github.io/ADI-OS/", "adityamehra007.github.io/ADI-OS", font_size_pt=8.2, color_hex=C_LINK, underline=True)

    r_sep2 = p_con2.add_run("  |  LinkedIn: ")
    r_sep2.font.name = "Calibri"; r_sep2.font.size = Pt(8.2); r_sep2.font.bold = True; r_sep2.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://www.linkedin.com/in/aditya-mehra-b8644b326", "linkedin.com/in/aditya-mehra-b8644b326", font_size_pt=8.2, color_hex=C_LINK, underline=True)

    r_sep3 = p_con2.add_run("  |  GitHub: ")
    r_sep3.font.name = "Calibri"; r_sep3.font.size = Pt(8.2); r_sep3.font.bold = True; r_sep3.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://github.com/AdityaMehra007", "github.com/AdityaMehra007", font_size_pt=8.2, color_hex=C_LINK, underline=True)

    def add_sec(title, before=3):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title.upper())
        r.font.name = "Calibri"
        r.font.size = Pt(9.2)
        r.font.bold = True
        r.font.color.rgb = C_NAVY
        add_clean_divider(p)

    # 2. CAREER OBJECTIVE (High Energy, Grounded, Builder Mindset)
    add_sec("Career Objective", before=1)
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(1)
    p_obj.paragraph_format.space_after = Pt(2.5)
    p_obj.paragraph_format.line_spacing = 1.08
    r_obj = p_obj.add_run(
        "Energetic and grounded BBA graduate eager to contribute to forward-thinking startups and established MNCs. "
        "Combines practical business knowledge with hands-on modern tech skills—actively vibe coding web applications, building live websites, and producing creative video content with AI tools. "
        "Brings proven internship experience in corporate client outreach, operations data verification, vendor coordination, and store retail sales. "
        "Low-profile attitude, fast learner, and determined to deliver real value and high reliability from day one."
    )
    r_obj.font.name = "Calibri"
    r_obj.font.size = Pt(8.4)
    r_obj.font.color.rgb = C_BODY

    # 3. PROOF OF WORK & AI CREATIVE PROJECTS (Stands Out!)
    add_sec("Proof of Work & AI Projects", before=3)

    pow_items = [
        {
            "title": "Interactive Web Applications & Vibe Coding",
            "context": "Personal Projects & Live Companion (adityamehra007.github.io/ADI-OS)",
            "bullets": [
                "Used AI coding tools and prompt engineering to vibe code interactive web apps, personal portfolio tools, and live responsive web pages.",
                "Explored modern UI designs, tested clean web layouts, and deployed working web prototypes to GitHub Pages."
            ]
        },
        {
            "title": "AI Video Generation & Digital Content Creation",
            "context": "Multimedia Brand Assets & Visual Storytelling",
            "bullets": [
                "Created engaging video clips, product concept visuals, and presentation materials using cutting-edge AI video and image tools.",
                "Produced creative content for digital brand showcases, marketing slide decks, and social media mockups."
            ]
        }
    ]

    for item in pow_items:
        p_h = doc.add_paragraph()
        p_h.paragraph_format.space_before = Pt(2)
        p_h.paragraph_format.space_after = Pt(0.5)
        p_h.paragraph_format.keep_with_next = True

        r_t = p_h.add_run(item["title"])
        r_t.font.name = "Calibri"; r_t.font.size = Pt(8.5); r_t.bold = True; r_t.font.color.rgb = C_NAVY

        r_ctx = p_h.add_run(f"  |  {item['context']}")
        r_ctx.font.name = "Calibri"; r_ctx.font.size = Pt(8.0); r_ctx.font.italic = True; r_ctx.font.color.rgb = C_SLATE

        for b in item["bullets"]:
            p_b = doc.add_paragraph()
            p_b.paragraph_format.left_indent = Inches(0.12)
            p_b.paragraph_format.space_before = Pt(0.2)
            p_b.paragraph_format.space_after = Pt(0.5)
            p_b.paragraph_format.line_spacing = 1.04

            rb = p_b.add_run("• ")
            rb.font.name = "Calibri"; rb.font.size = Pt(8.0); rb.bold = True; rb.font.color.rgb = C_SLATE

            rt = p_b.add_run(b)
            rt.font.name = "Calibri"; rt.font.size = Pt(8.0); rt.font.color.rgb = C_BODY

    # 4. WORK EXPERIENCE & INTERNSHIPS (All 5 Roles, Point-to-Point, No Numbers)
    add_sec("Work Experience & Internships", before=3)

    experiences = [
        {
            "role": "Business Development Intern",
            "company": "Pencil Mark Interior Solutions, Bengaluru",
            "bullets": [
                "Contacted corporate clients via phone and email to introduce office interior design solutions; prepared client proposals and meeting decks.",
                "Earned a written appreciation letter from senior leadership for strong dedication, consistent effort, and proactive client support."
            ]
        },
        {
            "role": "Operations Intern",
            "company": "Instawork Services India, Bengaluru",
            "bullets": [
                "Supported operations team with collecting, sorting, and verifying data for AI and robotics projects with high attention to detail."
            ]
        },
        {
            "role": "Stall Coordinator",
            "company": "AERO India Exhibition (Salt in My Coca), Bengaluru",
            "bullets": [
                "Managed exhibition stall setup, welcomed attendees, addressed customer questions, and maintained complete stock protection."
            ]
        },
        {
            "role": "Event Coordinator",
            "company": "Commercial & College Events, Bengaluru",
            "bullets": [
                "Coordinated vendor setups for stage, sound, and lighting, and managed artist travel, hospitality, and backstage arrangements."
            ]
        },
        {
            "role": "Store Assistant",
            "company": "Family Retail Store & Food Business, Bengaluru",
            "bullets": [
                "Handled daily customer billing, counter sales, inventory register updates, and supplier follow-ups for fresh materials."
            ]
        }
    ]

    for exp in experiences:
        p_h = doc.add_paragraph()
        p_h.paragraph_format.space_before = Pt(2)
        p_h.paragraph_format.space_after = Pt(0.5)
        p_h.paragraph_format.keep_with_next = True

        r_r = p_h.add_run(exp["role"])
        r_r.font.name = "Calibri"; r_r.font.size = Pt(8.5); r_r.bold = True; r_r.font.color.rgb = C_NAVY

        r_c = p_h.add_run(f"  |  {exp['company']}")
        r_c.font.name = "Calibri"; r_c.font.size = Pt(8.0); r_c.font.italic = True; r_c.font.color.rgb = C_SLATE

        for b in exp["bullets"]:
            p_b = doc.add_paragraph()
            p_b.paragraph_format.left_indent = Inches(0.12)
            p_b.paragraph_format.space_before = Pt(0.2)
            p_b.paragraph_format.space_after = Pt(0.5)
            p_b.paragraph_format.line_spacing = 1.04

            rb = p_b.add_run("• ")
            rb.font.name = "Calibri"; rb.font.size = Pt(8.0); rb.bold = True; rb.font.color.rgb = C_SLATE

            rt = p_b.add_run(b)
            rt.font.name = "Calibri"; rt.font.size = Pt(8.0); rt.font.color.rgb = C_BODY

    # 5. EDUCATION
    add_sec("Education", before=3)
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(1)
    p_edu.paragraph_format.space_after = Pt(1.5)
    p_edu.paragraph_format.line_spacing = 1.06
    p_edu.paragraph_format.keep_with_next = True

    r_deg = p_edu.add_run("Bachelor of Business Administration (BBA) — International Business  |  ")
    r_deg.font.name = "Calibri"; r_deg.font.size = Pt(8.6); r_deg.bold = True; r_deg.font.color.rgb = C_NAVY

    r_uni = p_edu.add_run("Dayananda Sagar University, Bengaluru\n")
    r_uni.font.name = "Calibri"; r_uni.font.size = Pt(8.2); r_uni.font.italic = True; r_uni.font.color.rgb = C_SLATE

    r_sub = p_edu.add_run("Key Subjects: International Trade, Marketing Management, Supply Chain, Business Communication, Accounting.")
    r_sub.font.name = "Calibri"; r_sub.font.size = Pt(8.0); r_sub.font.color.rgb = C_BODY

    # 6. KEY SKILLS (AI Tools, Office Tools, Business, Languages)
    add_sec("Key Skills", before=3)
    skills = [
        ("AI & Digital Builder Tools: ", "AI Vibe Coding, AI Video Creation, ChatGPT, Claude, Prompt Engineering, HTML/Web Basics, GitHub."),
        ("Computer & Office Tools: ", "MS Excel, MS Word, MS PowerPoint, Google Workspace (Sheets, Docs), Online Research, Data Entry."),
        ("Business & Workplace Skills: ", "Client Communication, Customer Service, Operations Support, Vendor Coordination, Teamwork, Problem Solving."),
        ("Languages: ", "English, Hindi.")
    ]
    for s_title, s_desc in skills:
        p_s = doc.add_paragraph()
        p_s.paragraph_format.left_indent = Inches(0.12)
        p_s.paragraph_format.space_before = Pt(0.2)
        p_s.paragraph_format.space_after = Pt(0.5)
        p_s.paragraph_format.line_spacing = 1.04

        rb = p_s.add_run("• ")
        rb.font.name = "Calibri"; rb.font.size = Pt(8.0); rb.bold = True; rb.font.color.rgb = C_SLATE

        rp = p_s.add_run(s_title)
        rp.font.name = "Calibri"; rp.font.size = Pt(8.0); rp.bold = True; rp.font.color.rgb = C_NAVY

        rd = p_s.add_run(s_desc)
        rd.font.name = "Calibri"; rd.font.size = Pt(8.0); rd.font.color.rgb = C_BODY

    # 7. CERTIFICATES & ACHIEVEMENTS
    add_sec("Certificates & Achievements", before=3)
    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.left_indent = Inches(0.12)
    p_cert.paragraph_format.space_before = Pt(1)
    p_cert.paragraph_format.space_after = Pt(1)
    p_cert.paragraph_format.line_spacing = 1.04

    certs = [
        "Appreciation Letter for Good Performance (Pencil Mark)",
        "Digital Marketing & E-commerce (Google)",
        "Service Marketing (NPTEL, IIT Kharagpur)",
        "Artificial Intelligence Tools Workshop (Outskill)"
    ]
    
    rb = p_cert.add_run("• ")
    rb.font.name = "Calibri"; rb.font.size = Pt(8.0); rb.bold = True; rb.font.color.rgb = C_SLATE
    rc = p_cert.add_run("   |   ".join(certs))
    rc.font.name = "Calibri"; rc.font.size = Pt(7.8); rc.font.color.rgb = C_BODY

    out_path = BASE_DIR / filename
    doc.save(out_path)
    return out_path

def generate_html_version(filename="ADITYA_MEHRA_STANDOUT_FRESHER_CV.html"):
    html_content = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Aditya Mehra — Standout Fresher CV</title>
<style>
  :root {
    --primary: #0f172a;
    --accent: #1e3a8a;
    --slate: #475569;
    --body: #1e293b;
    --border: #cbd5e1;
    --tag-bg: #f1f5f9;
  }
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body {
    font-family: Calibri, -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
    color: var(--body);
    background-color: #e2e8f0;
    line-height: 1.32;
    font-size: 8.8pt;
    padding: 20px;
  }
  .cv-page {
    max-width: 820px;
    margin: 0 auto;
    background: #ffffff;
    padding: 28px 36px;
    border-radius: 6px;
    box-shadow: 0 4px 20px rgba(0,0,0,0.08);
  }
  header {
    text-align: center;
    border-bottom: 2px solid var(--primary);
    padding-bottom: 8px;
    margin-bottom: 10px;
  }
  h1 {
    font-size: 19pt;
    font-weight: 800;
    color: var(--primary);
    letter-spacing: 0.5px;
    text-transform: uppercase;
    margin-bottom: 2px;
  }
  .subtitle {
    font-size: 9.2pt;
    font-weight: 700;
    color: var(--slate);
    margin-bottom: 4px;
  }
  .contact-bar {
    font-size: 8.2pt;
    color: var(--slate);
    display: flex;
    justify-content: center;
    flex-wrap: wrap;
    gap: 6px 12px;
  }
  .contact-bar a {
    color: var(--accent);
    text-decoration: underline;
    font-weight: 600;
  }
  section {
    margin-bottom: 8px;
  }
  h2 {
    font-size: 9.2pt;
    font-weight: 700;
    color: var(--primary);
    text-transform: uppercase;
    letter-spacing: 0.5px;
    border-bottom: 1px solid var(--border);
    padding-bottom: 2px;
    margin-bottom: 4px;
  }
  p.objective {
    font-size: 8.4pt;
    color: var(--body);
    line-height: 1.32;
    text-align: justify;
  }
  .edu-row {
    font-size: 8.4pt;
    margin-bottom: 2px;
  }
  .edu-title {
    font-weight: 700;
    color: var(--primary);
  }
  .edu-inst {
    font-style: italic;
    color: var(--slate);
  }
  .edu-sub {
    font-size: 8.0pt;
    color: var(--body);
  }
  .item-block {
    margin-bottom: 4px;
  }
  .item-header {
    display: flex;
    justify-content: space-between;
    align-items: baseline;
    font-size: 8.5pt;
  }
  .item-title {
    font-weight: 700;
    color: var(--primary);
  }
  .item-context {
    font-style: italic;
    color: var(--slate);
    font-size: 8.0pt;
  }
  ul {
    list-style-type: none;
    margin-top: 1px;
  }
  li {
    position: relative;
    padding-left: 14px;
    font-size: 8.0pt;
    line-height: 1.28;
    color: var(--body);
    margin-bottom: 1px;
  }
  li::before {
    content: "•";
    position: absolute;
    left: 2px;
    color: var(--slate);
    font-weight: bold;
  }
  .skills-row {
    font-size: 8.0pt;
    line-height: 1.28;
    margin-bottom: 2px;
    padding-left: 14px;
    position: relative;
  }
  .skills-row::before {
    content: "•";
    position: absolute;
    left: 2px;
    color: var(--slate);
    font-weight: bold;
  }
  .skills-row strong {
    color: var(--primary);
  }
  .cert-row {
    font-size: 7.8pt;
    line-height: 1.28;
    color: var(--body);
    padding-left: 14px;
    position: relative;
  }
  .cert-row::before {
    content: "•";
    position: absolute;
    left: 2px;
    color: var(--slate);
    font-weight: bold;
  }
  @media print {
    body { background: #fff; padding: 0; }
    .cv-page { box-shadow: none; padding: 0; max-width: 100%; }
  }
</style>
</head>
<body>
<div class="cv-page">
  <header>
    <h1>ADITYA MEHRA</h1>
    <div class="subtitle">BBA Graduate &nbsp;|&nbsp; Business Operations, AI-Driven Builder &amp; Associate</div>
    <div class="contact-bar">
      <span>Bengaluru, Karnataka, India</span>
      <span>•</span>
      <span><strong>Phone:</strong> <a href="tel:+917003456624">+91 7003456624</a></span>
      <span>•</span>
      <span><strong>Email:</strong> <a href="mailto:adityamehra799@gmail.com">adityamehra799@gmail.com</a></span>
    </div>
    <div class="contact-bar" style="margin-top: 2px;">
      <span><strong>Proof of Work &amp; Portfolio:</strong> <a href="https://adityamehra007.github.io/ADI-OS/" target="_blank">adityamehra007.github.io/ADI-OS</a></span>
      <span>•</span>
      <span><strong>LinkedIn:</strong> <a href="https://www.linkedin.com/in/aditya-mehra-b8644b326" target="_blank">linkedin.com/in/aditya-mehra-b8644b326</a></span>
      <span>•</span>
      <span><strong>GitHub:</strong> <a href="https://github.com/AdityaMehra007" target="_blank">github.com/AdityaMehra007</a></span>
    </div>
  </header>

  <section>
    <h2>Career Objective</h2>
    <p class="objective">
      Energetic and grounded BBA graduate eager to contribute to forward-thinking startups and established MNCs. 
      Combines practical business knowledge with hands-on modern tech skills—actively vibe coding web applications, building live websites, and producing creative video content with AI tools. 
      Brings proven internship experience in corporate client outreach, operations data verification, vendor coordination, and store retail sales. 
      Low-profile attitude, fast learner, and determined to deliver real value and high reliability from day one.
    </p>
  </section>

  <section>
    <h2>Proof of Work &amp; AI Projects</h2>
    
    <div class="item-block">
      <div class="item-header">
        <span class="item-title">Interactive Web Applications &amp; Vibe Coding</span>
        <span class="item-context">Personal Projects &amp; Live Companion (<a href="https://adityamehra007.github.io/ADI-OS/" target="_blank" style="color:var(--accent);">adityamehra007.github.io/ADI-OS</a>)</span>
      </div>
      <ul>
        <li>Used AI coding tools and prompt engineering to vibe code interactive web apps, personal portfolio tools, and live responsive web pages.</li>
        <li>Explored modern UI designs, tested clean web layouts, and deployed working web prototypes to GitHub Pages.</li>
      </ul>
    </div>

    <div class="item-block">
      <div class="item-header">
        <span class="item-title">AI Video Generation &amp; Digital Content Creation</span>
        <span class="item-context">Multimedia Brand Assets &amp; Visual Storytelling</span>
      </div>
      <ul>
        <li>Created engaging video clips, product concept visuals, and presentation materials using cutting-edge AI video and image tools.</li>
        <li>Produced creative content for digital brand showcases, marketing slide decks, and social media mockups.</li>
      </ul>
    </div>
  </section>

  <section>
    <h2>Work Experience &amp; Internships</h2>
    
    <div class="item-block">
      <div class="item-header">
        <span class="item-title">Business Development Intern</span>
        <span class="item-context">Pencil Mark Interior Solutions, Bengaluru</span>
      </div>
      <ul>
        <li>Contacted corporate clients via phone and email to introduce office interior design solutions; prepared client proposals and meeting decks.</li>
        <li>Earned a written appreciation letter from senior leadership for strong dedication, consistent effort, and proactive client support.</li>
      </ul>
    </div>

    <div class="item-block">
      <div class="item-header">
        <span class="item-title">Operations Intern</span>
        <span class="item-context">Instawork Services India, Bengaluru</span>
      </div>
      <ul>
        <li>Supported operations team with collecting, sorting, and verifying data for AI and robotics projects with high attention to detail.</li>
      </ul>
    </div>

    <div class="item-block">
      <div class="item-header">
        <span class="item-title">Stall Coordinator</span>
        <span class="item-context">AERO India Exhibition (Salt in My Coca), Bengaluru</span>
      </div>
      <ul>
        <li>Managed exhibition stall setup, welcomed attendees, addressed customer questions, and maintained complete stock protection.</li>
      </ul>
    </div>

    <div class="item-block">
      <div class="item-header">
        <span class="item-title">Event Coordinator</span>
        <span class="item-context">Commercial &amp; College Events, Bengaluru</span>
      </div>
      <ul>
        <li>Coordinated vendor setups for stage, sound, and lighting, and managed artist travel, hospitality, and backstage arrangements.</li>
      </ul>
    </div>

    <div class="item-block">
      <div class="item-header">
        <span class="item-title">Store Assistant</span>
        <span class="item-context">Family Retail Store &amp; Food Business, Bengaluru</span>
      </div>
      <ul>
        <li>Handled daily customer billing, counter sales, inventory register updates, and supplier follow-ups for fresh materials.</li>
      </ul>
    </div>
  </section>

  <section>
    <h2>Education</h2>
    <div class="edu-row">
      <span class="edu-title">Bachelor of Business Administration (BBA) — International Business</span> &nbsp;|&nbsp;
      <span class="edu-inst">Dayananda Sagar University, Bengaluru</span>
    </div>
    <div class="edu-sub">Key Subjects: International Trade, Marketing Management, Supply Chain, Business Communication, Accounting.</div>
  </section>

  <section>
    <h2>Key Skills</h2>
    <div class="skills-row">
      <strong>AI &amp; Digital Builder Tools:</strong> AI Vibe Coding, AI Video Creation, ChatGPT, Claude, Prompt Engineering, HTML/Web Basics, GitHub.
    </div>
    <div class="skills-row">
      <strong>Computer &amp; Office Tools:</strong> MS Excel, MS Word, MS PowerPoint, Google Workspace (Sheets, Docs), Online Research, Data Entry.
    </div>
    <div class="skills-row">
      <strong>Business &amp; Workplace Skills:</strong> Client Communication, Customer Service, Operations Support, Vendor Coordination, Teamwork, Problem Solving.
    </div>
    <div class="skills-row">
      <strong>Languages:</strong> English, Hindi.
    </div>
  </section>

  <section>
    <h2>Certificates &amp; Achievements</h2>
    <div class="cert-row">
      Appreciation Letter for Good Performance (Pencil Mark) &nbsp;|&nbsp;
      Digital Marketing &amp; E-commerce (Google) &nbsp;|&nbsp;
      Service Marketing (NPTEL, IIT Kharagpur) &nbsp;|&nbsp;
      Artificial Intelligence Tools Workshop (Outskill)
    </div>
  </section>
</div>
</body>
</html>
"""
    out_html = BASE_DIR / filename
    out_html.write_text(html_content, encoding="utf-8")
    return out_html

def generate_markdown_version(filename="ADITYA_MEHRA_STANDOUT_FRESHER_CV.md"):
    md_content = """# ADITYA MEHRA
**BBA Graduate | Business Operations, AI-Driven Builder & Associate**  
Bengaluru, Karnataka, India | Phone: [+91 7003456624](tel:+917003456624) | Email: [adityamehra799@gmail.com](mailto:adityamehra799@gmail.com)  
**Proof of Work & Portfolio:** [adityamehra007.github.io/ADI-OS](https://adityamehra007.github.io/ADI-OS/) | **LinkedIn:** [linkedin.com/in/aditya-mehra-b8644b326](https://www.linkedin.com/in/aditya-mehra-b8644b326) | **GitHub:** [github.com/AdityaMehra007](https://github.com/AdityaMehra007)

---

## CAREER OBJECTIVE
Energetic and grounded BBA graduate eager to contribute to forward-thinking startups and established MNCs. Combines practical business knowledge with hands-on modern tech skills—actively vibe coding web applications, building live websites, and producing creative video content with AI tools. Brings proven internship experience in corporate client outreach, operations data verification, vendor coordination, and store retail sales. Low-profile attitude, fast learner, and determined to deliver real value and high reliability from day one.

---

## PROOF OF WORK & AI PROJECTS

### Interactive Web Applications & Vibe Coding | Personal Projects & Live Companion (adityamehra007.github.io/ADI-OS)
- Used AI coding tools and prompt engineering to vibe code interactive web apps, personal portfolio tools, and live responsive web pages.
- Explored modern UI designs, tested clean web layouts, and deployed working web prototypes to GitHub Pages.

### AI Video Generation & Digital Content Creation | Multimedia Brand Assets & Visual Storytelling
- Created engaging video clips, product concept visuals, and presentation materials using cutting-edge AI video and image tools.
- Produced creative content for digital brand showcases, marketing slide decks, and social media mockups.

---

## WORK EXPERIENCE & INTERNSHIPS

### Business Development Intern | Pencil Mark Interior Solutions, Bengaluru
- Contacted corporate clients via phone and email to introduce office interior design solutions; prepared client proposals and meeting decks.
- Earned a written appreciation letter from senior leadership for strong dedication, consistent effort, and proactive client support.

### Operations Intern | Instawork Services India, Bengaluru
- Supported operations team with collecting, sorting, and verifying data for AI and robotics projects with high attention to detail.

### Stall Coordinator | AERO India Exhibition (Salt in My Coca), Bengaluru
- Managed exhibition stall setup, welcomed attendees, addressed customer questions, and maintained complete stock protection.

### Event Coordinator | Commercial & College Events, Bengaluru
- Coordinated vendor setups for stage, sound, and lighting, and managed artist travel, hospitality, and backstage arrangements.

### Store Assistant | Family Retail Store & Food Business, Bengaluru
- Handled daily customer billing, counter sales, inventory register updates, and supplier follow-ups for fresh materials.

---

## EDUCATION
**Bachelor of Business Administration (BBA) — International Business** | *Dayananda Sagar University, Bengaluru*  
- **Key Subjects:** International Trade, Marketing Management, Supply Chain, Business Communication, Accounting.

---

## KEY SKILLS
- **AI & Digital Builder Tools:** AI Vibe Coding, AI Video Creation, ChatGPT, Claude, Prompt Engineering, HTML/Web Basics, GitHub.
- **Computer & Office Tools:** MS Excel, MS Word, MS PowerPoint, Google Workspace (Sheets, Docs), Online Research, Data Entry.
- **Business & Workplace Skills:** Client Communication, Customer Service, Operations Support, Vendor Coordination, Teamwork, Problem Solving.
- **Languages:** English, Hindi.

---

## CERTIFICATES & ACHIEVEMENTS
- Appreciation Letter for Good Performance (Pencil Mark Interior Solutions)
- Digital Marketing & E-commerce Certificate (Google)
- Service Marketing Certificate (NPTEL, IIT Kharagpur)
- Artificial Intelligence Tools Workshop Certificate (Outskill)
"""
    out_md = BASE_DIR / filename
    out_md.write_text(md_content, encoding="utf-8")
    return out_md

if __name__ == "__main__":
    # 1. Build DOCX
    docx_file = create_standout_cv("ADITYA_MEHRA_STANDOUT_FRESHER_CV.docx")
    print(f"Created: {docx_file}")

    # Also build to ADITYA_MEHRA_PERFECT_FRESHER_CV.docx and ADITYA_MEHRA_SIMPLE_ENGLISH_CV.docx
    create_standout_cv("ADITYA_MEHRA_PERFECT_FRESHER_CV.docx")
    create_standout_cv("ADITYA_MEHRA_SIMPLE_ENGLISH_CV.docx")
    print("Updated all standard CV docx files.")

    # 2. Build HTML
    html_file = generate_html_version("ADITYA_MEHRA_STANDOUT_FRESHER_CV.html")
    generate_html_version("ADITYA_MEHRA_PERFECT_FRESHER_CV.html")
    generate_html_version("ADITYA_MEHRA_SIMPLE_ENGLISH_CV.html")
    print(f"Created HTML: {html_file}")

    # 3. Build Markdown
    md_file = generate_markdown_version("ADITYA_MEHRA_STANDOUT_FRESHER_CV.md")
    generate_markdown_version("ADITYA_MEHRA_PERFECT_FRESHER_CV.md")
    generate_markdown_version("ADITYA_MEHRA_SIMPLE_ENGLISH_CV.md")
    print(f"Created MD: {md_file}")

    # 4. Use Word COM to inspect page count and export PDF
    print("Opening Word COM to verify page count and export PDFs...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False

    for target_name in ["ADITYA_MEHRA_STANDOUT_FRESHER_CV", "ADITYA_MEHRA_PERFECT_FRESHER_CV", "ADITYA_MEHRA_SIMPLE_ENGLISH_CV"]:
        target_docx = BASE_DIR / f"{target_name}.docx"
        target_pdf = BASE_DIR / f"{target_name}.pdf"
        doc = word.Documents.Open(str(target_docx))
        pages = doc.ComputeStatistics(2)
        print(f"[{target_name}] Exact Page Count: {pages}")
        doc.ExportAsFixedFormat(str(target_pdf), 17) # 17 = wdExportFormatPDF
        doc.Close()
        print(f"[{target_name}] Exported PDF: {target_pdf}")

    word.Quit()
    print("ALL BUILDS AND VERIFICATIONS COMPLETED SUCCESSFULLY!")
