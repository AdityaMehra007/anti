import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client
import pypdf
from pathlib import Path

BASE_DIR = Path(r"e:\anti")

C_NAVY = RGBColor(15, 23, 42)
C_SLATE = RGBColor(51, 65, 85)
C_BODY = RGBColor(30, 41, 59)
C_MUTED = RGBColor(100, 116, 139)
C_LINK = "1E3A8A"

def add_clean_divider(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CBD5E1')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_hyperlink_run(paragraph, url, display_text, font_size_pt=8.8, color_hex="1E3A8A", bold=False, underline=True):
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'<w:color w:val="{color_hex}"/>'
        f'{"<w:b/>" if bold else ""}'
        f'{"<w:u w:val=\"single\"/>" if underline else ""}'
        f'<w:sz w:val="{int(font_size_pt * 2)}"/>'
        f'</w:rPr>'
        f'<w:t>{display_text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    paragraph._p.append(hyperlink)

def build_complete_cv(body_pt=8.8, head_pt=10.0, sec_b=8.5, bullet_a=2.8, line_sp=1.12, top_m=0.38, bot_m=0.38):
    doc = Document()
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(top_m)
        s.bottom_margin = Inches(bot_m)
        s.left_margin = Inches(0.48)
        s.right_margin = Inches(0.48)

    # 1. HEADER
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1.5)
    r = p_name.add_run("ADITYA MEHRA")
    r.font.name = "Calibri"; r.font.size = Pt(21.5); r.bold = True; r.font.color.rgb = C_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2.5)
    r = p_sub.add_run("BBA Graduate  •  Operations, Client Support & AI Productivity  •  Entry-Level")
    r.font.name = "Calibri"; r.font.size = Pt(9.5); r.bold = True; r.font.color.rgb = C_SLATE

    p_con1 = doc.add_paragraph()
    p_con1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con1.paragraph_format.space_before = Pt(0)
    p_con1.paragraph_format.space_after = Pt(1.5)
    r = p_con1.add_run("Bengaluru, Karnataka, India  •  ")
    r.font.name = "Calibri"; r.font.size = Pt(8.8); r.font.color.rgb = C_MUTED
    r = p_con1.add_run("Phone: ")
    r.font.name = "Calibri"; r.font.size = Pt(8.8); r.bold = True; r.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "tel:+917003456624", "+91 7003456624", font_size_pt=8.8, underline=False)
    r = p_con1.add_run("  •  Email: ")
    r.font.name = "Calibri"; r.font.size = Pt(8.8); r.bold = True; r.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "mailto:adityamehra799@gmail.com", "adityamehra799@gmail.com", font_size_pt=8.8, underline=False)

    p_con2 = doc.add_paragraph()
    p_con2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con2.paragraph_format.space_before = Pt(0)
    p_con2.paragraph_format.space_after = Pt(2.8)
    r = p_con2.add_run("Portfolio: ")
    r.font.name = "Calibri"; r.font.size = Pt(8.8); r.bold = True; r.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://adi-digital-universe.ai.studio/", "adi-digital-universe.ai.studio", font_size_pt=8.8, underline=True)
    r = p_con2.add_run("  •  LinkedIn: ")
    r.font.name = "Calibri"; r.font.size = Pt(8.8); r.bold = True; r.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://www.linkedin.com/in/aditya-mehra-b8644b326", "linkedin.com/in/aditya-mehra-b8644b326", font_size_pt=8.8, underline=True)
    r = p_con2.add_run("  •  GitHub: ")
    r.font.name = "Calibri"; r.font.size = Pt(8.8); r.bold = True; r.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://github.com/AdityaMehra007", "github.com/AdityaMehra007", font_size_pt=8.8, underline=True)

    def add_sec(title, before=sec_b):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title.upper())
        r.font.name = "Calibri"; r.font.size = Pt(head_pt); r.bold = True; r.font.color.rgb = C_NAVY
        add_clean_divider(p)

    def add_bullet(p, title_bold, text_body, space_after=bullet_a):
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_before = Pt(0.4)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_sp
        rb = p.add_run("• ")
        rb.font.name = "Calibri"; rb.font.size = Pt(body_pt); rb.bold = True; rb.font.color.rgb = C_SLATE
        if title_bold:
            rt = p.add_run(title_bold)
            rt.font.name = "Calibri"; rt.font.size = Pt(body_pt); rt.bold = True; rt.font.color.rgb = C_NAVY
        rc = p.add_run(text_body)
        rc.font.name = "Calibri"; rc.font.size = Pt(body_pt); rc.font.color.rgb = C_BODY

    # 2. CAREER OBJECTIVE
    add_sec("Career Objective", before=1.8)
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(0.8)
    p_obj.paragraph_format.space_after = Pt(1.8)
    p_obj.paragraph_format.line_spacing = line_sp
    r = p_obj.add_run(
        "Dedicated BBA graduate in International Business with hands-on experience across client outreach, operational coordination, and modern AI productivity workflows. "
        "Quick learner with strong computer proficiency, eager to support cross-functional teams in daily operations, customer service, vendor communication, and business administration."
    )
    r.font.name = "Calibri"; r.font.size = Pt(body_pt); r.font.color.rgb = C_BODY

    # 3. WORK EXPERIENCE & INTERNSHIPS (All 6 roles)
    add_sec("Work Experience & Internships")
    
    # 1. Pencil Mark Solutions
    p = doc.add_paragraph()
    add_bullet(p, "Business Development Intern | Pencil Mark Solutions: ",
               "Conducted targeted outreach to prospective corporate clients via phone and email to introduce commercial interior design services.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Assisted senior management with preparing service proposals, cost estimates, client meeting presentations, and follow-up schedules.")

    # 2. Instawork Services India
    p = doc.add_paragraph()
    add_bullet(p, "Operations Intern | Instawork Services India: ",
               "Organized, structured, and verified computer datasets to support engineering teams with robotics and AI operations.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Maintained strict data accuracy standards by cross-checking spreadsheets, documenting errors, and reporting workflow discrepancies promptly.")

    # 3. AERO India Exhibition
    p = doc.add_paragraph()
    add_bullet(p, "Stall Coordinator | AERO India Exhibition: ",
               "Managed exhibition booth setup, welcomed trade delegates and international visitors, answered inquiries, and maintained inventory control.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Monitored promotional merchandise and product catalogues under high footfall to ensure continuous availability throughout the event.")

    # 4. College & Music Events
    p = doc.add_paragraph()
    add_bullet(p, "Event Coordinator | College & Music Events: ",
               "Coordinated with technical vendors for stage, sound, and lighting setups while managing hospitality and travel logistics for performing artists.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Supervised on-ground event volunteers and maintained strict schedules to ensure seamless show operations and crowd management.")

    # 5. Store Assistant | Family Retail Store
    p = doc.add_paragraph()
    add_bullet(p, "Store Assistant | Family Retail Store: ",
               "Managed daily counter sales, customer billing, cash registers, and supplier restocking to maintain seamless store operations.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Communicated regularly with wholesale distributors to place repeat merchandise orders and verify incoming delivery receipts.")

    # 6. Community Operations | Social Initiative
    p = doc.add_paragraph()
    add_bullet(p, "Community Operations Volunteer | Social Initiative: ",
               "Coordinated volunteer shift schedules, attendee check-ins, and on-ground logistics for community educational drives in Bengaluru.")

    # 4. PROJECTS & AI PROOF OF WORK (All 3 Projects)
    add_sec("Projects & AI Proof of Work")
    p = doc.add_paragraph()
    add_bullet(p, "Web Apps & Digital Portfolio (adi-digital-universe.ai.studio): ",
               "Built and deployed interactive web applications and a personal portfolio website using modern AI coding workflows and cloud hosting.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Maintained active public GitHub repositories showcasing code implementations, interface prototypes, and practical automation scripts.")

    p = doc.add_paragraph()
    add_bullet(p, "Operations Automation & Workflow Tools: ",
               "Prototyped operations automation scripts, interactive status dashboards, and digital checklists to eliminate repetitive manual tracking.")

    p = doc.add_paragraph()
    add_bullet(p, "AI Multimedia & Content Production: ",
               "Produced promotional video assets, visual marketing collateral, and digital presentation decks using modern generative AI tools.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Accelerated creative content delivery by combining AI image generation workflows, video storyboard drafting, and prompt refinement.")

    # 5. EDUCATION
    add_sec("Education")
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(0.8)
    p_edu.paragraph_format.space_after = Pt(1.5)
    p_edu.paragraph_format.line_spacing = line_sp
    p_edu.paragraph_format.keep_with_next = True
    r = p_edu.add_run("Bachelor of Business Administration (BBA) — International Business (2023 – 2026)  •  ")
    r.font.name = "Calibri"; r.font.size = Pt(body_pt + 0.2); r.bold = True; r.font.color.rgb = C_NAVY
    r = p_edu.add_run("Dayananda Sagar University, Bengaluru\n")
    r.font.name = "Calibri"; r.font.size = Pt(body_pt); r.font.italic = True; r.font.color.rgb = C_SLATE
    r = p_edu.add_run("Relevant Coursework: International Trade, Marketing Management, Supply Chain Logistics, Business Communication, Accounting.")
    r.font.name = "Calibri"; r.font.size = Pt(body_pt - 0.4); r.font.color.rgb = C_BODY

    # 6. KEY SKILLS & COMPETENCIES
    add_sec("Key Skills & Competencies")
    skills = [
        ("Computer & AI Tools: ", "ChatGPT, Generative AI Platforms, MS Excel (Formulas & Data Verification), MS Word, MS PowerPoint, Google Workspace, GitHub, Web Research."),
        ("Workplace Competencies: ", "Client Communication, Customer Relations, Operations Support, Vendor Coordination, Team Collaboration, Problem Solving, Time Management."),
        ("Languages: ", "English, Hindi.")
    ]
    for s_title, s_desc in skills:
        p = doc.add_paragraph()
        add_bullet(p, s_title, s_desc, space_after=2.0)

    # 7. CERTIFICATES & HONORS
    add_sec("Certificates & Honors")
    certs = [
        ("Internship Certificate: ", "Pencil Mark Solutions (Client Outreach & Business Operations)."),
        ("Fundamentals of Digital Marketing: ", "Google (Online Marketing, SEO & Web Analytics)."),
        ("Marketing Management Certification: ", "NPTEL, IIT Kharagpur (Marketing Principles & Consumer Behavior)."),
        ("AI Productivity Workshop: ", "Outskill (Generative AI Workflows & Modern Office Automation).")
    ]
    for c_title, c_desc in certs:
        p = doc.add_paragraph()
        add_bullet(p, c_title, c_desc, space_after=1.5)

    test_docx = BASE_DIR / "test_complete_all.docx"
    doc.save(test_docx)
    return test_docx

if __name__ == "__main__":
    for top_m, bot_m, body_pt, head_pt, sec_b, bullet_a, line_sp in [
        (0.38, 0.38, 9.0, 10.2, 8.5, 2.8, 1.14),
        (0.38, 0.38, 8.9, 10.0, 8.0, 2.5, 1.12),
        (0.35, 0.35, 8.9, 10.0, 8.5, 2.8, 1.13),
        (0.35, 0.35, 9.0, 10.0, 9.0, 3.0, 1.14),
    ]:
        docx_file = build_complete_cv(body_pt, head_pt, sec_b, bullet_a, line_sp, top_m, bot_m)
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(docx_file))
        pages = doc.ComputeStatistics(2)
        pdf_file = BASE_DIR / "test_complete_all.pdf"
        doc.ExportAsFixedFormat(str(pdf_file), 17)
        doc.Close()
        word.Quit()
        
        reader = pypdf.PdfReader(str(pdf_file))
        p_cnt = len(reader.pages)
        
        y_positions = []
        def visitor(text, cm, tm, fontDict, fontSize):
            if text.strip():
                y_positions.append(tm[5])
        reader.pages[0].extract_text(visitor_text=visitor)
        bottom_y = min(y_positions)
        print(f"top_m={top_m}, body_pt={body_pt}, sec_b={sec_b}, bullet_a={bullet_a}, line_sp={line_sp} => Word: {pages}, PyPDF: {p_cnt}, Bottom y: {bottom_y:.1f} pt (rem: {bottom_y/72.0:.2f} in / {bottom_y/841.68*100:.1f}%)")
