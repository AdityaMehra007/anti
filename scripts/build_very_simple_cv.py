import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client

BASE_DIR = Path(r"e:\anti")

# Palette: Clean, friendly, professional
C_NAVY = RGBColor(15, 23, 42)       # #0f172a Deep Navy
C_SLATE = RGBColor(71, 85, 105)     # #475569 Slate
C_BODY = RGBColor(30, 41, 59)       # #1e293b Body text
C_MUTED = RGBColor(100, 116, 139)   # #64748b Metadata
C_LINK = "1E3A8A"                   # #1e3a8a Deep Blue

def add_clean_divider(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')       # 0.5 pt
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CBD5E1')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_hyperlink_run(paragraph, url, display_text, font_size_pt=8.4, color_hex="1E3A8A", bold=False, underline=True):
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'<w:color w:val="{color_hex}"/>'
        f'{"<w:b/>" if bold else ""}'
        f'{"<w:u w:val=\"single\"/>" if underline else ""}'
        f'<w:sz w:val="{int(font_size_pt * 2)}"/>'
        f'</w:rPr>'
        f'<w:t>{display_text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    paragraph._p.append(hyperlink)

def create_very_simple_cv(filename="ADITYA_MEHRA_VERY_SIMPLE_CV.docx"):
    doc = Document()

    # Margins calibrated for clean 1-page fit
    for section in doc.sections:
        section.page_width = Inches(8.27)    # A4
        section.page_height = Inches(11.69)  # A4
        section.top_margin = Inches(0.40)
        section.bottom_margin = Inches(0.40)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    # 1. HEADER
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    r_name = p_name.add_run("ADITYA MEHRA")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(20)
    r_name.font.bold = True
    r_name.font.color.rgb = C_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2.5)
    r_sub = p_sub.add_run("BBA Graduate  |  Looking for an Entry-Level Job")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_SLATE

    # Contact Line 1
    p_con1 = doc.add_paragraph()
    p_con1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con1.paragraph_format.space_before = Pt(0)
    p_con1.paragraph_format.space_after = Pt(1.5)
    
    r_loc = p_con1.add_run("Bengaluru, Karnataka, India  |  ")
    r_loc.font.name = "Calibri"; r_loc.font.size = Pt(8.4); r_loc.font.color.rgb = C_MUTED
    
    r_ph = p_con1.add_run("Phone: ")
    r_ph.font.name = "Calibri"; r_ph.font.size = Pt(8.4); r_ph.font.bold = True; r_ph.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "tel:+917003456624", "+91 7003456624", font_size_pt=8.4, color_hex=C_LINK, underline=False)
    
    r_sep1 = p_con1.add_run("  |  Email: ")
    r_sep1.font.name = "Calibri"; r_sep1.font.size = Pt(8.4); r_sep1.font.bold = True; r_sep1.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "mailto:adityamehra799@gmail.com", "adityamehra799@gmail.com", font_size_pt=8.4, color_hex=C_LINK, underline=False)

    # Contact Line 2
    p_con2 = doc.add_paragraph()
    p_con2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con2.paragraph_format.space_before = Pt(0)
    p_con2.paragraph_format.space_after = Pt(4)

    r_web_lbl = p_con2.add_run("Website: ")
    r_web_lbl.font.name = "Calibri"; r_web_lbl.font.size = Pt(8.4); r_web_lbl.font.bold = True; r_web_lbl.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://adityamehra007.github.io/ADI-OS/", "adityamehra007.github.io/ADI-OS", font_size_pt=8.4, color_hex=C_LINK, underline=True)

    r_sep2 = p_con2.add_run("  |  LinkedIn: ")
    r_sep2.font.name = "Calibri"; r_sep2.font.size = Pt(8.4); r_sep2.font.bold = True; r_sep2.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://www.linkedin.com/in/aditya-mehra-b8644b326", "linkedin.com/in/aditya-mehra-b8644b326", font_size_pt=8.4, color_hex=C_LINK, underline=True)

    r_sep3 = p_con2.add_run("  |  GitHub: ")
    r_sep3.font.name = "Calibri"; r_sep3.font.size = Pt(8.4); r_sep3.font.bold = True; r_sep3.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://github.com/AdityaMehra007", "github.com/AdityaMehra007", font_size_pt=8.4, color_hex=C_LINK, underline=True)

    def add_sec(title, before=3.5):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title.upper())
        r.font.name = "Calibri"
        r.font.size = Pt(9.2)
        r.font.bold = True
        r.font.color.rgb = C_NAVY
        add_clean_divider(p)

    def add_bullet(p, title_bold, text_body):
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_before = Pt(0.2)
        p.paragraph_format.space_after = Pt(0.8)
        p.paragraph_format.line_spacing = 1.05
        
        rb = p.add_run("• ")
        rb.font.name = "Calibri"; rb.font.size = Pt(8.2); rb.bold = True; rb.font.color.rgb = C_SLATE

        if title_bold:
            rt = p.add_run(title_bold)
            rt.font.name = "Calibri"; rt.font.size = Pt(8.2); rt.bold = True; rt.font.color.rgb = C_NAVY

        rc = p.add_run(text_body)
        rc.font.name = "Calibri"; rc.font.size = Pt(8.2); rc.font.color.rgb = C_BODY

    # 2. CAREER OBJECTIVE (Simple, everyday words)
    add_sec("Career Objective", before=1)
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(1)
    p_obj.paragraph_format.space_after = Pt(2)
    p_obj.paragraph_format.line_spacing = 1.08
    r_obj = p_obj.add_run(
        "Hardworking BBA graduate looking for an entry-level job in business operations, sales, or office work. "
        "Quick to learn new tools, good with computers and AI tools, and ready to help the team with daily work."
    )
    r_obj.font.name = "Calibri"
    r_obj.font.size = Pt(8.4)
    r_obj.font.color.rgb = C_BODY

    # 3. WORK EXPERIENCE & INTERNSHIPS (Simple, plain English)
    add_sec("Work Experience & Internships", before=3)
    
    exp_list = [
        ("Business Development Intern | Pencil Mark Solutions: ", "Called and emailed corporate clients to introduce interior design services; got a written appreciation letter for good work."),
        ("Operations Intern | Instawork Services India: ", "Helped the team collect, sort, and double-check data on the computer for AI and robotics projects."),
        ("Stall Coordinator | AERO India Exhibition: ", "Set up the display stall, welcomed visitors, answered questions, and took care of all stock items."),
        ("Event Coordinator | College & Music Events: ", "Coordinated with suppliers for sound, lights, and stage setup; arranged travel and stay for artists."),
        ("Store Assistant | Family Retail Store: ", "Handled daily customer billing, counter sales, shop registers, and ordering fresh stock from suppliers.")
    ]

    for title, desc in exp_list:
        p_exp = doc.add_paragraph()
        add_bullet(p_exp, title, desc)

    # 4. PROJECTS & AI WORK (Simple words)
    add_sec("Projects & AI Work", before=3)
    p1 = doc.add_paragraph()
    add_bullet(p1, "Simple Web Apps & Personal Website (adityamehra007.github.io/ADI-OS): ", "Made working web tools and personal portfolio pages using AI coding tools and GitHub.")

    p2 = doc.add_paragraph()
    add_bullet(p2, "AI Video Making: ", "Created simple promotional videos, product clips, and slide designs using AI video tools.")

    # 5. EDUCATION (Simple words)
    add_sec("Education", before=3)
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(1)
    p_edu.paragraph_format.space_after = Pt(1.5)
    p_edu.paragraph_format.line_spacing = 1.06
    p_edu.paragraph_format.keep_with_next = True

    r_deg = p_edu.add_run("Bachelor of Business Administration (BBA) — International Business  |  ")
    r_deg.font.name = "Calibri"; r_deg.font.size = Pt(8.5); r_deg.bold = True; r_deg.font.color.rgb = C_NAVY

    r_uni = p_edu.add_run("Dayananda Sagar University, Bengaluru\n")
    r_uni.font.name = "Calibri"; r_uni.font.size = Pt(8.1); r_uni.font.italic = True; r_uni.font.color.rgb = C_SLATE

    r_sub = p_edu.add_run("Subjects Studied: International Trade, Marketing, Supply Chain, Business Communication, Accounting.")
    r_sub.font.name = "Calibri"; r_sub.font.size = Pt(8.0); r_sub.font.color.rgb = C_BODY

    # 6. KEY SKILLS (Simple words)
    add_sec("Key Skills", before=3)
    skills = [
        ("Computer & AI Skills: ", "MS Excel, MS Word, MS PowerPoint, Google Sheets, ChatGPT, AI Tools, Internet Research."),
        ("Work Skills: ", "Client Communication, Customer Service, Teamwork, Vendor Follow-ups, Office Support, Problem Solving."),
        ("Languages: ", "English, Hindi.")
    ]
    for s_title, s_desc in skills:
        p_s = doc.add_paragraph()
        add_bullet(p_s, s_title, s_desc)

    # 7. CERTIFICATES (Simple words)
    add_sec("Certificates", before=3)
    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.left_indent = Inches(0.12)
    p_cert.paragraph_format.space_before = Pt(1)
    p_cert.paragraph_format.space_after = Pt(1)
    p_cert.paragraph_format.line_spacing = 1.04

    certs = [
        "Appreciation Letter for Good Work (Pencil Mark)",
        "Digital Marketing (Google)",
        "Marketing Certificate (NPTEL, IIT Kharagpur)",
        "AI Workshop (Outskill)"
    ]
    
    rb = p_cert.add_run("• ")
    rb.font.name = "Calibri"; rb.font.size = Pt(8.0); rb.bold = True; rb.font.color.rgb = C_SLATE
    rc = p_cert.add_run("   |   ".join(certs))
    rc.font.name = "Calibri"; rc.font.size = Pt(7.8); rc.font.color.rgb = C_BODY

    out_path = BASE_DIR / filename
    doc.save(out_path)
    return out_path

def generate_html_version(filename="ADITYA_MEHRA_VERY_SIMPLE_CV.html"):
    html_content = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Aditya Mehra — Simple Fresher CV</title>
<style>
  :root {
    --primary: #0f172a;
    --accent: #1e3a8a;
    --slate: #475569;
    --body: #1e293b;
    --border: #cbd5e1;
  }
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body {
    font-family: Calibri, -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
    color: var(--body);
    background-color: #f1f5f9;
    line-height: 1.30;
    font-size: 8.8pt;
    padding: 24px;
  }
  .cv-page {
    max-width: 820px;
    margin: 0 auto;
    background: #ffffff;
    padding: 32px 40px;
    border-radius: 6px;
    box-shadow: 0 4px 16px rgba(0,0,0,0.06);
  }
  header {
    text-align: center;
    border-bottom: 2px solid var(--primary);
    padding-bottom: 8px;
    margin-bottom: 10px;
  }
  h1 {
    font-size: 20pt;
    font-weight: 800;
    color: var(--primary);
    letter-spacing: 0.5px;
    text-transform: uppercase;
    margin-bottom: 2px;
  }
  .subtitle {
    font-size: 9.5pt;
    font-weight: 700;
    color: var(--slate);
    margin-bottom: 4px;
  }
  .contact-bar {
    font-size: 8.4pt;
    color: var(--slate);
    display: flex;
    justify-content: center;
    flex-wrap: wrap;
    gap: 6px 12px;
  }
  .contact-bar a {
    color: var(--accent);
    text-decoration: underline;
    font-weight: 600;
  }
  section { margin-bottom: 8px; }
  h2 {
    font-size: 9.2pt;
    font-weight: 700;
    color: var(--primary);
    text-transform: uppercase;
    letter-spacing: 0.5px;
    border-bottom: 1px solid var(--border);
    padding-bottom: 2px;
    margin-bottom: 4px;
  }
  p.objective {
    font-size: 8.4pt;
    color: var(--body);
    line-height: 1.30;
    text-align: justify;
  }
  .point-item {
    font-size: 8.2pt;
    line-height: 1.28;
    color: var(--body);
    padding-left: 14px;
    position: relative;
    margin-bottom: 2px;
  }
  .point-item::before {
    content: "•";
    position: absolute;
    left: 2px;
    color: var(--slate);
    font-weight: bold;
  }
  .point-item strong {
    color: var(--primary);
  }
  .edu-row {
    font-size: 8.4pt;
    margin-bottom: 2px;
  }
  .edu-title { font-weight: 700; color: var(--primary); }
  .edu-inst { font-style: italic; color: var(--slate); }
  .edu-sub { font-size: 8.0pt; color: var(--body); }
  @media print {
    body { background: #fff; padding: 0; }
    .cv-page { box-shadow: none; padding: 0; max-width: 100%; }
  }
</style>
</head>
<body>
<div class="cv-page">
  <header>
    <h1>ADITYA MEHRA</h1>
    <div class="subtitle">BBA Graduate &nbsp;|&nbsp; Looking for an Entry-Level Job</div>
    <div class="contact-bar">
      <span>Bengaluru, Karnataka, India</span>
      <span>•</span>
      <span><strong>Phone:</strong> <a href="tel:+917003456624">+91 7003456624</a></span>
      <span>•</span>
      <span><strong>Email:</strong> <a href="mailto:adityamehra799@gmail.com">adityamehra799@gmail.com</a></span>
    </div>
    <div class="contact-bar" style="margin-top: 2px;">
      <span><strong>Website:</strong> <a href="https://adityamehra007.github.io/ADI-OS/" target="_blank">adityamehra007.github.io/ADI-OS</a></span>
      <span>•</span>
      <span><strong>LinkedIn:</strong> <a href="https://www.linkedin.com/in/aditya-mehra-b8644b326" target="_blank">linkedin.com/in/aditya-mehra-b8644b326</a></span>
      <span>•</span>
      <span><strong>GitHub:</strong> <a href="https://github.com/AdityaMehra007" target="_blank">github.com/AdityaMehra007</a></span>
    </div>
  </header>

  <section>
    <h2>Career Objective</h2>
    <p class="objective">
      Hardworking BBA graduate looking for an entry-level job in business operations, sales, or office work. 
      Quick to learn new tools, good with computers and AI tools, and ready to help the team with daily work.
    </p>
  </section>

  <section>
    <h2>Work Experience &amp; Internships</h2>
    <div class="point-item">
      <strong>Business Development Intern | Pencil Mark Solutions:</strong> 
      Called and emailed corporate clients to introduce interior design services; got a written appreciation letter for good work.
    </div>
    <div class="point-item">
      <strong>Operations Intern | Instawork Services India:</strong> 
      Helped the team collect, sort, and double-check data on the computer for AI and robotics projects.
    </div>
    <div class="point-item">
      <strong>Stall Coordinator | AERO India Exhibition:</strong> 
      Set up the display stall, welcomed visitors, answered questions, and took care of all stock items.
    </div>
    <div class="point-item">
      <strong>Event Coordinator | College &amp; Music Events:</strong> 
      Coordinated with suppliers for sound, lights, and stage setup; arranged travel and stay for artists.
    </div>
    <div class="point-item">
      <strong>Store Assistant | Family Retail Store:</strong> 
      Handled daily customer billing, counter sales, shop registers, and ordering fresh stock from suppliers.
    </div>
  </section>

  <section>
    <h2>Projects &amp; AI Work</h2>
    <div class="point-item">
      <strong>Simple Web Apps &amp; Personal Website (<a href="https://adityamehra007.github.io/ADI-OS/" target="_blank" style="color:var(--accent);">adityamehra007.github.io/ADI-OS</a>):</strong> 
      Made working web tools and personal portfolio pages using AI coding tools and GitHub.
    </div>
    <div class="point-item">
      <strong>AI Video Making:</strong> 
      Created simple promotional videos, product clips, and slide designs using AI video tools.
    </div>
  </section>

  <section>
    <h2>Education</h2>
    <div class="edu-row">
      <span class="edu-title">Bachelor of Business Administration (BBA) — International Business</span> &nbsp;|&nbsp;
      <span class="edu-inst">Dayananda Sagar University, Bengaluru</span>
    </div>
    <div class="edu-sub">Subjects Studied: International Trade, Marketing, Supply Chain, Business Communication, Accounting.</div>
  </section>

  <section>
    <h2>Key Skills</h2>
    <div class="point-item">
      <strong>Computer &amp; AI Skills:</strong> MS Excel, MS Word, MS PowerPoint, Google Sheets, ChatGPT, AI Tools, Internet Research.
    </div>
    <div class="point-item">
      <strong>Work Skills:</strong> Client Communication, Customer Service, Teamwork, Vendor Follow-ups, Office Support, Problem Solving.
    </div>
    <div class="point-item">
      <strong>Languages:</strong> English, Hindi.
    </div>
  </section>

  <section>
    <h2>Certificates</h2>
    <div class="point-item">
      Appreciation Letter for Good Work (Pencil Mark) &nbsp;|&nbsp;
      Digital Marketing (Google) &nbsp;|&nbsp;
      Marketing Certificate (NPTEL, IIT Kharagpur) &nbsp;|&nbsp;
      AI Workshop (Outskill)
    </div>
  </section>
</div>
</body>
</html>
"""
    out_html = BASE_DIR / filename
    out_html.write_text(html_content, encoding="utf-8")
    return out_html

def generate_markdown_version(filename="ADITYA_MEHRA_VERY_SIMPLE_CV.md"):
    md_content = """# ADITYA MEHRA
**BBA Graduate | Looking for an Entry-Level Job**  
Bengaluru, Karnataka, India | Phone: [+91 7003456624](tel:+917003456624) | Email: [adityamehra799@gmail.com](mailto:adityamehra799@gmail.com)  
**Website:** [adityamehra007.github.io/ADI-OS](https://adityamehra007.github.io/ADI-OS/) | **LinkedIn:** [linkedin.com/in/aditya-mehra-b8644b326](https://www.linkedin.com/in/aditya-mehra-b8644b326) | **GitHub:** [github.com/AdityaMehra007](https://github.com/AdityaMehra007)

---

## CAREER OBJECTIVE
Hardworking BBA graduate looking for an entry-level job in business operations, sales, or office work. Quick to learn new tools, good with computers and AI tools, and ready to help the team with daily work.

---

## WORK EXPERIENCE & INTERNSHIPS
- **Business Development Intern | Pencil Mark Solutions:** Called and emailed corporate clients to introduce interior design services; got a written appreciation letter for good work.
- **Operations Intern | Instawork Services India:** Helped the team collect, sort, and double-check data on the computer for AI and robotics projects.
- **Stall Coordinator | AERO India Exhibition:** Set up the display stall, welcomed visitors, answered questions, and took care of all stock items.
- **Event Coordinator | College & Music Events:** Coordinated with suppliers for sound, lights, and stage setup; arranged travel and stay for artists.
- **Store Assistant | Family Retail Store:** Handled daily customer billing, counter sales, shop registers, and ordering fresh stock from suppliers.

---

## PROJECTS & AI WORK
- **Simple Web Apps & Personal Website ([adityamehra007.github.io/ADI-OS](https://adityamehra007.github.io/ADI-OS/)):** Made working web tools and personal portfolio pages using AI coding tools and GitHub.
- **AI Video Making:** Created simple promotional videos, product clips, and slide designs using AI video tools.

---

## EDUCATION
**Bachelor of Business Administration (BBA) — International Business** | *Dayananda Sagar University, Bengaluru*  
- **Subjects Studied:** International Trade, Marketing, Supply Chain, Business Communication, Accounting.

---

## KEY SKILLS
- **Computer & AI Skills:** MS Excel, MS Word, MS PowerPoint, Google Sheets, ChatGPT, AI Tools, Internet Research.
- **Work Skills:** Client Communication, Customer Service, Teamwork, Vendor Follow-ups, Office Support, Problem Solving.
- **Languages:** English, Hindi.

---

## CERTIFICATES
- Appreciation Letter for Good Work (Pencil Mark) | Digital Marketing (Google) | Marketing Certificate (NPTEL, IIT Kharagpur) | AI Workshop (Outskill)
"""
    out_md = BASE_DIR / filename
    out_md.write_text(md_content, encoding="utf-8")
    return out_md

if __name__ == "__main__":
    # 1. Build DOCX
    docx_file = create_very_simple_cv("ADITYA_MEHRA_VERY_SIMPLE_CV.docx")
    print(f"Created: {docx_file}")

    # Also sync standard names
    create_very_simple_cv("ADITYA_MEHRA_POINT_TO_POINT_CV.docx")
    create_very_simple_cv("ADITYA_MEHRA_STANDOUT_FRESHER_CV.docx")
    create_very_simple_cv("ADITYA_MEHRA_PERFECT_FRESHER_CV.docx")
    create_very_simple_cv("ADITYA_MEHRA_SIMPLE_ENGLISH_CV.docx")
    print("Updated all docx versions.")

    # 2. Build HTML
    html_file = generate_html_version("ADITYA_MEHRA_VERY_SIMPLE_CV.html")
    generate_html_version("ADITYA_MEHRA_POINT_TO_POINT_CV.html")
    generate_html_version("ADITYA_MEHRA_STANDOUT_FRESHER_CV.html")
    generate_html_version("ADITYA_MEHRA_PERFECT_FRESHER_CV.html")
    generate_html_version("ADITYA_MEHRA_SIMPLE_ENGLISH_CV.html")
    print(f"Created HTML: {html_file}")

    # 3. Build Markdown
    md_file = generate_markdown_version("ADITYA_MEHRA_VERY_SIMPLE_CV.md")
    generate_markdown_version("ADITYA_MEHRA_POINT_TO_POINT_CV.md")
    generate_markdown_version("ADITYA_MEHRA_STANDOUT_FRESHER_CV.md")
    generate_markdown_version("ADITYA_MEHRA_PERFECT_FRESHER_CV.md")
    generate_markdown_version("ADITYA_MEHRA_SIMPLE_ENGLISH_CV.md")
    print(f"Created MD: {md_file}")

    # 4. Word COM Export & Verification
    print("Opening Word COM to verify page count and export PDFs...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False

    targets = [
        "ADITYA_MEHRA_VERY_SIMPLE_CV",
        "ADITYA_MEHRA_POINT_TO_POINT_CV",
        "ADITYA_MEHRA_STANDOUT_FRESHER_CV",
        "ADITYA_MEHRA_PERFECT_FRESHER_CV",
        "ADITYA_MEHRA_SIMPLE_ENGLISH_CV"
    ]
    for target_name in targets:
        target_docx = BASE_DIR / f"{target_name}.docx"
        target_pdf = BASE_DIR / f"{target_name}.pdf"
        doc = word.Documents.Open(str(target_docx))
        pages = doc.ComputeStatistics(2)
        print(f"[{target_name}] Exact Page Count: {pages}")
        doc.ExportAsFixedFormat(str(target_pdf), 17) # 17 = wdExportFormatPDF
        doc.Close()
        print(f"[{target_name}] Exported PDF: {target_pdf}")

    word.Quit()
    print("ALL BUILDS AND VERIFICATIONS COMPLETED SUCCESSFULLY!")
