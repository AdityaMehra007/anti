# -*- coding: utf-8 -*-
"""
ANTIGRAVITY BATCH COMPILER: Tailored CVs and Cover Letters to DOCX
Compiles all 61 company-tailored CVs and Cover Letters into executive-grade .docx documents.
"""

import os
import sys
import re
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')
if hasattr(sys.stderr, 'reconfigure'):
    sys.stderr.reconfigure(encoding='utf-8')

BASE_DIR = Path('e:/anti')
CV_DIR = BASE_DIR / 'Company_Tailored_CVs'
OUTPUT_DIR = CV_DIR / 'docx'
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

# Styling constants
COLOR_NAVY = RGBColor(15, 23, 42)      # #0F172A
COLOR_SLATE = RGBColor(51, 65, 85)     # #334155
COLOR_MUTED = RGBColor(100, 116, 139)  # #64748B
COLOR_TEXT = RGBColor(30, 41, 59)      # #1E293B

def clean_text(text: str) -> str:
    """Cleans up encoding artifacts like rogue question marks."""
    text = text.replace(' ? ', ' – ')
    text = text.replace(' ?', ' –')
    text = text.replace('??? ', '')
    text = text.replace('?? ', '')
    return text.strip()

def compile_cv_to_docx(md_path: Path, out_path: Path):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.65)
        s.bottom_margin = Inches(0.65)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    with open(md_path, 'r', encoding='utf-8', errors='replace') as f:
        lines = f.readlines()

    for line in lines:
        raw = line.rstrip()
        clean = clean_text(raw)
        if not clean:
            continue

        if clean.startswith('# '):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(clean[2:].strip().upper())
            run.font.name = 'Calibri'
            run.font.size = Pt(20)
            run.font.bold = True
            run.font.color.rgb = COLOR_NAVY

        elif clean.startswith('**Bengaluru, Karnataka') or ('+91 7003456624' in clean and 'ashishiash007@gmail.com' in clean):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            clean_contact = clean.replace('**', '')
            run = p.add_run(clean_contact)
            run.font.name = 'Calibri'
            run.font.size = Pt(9.5)
            run.font.color.rgb = COLOR_MUTED

        elif clean.startswith('## TARGET POSITION:'):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(clean.replace('## ', ''))
            run.font.name = 'Calibri'
            run.font.size = Pt(10)
            run.font.bold = True
            run.font.color.rgb = RGBColor(14, 116, 144)

        elif clean.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(clean[3:].strip().upper())
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = COLOR_NAVY

        elif clean.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(1)
            heading_text = clean[4:].strip()
            run = p.add_run(heading_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(10.5)
            run.font.bold = True
            run.font.color.rgb = COLOR_SLATE

        elif clean.startswith('*') and clean.endswith('*') and len(clean) > 2:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            meta_text = clean.strip('*').strip()
            run = p.add_run(meta_text)
            run.font.name = 'Calibri'
            run.font.size = Pt(9)
            run.font.italic = True
            run.font.color.rgb = COLOR_MUTED

        elif clean.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.12
            bullet_text = clean[2:].strip()
            
            parts = re.split(r'(\*\*.*?\*\*)', bullet_text)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9.5)
                    r.font.bold = True
                    r.font.color.rgb = COLOR_TEXT
                else:
                    r = p.add_run(part)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = COLOR_TEXT

        elif clean == '---':
            continue

        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            parts = re.split(r'(\*\*.*?\*\*)', clean)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9.5)
                    r.font.bold = True
                    r.font.color.rgb = COLOR_TEXT
                else:
                    r = p.add_run(part)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = COLOR_TEXT

    doc.save(str(out_path))

def compile_cover_letter_to_docx(md_path: Path, out_path: Path):
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.85)
        s.right_margin = Inches(0.85)

    with open(md_path, 'r', encoding='utf-8', errors='replace') as f:
        lines = f.readlines()

    for line in lines:
        raw = line.rstrip()
        clean = clean_text(raw)
        if not clean or clean == '---':
            continue

        if clean.startswith('# COVER LETTER:'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            run = p.add_run(clean.replace('# COVER LETTER:', 'APPLICATION COVER LETTER:').strip().upper())
            run.font.name = 'Calibri'
            run.font.size = Pt(13)
            run.font.bold = True
            run.font.color.rgb = COLOR_NAVY

        elif clean.startswith('**Candidate:**') or clean.startswith('**Contact:**') or clean.startswith('**Target Organization:**') or clean.startswith('**Date:**'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            parts = re.split(r'(\*\*.*?\*\*)', clean)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9.5)
                    r.font.bold = True
                    r.font.color.rgb = COLOR_SLATE
                else:
                    r = p.add_run(part)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(9.5)
                    r.font.color.rgb = COLOR_MUTED

        elif clean.startswith('Dear Hiring'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(clean)
            r.font.name = 'Calibri'
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = COLOR_NAVY

        elif re.match(r'^\d+\.\s+\*\*', clean):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            content = re.sub(r'^\d+\.\s+', '', clean)
            parts = re.split(r'(\*\*.*?\*\*)', content)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Calibri'
                    r.font.size = Pt(10)
                    r.font.bold = True
                    r.font.color.rgb = COLOR_TEXT
                else:
                    r = p.add_run(part)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(10)
                    r.font.color.rgb = COLOR_TEXT

        elif clean == 'Sincerely,':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(clean)
            r.font.name = 'Calibri'
            r.font.size = Pt(10.5)
            r.font.color.rgb = COLOR_TEXT

        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.18
            parts = re.split(r'(\*\*.*?\*\*)', clean)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    r = p.add_run(part[2:-2])
                    r.font.name = 'Calibri'
                    r.font.size = Pt(10)
                    r.font.bold = True
                    r.font.color.rgb = COLOR_TEXT
                else:
                    r = p.add_run(part)
                    r.font.name = 'Calibri'
                    r.font.size = Pt(10)
                    r.font.color.rgb = COLOR_TEXT

    doc.save(str(out_path))

def main():
    print("=" * 80)
    print("    COMPILING ALL 61 COMPANY-TAILORED CVS & COVER LETTERS TO DOCX    ")
    print("=" * 80)

    cv_files = sorted(list(CV_DIR.glob('CV_BLR-JOB-*.md')))
    cl_files = sorted(list(CV_DIR.glob('Cover_Letter_BLR-JOB-*.md')))

    print(f"Found {len(cv_files)} Tailored CVs and {len(cl_files)} Cover Letters.")

    compiled_cvs = 0
    for cv_path in cv_files:
        out_name = cv_path.stem + '.docx'
        out_path = OUTPUT_DIR / out_name
        try:
            compile_cv_to_docx(cv_path, out_path)
            compiled_cvs += 1
        except Exception as e:
            print(f"  [ERROR] Failed to compile {cv_path.name}: {e}")

    print(f"  -> Successfully compiled {compiled_cvs}/{len(cv_files)} CVs to DOCX.")

    compiled_cls = 0
    for cl_path in cl_files:
        out_name = cl_path.stem + '.docx'
        out_path = OUTPUT_DIR / out_name
        try:
            compile_cover_letter_to_docx(cl_path, out_path)
            compiled_cls += 1
        except Exception as e:
            print(f"  [ERROR] Failed to compile {cl_path.name}: {e}")

    print(f"  -> Successfully compiled {compiled_cls}/{len(cl_files)} Cover Letters to DOCX.")
    print(f"\nAll documents saved to: {OUTPUT_DIR}")
    print("=" * 80)

if __name__ == '__main__':
    main()
