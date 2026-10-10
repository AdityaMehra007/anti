import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client
import pypdf

BASE_DIR = Path(r"e:\anti")

# Palette: World-Class Executive Typography (Harvard / McKinsey Standard)
C_NAVY = RGBColor(15, 23, 42)       # #0f172a Deep Slate Navy
C_SLATE = RGBColor(51, 65, 85)      # #334155 Professional Slate
C_BODY = RGBColor(30, 41, 59)       # #1e293b High contrast body charcoal
C_MUTED = RGBColor(100, 116, 139)   # #64748b Metadata slate
C_LINK = "1E3A8A"                   # #1e3a8a Deep Royal Blue for hyperlinks

def add_clean_divider(paragraph):
    """Adds an elegant 0.5pt hairline horizontal rule below the section header."""
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')       # 0.5 pt clean hairline
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CBD5E1')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_hyperlink_run(paragraph, url, display_text, font_size_pt=8.8, color_hex="1E3A8A", bold=False, underline=True):
    """Injects an active OpenXML hyperlink run that Word converts into an active PDF URI annotation."""
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'<w:color w:val="{color_hex}"/>'
        f'{"<w:b/>" if bold else ""}'
        f'{"<w:u w:val=\"single\"/>" if underline else ""}'
        f'<w:sz w:val="{int(font_size_pt * 2)}"/>'
        f'</w:rPr>'
        f'<w:t>{display_text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    paragraph._p.append(hyperlink)

def create_world_class_cv(filename="ADITYA_MEHRA_BEST_OF_ALL_CV.docx"):
    """
    Builds the ultimate 1-page CV capturing everything about Aditya Mehra.
    - All 6 real-world roles (Pencil Mark, Instawork, AERO India, Events, Retail, Community NGO)
    - All 3 tech & AI proof-of-work projects (Digital Universe, Automation Tools, Multimedia)
    - Full education details, coursework, skills, and certifications
    - Calibrated to fill 94.2% of the A4 page height (bottom margin ~0.68 in) with strict 1-page invariant.
    """
    doc = Document()

    # Exact page margins
    for section in doc.sections:
        section.page_width = Inches(8.27)    # 210 mm
        section.page_height = Inches(11.69)  # 297 mm
        section.top_margin = Inches(0.38)
        section.bottom_margin = Inches(0.38)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)

    # 1. HEADER - NAME
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1.5)
    r_name = p_name.add_run("ADITYA MEHRA")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(21.5)
    r_name.font.bold = True
    r_name.font.color.rgb = C_NAVY

    # SUBTITLE: Grounded Executive Fresher Branding
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2.5)
    r_sub = p_sub.add_run("BBA Graduate  •  Operations, Client Support & AI Productivity  •  Entry-Level")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_SLATE

    # CONTACT ROW 1: Location • Phone • Email
    p_con1 = doc.add_paragraph()
    p_con1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con1.paragraph_format.space_before = Pt(0)
    p_con1.paragraph_format.space_after = Pt(1.5)
    
    r_loc = p_con1.add_run("Bengaluru, Karnataka, India  •  ")
    r_loc.font.name = "Calibri"; r_loc.font.size = Pt(8.8); r_loc.font.color.rgb = C_MUTED
    
    r_ph = p_con1.add_run("Phone: ")
    r_ph.font.name = "Calibri"; r_ph.font.size = Pt(8.8); r_ph.font.bold = True; r_ph.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "tel:+917003456624", "+91 7003456624", font_size_pt=8.8, color_hex=C_LINK, underline=False)
    
    r_sep1 = p_con1.add_run("  •  Email: ")
    r_sep1.font.name = "Calibri"; r_sep1.font.size = Pt(8.8); r_sep1.font.bold = True; r_sep1.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "mailto:adityamehra799@gmail.com", "adityamehra799@gmail.com", font_size_pt=8.8, color_hex=C_LINK, underline=False)

    # CONTACT ROW 2: Portfolio • LinkedIn • GitHub
    p_con2 = doc.add_paragraph()
    p_con2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con2.paragraph_format.space_before = Pt(0)
    p_con2.paragraph_format.space_after = Pt(2.8)

    r_web_lbl = p_con2.add_run("Portfolio: ")
    r_web_lbl.font.name = "Calibri"; r_web_lbl.font.size = Pt(8.8); r_web_lbl.font.bold = True; r_web_lbl.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://adi-digital-universe.ai.studio/", "adi-digital-universe.ai.studio", font_size_pt=8.8, color_hex=C_LINK, underline=True)

    r_sep2 = p_con2.add_run("  •  LinkedIn: ")
    r_sep2.font.name = "Calibri"; r_sep2.font.size = Pt(8.8); r_sep2.font.bold = True; r_sep2.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://www.linkedin.com/in/aditya-mehra-b8644b326", "linkedin.com/in/aditya-mehra-b8644b326", font_size_pt=8.8, color_hex=C_LINK, underline=True)

    r_sep3 = p_con2.add_run("  •  GitHub: ")
    r_sep3.font.name = "Calibri"; r_sep3.font.size = Pt(8.8); r_sep3.font.bold = True; r_sep3.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://github.com/AdityaMehra007", "github.com/AdityaMehra007", font_size_pt=8.8, color_hex=C_LINK, underline=True)

    # Master calibrated typography
    body_pt = 9.0
    head_pt = 10.2
    sec_before = 8.5
    bullet_after = 2.8
    line_sp = 1.14

    def add_sec(title, before=sec_before):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title.upper())
        r.font.name = "Calibri"
        r.font.size = Pt(head_pt)
        r.font.bold = True
        r.font.color.rgb = C_NAVY
        add_clean_divider(p)

    def add_bullet(p, title_bold, text_body, space_after=bullet_after):
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_before = Pt(0.4)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_sp
        
        rb = p.add_run("• ")
        rb.font.name = "Calibri"; rb.font.size = Pt(body_pt); rb.bold = True; rb.font.color.rgb = C_SLATE

        if title_bold:
            rt = p.add_run(title_bold)
            rt.font.name = "Calibri"; rt.font.size = Pt(body_pt); rt.bold = True; rt.font.color.rgb = C_NAVY

        rc = p.add_run(text_body)
        rc.font.name = "Calibri"; rc.font.size = Pt(body_pt); rc.font.color.rgb = C_BODY

    # 2. CAREER OBJECTIVE
    add_sec("Career Objective", before=1.8)
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(0.8)
    p_obj.paragraph_format.space_after = Pt(1.8)
    p_obj.paragraph_format.line_spacing = line_sp
    r_obj = p_obj.add_run(
        "Dedicated BBA graduate in International Business with practical experience across corporate outreach, operational workflows, and modern AI productivity tools. "
        "Fast learner with strong computer proficiency, eager to support cross-functional teams in daily operations, customer relations, vendor communication, and business administration."
    )
    r_obj.font.name = "Calibri"
    r_obj.font.size = Pt(body_pt)
    r_obj.font.color.rgb = C_BODY

    # 3. WORK EXPERIENCE & INTERNSHIPS (All 6 verified roles)
    add_sec("Work Experience & Internships")
    
    # 1. Pencil Mark Solutions
    p = doc.add_paragraph()
    add_bullet(p, "Business Development Intern | Pencil Mark Solutions: ",
               "Conducted targeted outreach to prospective corporate clients via phone and email to introduce commercial interior design services.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Assisted senior management with preparing service proposals, cost estimates, client meeting presentations, and follow-up schedules.")

    # 2. Instawork Services India
    p = doc.add_paragraph()
    add_bullet(p, "Operations Intern | Instawork Services India: ",
               "Organized, structured, and verified computer datasets to support engineering teams with robotics and AI operations.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Maintained strict data accuracy standards by cross-checking spreadsheets, documenting errors, and reporting workflow discrepancies promptly.")

    # 3. AERO India Exhibition
    p = doc.add_paragraph()
    add_bullet(p, "Stall Coordinator | AERO India Exhibition: ",
               "Managed exhibition booth setup, welcomed trade delegates and international visitors, answered inquiries, and maintained inventory control.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Monitored promotional merchandise and product catalogues under high footfall to ensure continuous availability throughout the event.")

    # 4. College & Music Events
    p = doc.add_paragraph()
    add_bullet(p, "Event Coordinator | College & Music Events: ",
               "Coordinated with technical vendors for stage, sound, and lighting setups while managing hospitality and travel logistics for performing artists.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Supervised on-ground event volunteers and maintained strict schedules to ensure seamless show operations and crowd management.")

    # 5. Store Assistant | Family Retail Store
    p = doc.add_paragraph()
    add_bullet(p, "Store Assistant | Family Retail Store: ",
               "Managed daily counter sales, customer billing, cash registers, and supplier restocking to maintain seamless store operations.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Communicated regularly with wholesale distributors to place repeat merchandise orders and verify incoming delivery receipts.")

    # 6. Community Operations Volunteer | Social Initiative
    p = doc.add_paragraph()
    add_bullet(p, "Community Operations Volunteer | Social Initiative: ",
               "Coordinated volunteer shift schedules, attendee check-ins, and on-ground logistics for community educational drives in Bengaluru.")

    # 4. PROJECTS & AI PROOF OF WORK (All 3 projects)
    add_sec("Projects & AI Proof of Work")
    p = doc.add_paragraph()
    add_bullet(p, "Web Apps & Digital Portfolio (adi-digital-universe.ai.studio): ",
               "Built and deployed interactive web applications and a personal portfolio website using modern AI coding workflows and cloud hosting.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Maintained active public GitHub repositories showcasing code implementations, interface prototypes, and practical automation scripts.")

    p = doc.add_paragraph()
    add_bullet(p, "Operations Automation & Workflow Tools: ",
               "Prototyped operations automation scripts, interactive status dashboards, and digital checklists to eliminate repetitive manual tracking.")

    p = doc.add_paragraph()
    add_bullet(p, "AI Multimedia & Content Production: ",
               "Produced promotional video assets, visual marketing collateral, and digital presentation decks using modern generative AI tools.")
    p = doc.add_paragraph()
    add_bullet(p, "", "Accelerated creative content delivery by combining AI image generation workflows, video storyboard drafting, and prompt refinement.")

    # 5. EDUCATION
    add_sec("Education")
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(0.8)
    p_edu.paragraph_format.space_after = Pt(1.5)
    p_edu.paragraph_format.line_spacing = line_sp
    p_edu.paragraph_format.keep_with_next = True

    r_deg = p_edu.add_run("Bachelor of Business Administration (BBA) — International Business (2023 – 2026)  •  ")
    r_deg.font.name = "Calibri"; r_deg.font.size = Pt(body_pt + 0.2); r_deg.bold = True; r_deg.font.color.rgb = C_NAVY

    r_uni = p_edu.add_run("Dayananda Sagar University, Bengaluru\n")
    r_uni.font.name = "Calibri"; r_uni.font.size = Pt(body_pt); r_uni.font.italic = True; r_uni.font.color.rgb = C_SLATE

    r_sub = p_edu.add_run("Relevant Coursework: International Trade, Marketing Management, Supply Chain Logistics, Business Communication, Accounting.")
    r_sub.font.name = "Calibri"; r_sub.font.size = Pt(body_pt - 0.4); r_sub.font.color.rgb = C_BODY

    # 6. KEY SKILLS & COMPETENCIES
    add_sec("Key Skills & Competencies")
    skills = [
        ("Computer & AI Tools: ", "ChatGPT, Generative AI Platforms, MS Excel (Formulas & Data Verification), MS Word, MS PowerPoint, Google Workspace, GitHub, Web Research."),
        ("Workplace Competencies: ", "Client Communication, Customer Relations, Operations Support, Vendor Coordination, Team Collaboration, Problem Solving, Time Management."),
        ("Languages: ", "English, Hindi.")
    ]
    for s_title, s_desc in skills:
        p_s = doc.add_paragraph()
        add_bullet(p_s, s_title, s_desc, space_after=2.0)

    # 7. CERTIFICATES & HONORS
    add_sec("Certificates & Honors")
    certs = [
        ("Internship Certificate: ", "Pencil Mark Solutions (Client Outreach & Business Operations)."),
        ("Fundamentals of Digital Marketing: ", "Google (Online Marketing, SEO & Web Analytics)."),
        ("Marketing Management Certification: ", "NPTEL, IIT Kharagpur (Marketing Principles & Consumer Behavior)."),
        ("AI Productivity Workshop: ", "Outskill (Generative AI Workflows & Modern Office Automation).")
    ]
    for c_title, c_desc in certs:
        p_c = doc.add_paragraph()
        add_bullet(p_c, c_title, c_desc, space_after=1.6)

    out_path = BASE_DIR / filename
    doc.save(out_path)
    return out_path

def generate_html_version(filename="ADITYA_MEHRA_BEST_OF_ALL_CV.html"):
    html_content = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Aditya Mehra — The Best of All CV (Master Complete Edition)</title>
<style>
  :root {
    --primary: #0f172a;
    --accent: #1e3a8a;
    --slate: #334155;
    --body: #1e293b;
    --border: #cbd5e1;
  }
  * { box-sizing: border-box; margin: 0; padding: 0; }
  body {
    font-family: Calibri, -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
    color: var(--body);
    background-color: #f1f5f9;
    line-height: 1.34;
    padding: 24px;
    display: flex;
    justify-content: center;
  }
  .cv-page {
    background: #ffffff;
    width: 210mm;
    min-height: 297mm;
    padding: 10mm 12mm;
    box-shadow: 0 4px 24px rgba(0, 0, 0, 0.08);
    border-radius: 4px;
  }
  header { text-align: center; margin-bottom: 10px; }
  h1 {
    font-size: 21.5pt;
    font-weight: 800;
    letter-spacing: 0.5px;
    color: var(--primary);
    margin-bottom: 2px;
  }
  .subtitle {
    font-size: 9.5pt;
    font-weight: 700;
    color: var(--slate);
    margin-bottom: 4px;
  }
  .contact-bar {
    font-size: 8.8pt;
    color: #64748b;
    display: flex;
    justify-content: center;
    flex-wrap: wrap;
    gap: 8px;
    line-height: 1.35;
  }
  .contact-bar span { display: inline-flex; align-items: center; }
  .contact-bar strong { color: var(--primary); }
  .contact-bar a {
    color: var(--accent);
    text-decoration: underline;
    font-weight: 600;
  }
  section { margin-bottom: 9px; }
  h2 {
    font-size: 10.2pt;
    font-weight: 700;
    color: var(--primary);
    text-transform: uppercase;
    letter-spacing: 0.5px;
    border-bottom: 1px solid var(--border);
    padding-bottom: 2px;
    margin-bottom: 5px;
  }
  p.objective {
    font-size: 9.0pt;
    color: var(--body);
    line-height: 1.36;
    text-align: justify;
  }
  .point-item {
    font-size: 9.0pt;
    line-height: 1.32;
    color: var(--body);
    padding-left: 14px;
    position: relative;
    margin-bottom: 3.0px;
  }
  .point-item::before {
    content: "•";
    position: absolute;
    left: 2px;
    color: var(--slate);
    font-weight: bold;
  }
  .point-item strong {
    color: var(--primary);
  }
  .edu-row {
    font-size: 9.0pt;
    margin-bottom: 2px;
  }
  .edu-title { font-weight: 700; color: var(--primary); }
  .edu-inst { font-style: italic; color: var(--slate); }
  .edu-sub { font-size: 8.6pt; color: var(--body); }
  @media print {
    body { background: #fff; padding: 0; }
    .cv-page { box-shadow: none; padding: 10mm 12mm; width: 100%; min-height: auto; }
  }
</style>
</head>
<body>
<div class="cv-page">
  <header>
    <h1>ADITYA MEHRA</h1>
    <div class="subtitle">BBA Graduate &nbsp;•&nbsp; Operations, Client Support &amp; AI Productivity &nbsp;•&nbsp; Entry-Level</div>
    <div class="contact-bar">
      <span>Bengaluru, Karnataka, India</span>
      <span>•</span>
      <span><strong>Phone:</strong> <a href="tel:+917003456624">+91 7003456624</a></span>
      <span>•</span>
      <span><strong>Email:</strong> <a href="mailto:adityamehra799@gmail.com">adityamehra799@gmail.com</a></span>
    </div>
    <div class="contact-bar" style="margin-top: 2.5px;">
      <span><strong>Portfolio:</strong> <a href="https://adi-digital-universe.ai.studio/" target="_blank">adi-digital-universe.ai.studio</a></span>
      <span>•</span>
      <span><strong>LinkedIn:</strong> <a href="https://www.linkedin.com/in/aditya-mehra-b8644b326" target="_blank">linkedin.com/in/aditya-mehra-b8644b326</a></span>
      <span>•</span>
      <span><strong>GitHub:</strong> <a href="https://github.com/AdityaMehra007" target="_blank">github.com/AdityaMehra007</a></span>
    </div>
  </header>

  <section>
    <h2>Career Objective</h2>
    <p class="objective">
      Dedicated BBA graduate in International Business with practical experience across corporate outreach, operational workflows, and modern AI productivity tools. 
      Fast learner with strong computer proficiency, eager to support cross-functional teams in daily operations, customer relations, vendor communication, and business administration.
    </p>
  </section>

  <section>
    <h2>Work Experience &amp; Internships</h2>
    <div class="point-item">
      <strong>Business Development Intern | Pencil Mark Solutions:</strong> 
      Conducted targeted outreach to prospective corporate clients via phone and email to introduce commercial interior design services.
    </div>
    <div class="point-item">
      Assisted senior management with preparing service proposals, cost estimates, client meeting presentations, and follow-up schedules.
    </div>

    <div class="point-item">
      <strong>Operations Intern | Instawork Services India:</strong> 
      Organized, structured, and verified computer datasets to support engineering teams with robotics and AI operations.
    </div>
    <div class="point-item">
      Maintained strict data accuracy standards by cross-checking spreadsheets, documenting errors, and reporting workflow discrepancies promptly.
    </div>

    <div class="point-item">
      <strong>Stall Coordinator | AERO India Exhibition:</strong> 
      Managed exhibition booth setup, welcomed trade delegates and international visitors, answered inquiries, and maintained inventory control.
    </div>
    <div class="point-item">
      Monitored promotional merchandise and product catalogues under high footfall to ensure continuous availability throughout the event.
    </div>

    <div class="point-item">
      <strong>Event Coordinator | College &amp; Music Events:</strong> 
      Coordinated with technical vendors for stage, sound, and lighting setups while managing hospitality and travel logistics for performing artists.
    </div>
    <div class="point-item">
      Supervised on-ground event volunteers and maintained strict schedules to ensure seamless show operations and crowd management.
    </div>

    <div class="point-item">
      <strong>Store Assistant | Family Retail Store:</strong> 
      Managed daily counter sales, customer billing, cash registers, and supplier restocking to maintain seamless store operations.
    </div>
    <div class="point-item">
      Communicated regularly with wholesale distributors to place repeat merchandise orders and verify incoming delivery receipts.
    </div>

    <div class="point-item">
      <strong>Community Operations Volunteer | Social Initiative:</strong> 
      Coordinated volunteer shift schedules, attendee check-ins, and on-ground logistics for community educational drives in Bengaluru.
    </div>
  </section>

  <section>
    <h2>Projects &amp; AI Proof of Work</h2>
    <div class="point-item">
      <strong>Web Apps &amp; Digital Portfolio (<a href="https://adi-digital-universe.ai.studio/" target="_blank" style="color:var(--accent);">adi-digital-universe.ai.studio</a>):</strong> 
      Built and deployed interactive web applications and a personal portfolio website using modern AI coding workflows and cloud hosting.
    </div>
    <div class="point-item">
      Maintained active public GitHub repositories showcasing code implementations, interface prototypes, and practical automation scripts.
    </div>

    <div class="point-item">
      <strong>Operations Automation &amp; Workflow Tools:</strong> 
      Prototyped operations automation scripts, interactive status dashboards, and digital checklists to eliminate repetitive manual tracking.
    </div>

    <div class="point-item">
      <strong>AI Multimedia &amp; Content Production:</strong> 
      Produced promotional video assets, visual marketing collateral, and digital presentation decks using modern generative AI tools.
    </div>
    <div class="point-item">
      Accelerated creative content delivery by combining AI image generation workflows, video storyboard drafting, and prompt refinement.
    </div>
  </section>

  <section>
    <h2>Education</h2>
    <div class="edu-row">
      <span class="edu-title">Bachelor of Business Administration (BBA) — International Business (2023 – 2026)</span> &nbsp;•&nbsp;
      <span class="edu-inst">Dayananda Sagar University, Bengaluru</span>
    </div>
    <div class="edu-sub">Relevant Coursework: International Trade, Marketing Management, Supply Chain Logistics, Business Communication, Accounting.</div>
  </section>

  <section>
    <h2>Key Skills &amp; Competencies</h2>
    <div class="point-item">
      <strong>Computer &amp; AI Tools:</strong> ChatGPT, Generative AI Platforms, MS Excel (Formulas &amp; Data Verification), MS Word, MS PowerPoint, Google Workspace, GitHub, Web Research.
    </div>
    <div class="point-item">
      <strong>Workplace Competencies:</strong> Client Communication, Customer Relations, Operations Support, Vendor Coordination, Team Collaboration, Problem Solving, Time Management.
    </div>
    <div class="point-item">
      <strong>Languages:</strong> English, Hindi.
    </div>
  </section>

  <section>
    <h2>Certificates &amp; Honors</h2>
    <div class="point-item">
      <strong>Internship Certificate:</strong> Pencil Mark Solutions (Client Outreach &amp; Business Operations).
    </div>
    <div class="point-item">
      <strong>Fundamentals of Digital Marketing:</strong> Google (Online Marketing, SEO &amp; Web Analytics).
    </div>
    <div class="point-item">
      <strong>Marketing Management Certification:</strong> NPTEL, IIT Kharagpur (Marketing Principles &amp; Consumer Behavior).
    </div>
    <div class="point-item">
      <strong>AI Productivity Workshop:</strong> Outskill (Generative AI Workflows &amp; Modern Office Automation).
    </div>
  </section>
</div>
</body>
</html>
"""
    out_html = BASE_DIR / filename
    out_html.write_text(html_content, encoding="utf-8")
    return out_html

def generate_markdown_version(filename="ADITYA_MEHRA_BEST_OF_ALL_CV.md"):
    md_content = """# ADITYA MEHRA
**BBA Graduate | Operations, Client Support & AI Productivity | Entry-Level**  
Bengaluru, Karnataka, India • Phone: [+91 7003456624](tel:+917003456624) • Email: [adityamehra799@gmail.com](mailto:adityamehra799@gmail.com)  
**Portfolio:** [adi-digital-universe.ai.studio](https://adi-digital-universe.ai.studio/) • **LinkedIn:** [linkedin.com/in/aditya-mehra-b8644b326](https://www.linkedin.com/in/aditya-mehra-b8644b326) • **GitHub:** [github.com/AdityaMehra007](https://github.com/AdityaMehra007)

---

## CAREER OBJECTIVE
Dedicated BBA graduate in International Business with practical experience across corporate outreach, operational workflows, and modern AI productivity tools. Fast learner with strong computer proficiency, eager to support cross-functional teams in daily operations, customer relations, vendor communication, and business administration.

---

## WORK EXPERIENCE & INTERNSHIPS
- **Business Development Intern | Pencil Mark Solutions:** Conducted targeted outreach to prospective corporate clients via phone and email to introduce commercial interior design services.
- Assisted senior management with preparing service proposals, cost estimates, client meeting presentations, and follow-up schedules.
- **Operations Intern | Instawork Services India:** Organized, structured, and verified computer datasets to support engineering teams with robotics and AI operations.
- Maintained strict data accuracy standards by cross-checking spreadsheets, documenting errors, and reporting workflow discrepancies promptly.
- **Stall Coordinator | AERO India Exhibition:** Managed exhibition booth setup, welcomed trade delegates and international visitors, answered inquiries, and maintained inventory control.
- Monitored promotional merchandise and product catalogues under high footfall to ensure continuous availability throughout the event.
- **Event Coordinator | College & Music Events:** Coordinated with technical vendors for stage, sound, and lighting setups while managing hospitality and travel logistics for performing artists.
- Supervised on-ground event volunteers and maintained strict schedules to ensure seamless show operations and crowd management.
- **Store Assistant | Family Retail Store:** Managed daily counter sales, customer billing, cash registers, and supplier restocking to maintain seamless store operations.
- Communicated regularly with wholesale distributors to place repeat merchandise orders and verify incoming delivery receipts.
- **Community Operations Volunteer | Social Initiative:** Coordinated volunteer shift schedules, attendee check-ins, and on-ground logistics for community educational drives in Bengaluru.

---

## PROJECTS & AI PROOF OF WORK
- **Web Apps & Digital Portfolio ([adi-digital-universe.ai.studio](https://adi-digital-universe.ai.studio/)):** Built and deployed interactive web applications and a personal portfolio website using modern AI coding workflows and cloud hosting.
- Maintained active public GitHub repositories showcasing code implementations, interface prototypes, and practical automation scripts.
- **Operations Automation & Workflow Tools:** Prototyped operations automation scripts, interactive status dashboards, and digital checklists to eliminate repetitive manual tracking.
- **AI Multimedia & Content Production:** Produced promotional video assets, visual marketing collateral, and digital presentation decks using modern generative AI tools.
- Accelerated creative content delivery by combining AI image generation workflows, video storyboard drafting, and prompt refinement.

---

## EDUCATION
**Bachelor of Business Administration (BBA) — International Business (2023 – 2026)** • *Dayananda Sagar University, Bengaluru*  
- **Relevant Coursework:** International Trade, Marketing Management, Supply Chain Logistics, Business Communication, Accounting.

---

## KEY SKILLS & COMPETENCIES
- **Computer & AI Tools:** ChatGPT, Generative AI Platforms, MS Excel (Formulas & Data Verification), MS Word, MS PowerPoint, Google Workspace, GitHub, Web Research.
- **Workplace Competencies:** Client Communication, Customer Relations, Operations Support, Vendor Coordination, Team Collaboration, Problem Solving, Time Management.
- **Languages:** English, Hindi.

---

## CERTIFICATES & HONORS
- **Internship Certificate:** Pencil Mark Solutions (Client Outreach & Business Operations).
- **Fundamentals of Digital Marketing:** Google (Online Marketing, SEO & Web Analytics).
- **Marketing Management Certification:** NPTEL, IIT Kharagpur (Marketing Principles & Consumer Behavior).
- **AI Productivity Workshop:** Outskill (Generative AI Workflows & Modern Office Automation).
"""
    out_md = BASE_DIR / filename
    out_md.write_text(md_content, encoding="utf-8")
    return out_md

if __name__ == "__main__":
    canonical_names = [
        "ADITYA_MEHRA_BEST_OF_ALL_CV",
        "ADITYA_MEHRA_FLAWLESS_CV",
        "ADITYA_MEHRA_PERFECT_FRESHER_CV",
        "ADITYA_MEHRA_STANDOUT_FRESHER_CV",
        "ADITYA_MEHRA_POINT_TO_POINT_CV",
        "ADITYA_MEHRA_VERY_SIMPLE_CV",
        "ADITYA_MEHRA_SIMPLE_ENGLISH_CV"
    ]
    for c_name in canonical_names:
        create_world_class_cv(f"{c_name}.docx")
        generate_html_version(f"{c_name}.html")
        generate_markdown_version(f"{c_name}.md")
    print("Synchronized all document formats.")

    # Word COM Export & Strict Verification
    print("Starting Word COM export & verification...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False

    for c_name in canonical_names:
        target_docx = BASE_DIR / f"{c_name}.docx"
        target_pdf = BASE_DIR / f"{c_name}.pdf"
        doc = word.Documents.Open(str(target_docx))
        pages = doc.ComputeStatistics(2)
        print(f"[{c_name}] Word COM Statistics - Pages: {pages}")
        assert pages == 1, f"ERROR: {c_name} exceeded 1 page! Count: {pages}"
        doc.ExportAsFixedFormat(str(target_pdf), 17) # 17 = wdExportFormatPDF
        doc.Close()
        print(f"[{c_name}] Exported PDF: {target_pdf}")

    word.Quit()
    print("Word COM export completed.")

    # Independent PyPDF Verification of all PDFs
    print("Verifying PDFs with PyPDF...")
    for c_name in canonical_names:
        pdf_path = BASE_DIR / f"{c_name}.pdf"
        reader = pypdf.PdfReader(str(pdf_path))
        num_pages = len(reader.pages)
        print(f"[{c_name}.pdf] Total PyPDF Pages: {num_pages}")
        assert num_pages == 1, f"PyPDF Page Count Failed for {c_name}: {num_pages}"
        
        # Verify links
        page = reader.pages[0]
        links = []
        if '/Annots' in page:
            for annot in page['/Annots']:
                obj = annot.get_object()
                if '/A' in obj and '/URI' in obj['/A']:
                    links.append(obj['/A']['/URI'])
        print(f"[{c_name}.pdf] Verified Links ({len(links)}): {links}")
        assert len(links) == 5, f"Expected 5 active links, found {len(links)}"

        # Measure vertical coverage
        y_positions = []
        def visitor_body(text, cm, tm, fontDict, fontSize):
            if text.strip():
                y_positions.append(tm[5])
        page.extract_text(visitor_text=visitor_body)
        bottom_y = min(y_positions)
        print(f"[{c_name}.pdf] Lowest text coordinate y: {bottom_y:.1f} pt (bottom margin: {bottom_y/72.0:.2f} in / {bottom_y/841.68*100:.1f}%)")

    print("\n>>> ALL CHECKS PASSED: 100% COMPLETE MASTER 1-PAGE CV READY! <<<")
