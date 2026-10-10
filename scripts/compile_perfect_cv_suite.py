import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import win32com.client

BASE_DIR = Path(r"e:\anti")

# Brand colors
C_NAVY = RGBColor(15, 23, 42)       # #0F172A
C_ACCENT = RGBColor(30, 58, 138)    # #1E3A8A
C_TEAL = RGBColor(13, 148, 136)     # #0D9488
C_MUTED = RGBColor(71, 85, 105)     # #475569
C_BODY = RGBColor(30, 41, 59)       # #1E293B

def set_cell_margins(cell, top=50, bottom=50, left=50, right=50):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_clean_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '6')
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CBD5E1')
    pBdr.append(bottom)
    pPr.append(pBdr)

# ==========================================
# 1. BUILD COMPREHENSIVE 2-PAGE MASTER CV
# ==========================================
def build_2page_cv():
    doc = Document()
    
    # Configure A4 dimensions & precise margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.42)
        section.bottom_margin = Inches(0.42)
        section.left_margin = Inches(0.55)
        section.right_margin = Inches(0.55)

    # Header
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    run_name = p_name.add_run("ADITYA MEHRA")
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(20)
    run_name.font.bold = True
    run_name.font.color.rgb = C_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(3)
    run_sub = p_sub.add_run("BUSINESS OPERATIONS  |  COMMERCIAL GROWTH  |  AI-ENABLED AUTOMATION")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(9.5)
    run_sub.font.bold = True
    run_sub.font.color.rgb = C_TEAL

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(6)
    p_contact.paragraph_format.line_spacing = 1.05
    run_contact = p_contact.add_run(
        "Bengaluru, Karnataka, India  |  +91 7003456624  |  adityamehra799@gmail.com\n"
        "LinkedIn: linkedin.com/in/aditya-mehra-b8644b326  |  GitHub: github.com/AdityaMehra007"
    )
    run_contact.font.name = "Calibri"
    run_contact.font.size = Pt(8.5)
    run_contact.font.color.rgb = C_MUTED

    def add_section(title, space_before=6):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title.upper())
        run.font.name = "Calibri"
        run.font.size = Pt(10)
        run.font.bold = True
        run.font.color.rgb = C_NAVY
        add_clean_border(p)

    # Professional Summary
    add_section("Professional Summary", space_before=4)
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(1)
    p_sum.paragraph_format.space_after = Pt(4)
    p_sum.paragraph_format.line_spacing = 1.1
    p_sum.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_sum = p_sum.add_run(
        "Results-oriented Business Operations & Commercial Growth Specialist with extensive hands-on experience directing 300+ live commercial activations, "
        "governing high-stakes operations at AERO India 2025 (100k+ visitors, 0 shrinkage), and curating enterprise robotics AI training pipelines at Instawork "
        "with a verified 99%+ QA benchmark. Combines an international business foundation (BBA, Dayananda Sagar University '26) with proven execution across "
        "direct supplier SLA governance, 15–20% procurement cost reduction, multi-thread B2B deal negotiation, and AI-enabled process automation. "
        "Awarded an official written management commendation for commercial excellence. Seeking to deploy operational rigor within top-tier MNCs, GCCs, and high-growth enterprises."
    )
    r_sum.font.name = "Calibri"
    r_sum.font.size = Pt(8.8)
    r_sum.font.color.rgb = C_BODY

    # Core Competencies
    add_section("Core Competencies", space_before=5)
    competencies = [
        ("Operations & Vendor Governance: ", "End-to-End Project Delivery (300+ Events), Direct Supplier SLA Governance, Landed Cost Modeling, 15–20% Procurement Cost Optimization, High-Pressure Crisis Triage."),
        ("Commercial & B2B Growth: ", "Enterprise Client Engagement (Tata Communications, Puma India), Pipeline Management (15+ Concurrent Threads), Deal Qualification, Commercial Rate Structuring, Contract Negotiation."),
        ("AI & Process Automation: ", "Robotics Training Data Curation, Multi-Modal Annotation QA, Generative AI Prompt Engineering, Autonomous Workflow Design, Python & LLM Data Synthesis, Clay AI, Notion AI."),
        ("Global Trade & Supply Chain: ", "Incoterms 2020 (FOB, CIF, DDP), Letter of Credit (UCP 600) Principles, 3PL Freight Logistics, Inventory Buffer Management, Customs & EXIM Compliance Frameworks."),
        ("Analytics & Systems: ", "Advanced MS Excel (XLOOKUP, Power Query, Dynamic Pivot Tables), Salesforce & Zoho CRM, SQL Querying Fundamentals, Unit Economics, Margin Modeling."),
        ("Multilingual Fluency: ", "English (Fluent / Professional), Hindi (Native), Punjabi (Native), Bengali (Proficient), French (Conversational).")
    ]
    for prefix, details in competencies:
        p_c = doc.add_paragraph()
        p_c.paragraph_format.left_indent = Inches(0.15)
        p_c.paragraph_format.space_before = Pt(0.5)
        p_c.paragraph_format.space_after = Pt(1)
        p_c.paragraph_format.line_spacing = 1.08
        
        rb = p_c.add_run("• ")
        rb.font.name = "Calibri"
        rb.font.size = Pt(8.8)
        rb.font.bold = True
        rb.font.color.rgb = C_TEAL

        rp = p_c.add_run(prefix)
        rp.font.name = "Calibri"
        rp.font.size = Pt(8.8)
        rp.font.bold = True
        rp.font.color.rgb = C_NAVY

        rt = p_c.add_run(details)
        rt.font.name = "Calibri"
        rt.font.size = Pt(8.8)
        rt.font.color.rgb = C_BODY

    # Professional Experience
    add_section("Professional Experience", space_before=6)

    jobs = [
        {
            "role": "Independent Event Director & Commercial Operations Lead",
            "company": "Independent Commercial Operations | Bengaluru, India",
            "dates": "2019 – Present",
            "bullets": [
                ("Scaled Multi-Format Operations: ", "Directed end-to-end production and delivery for 300+ commercial events, corporate launches, and experiential campaigns for premier brands including Tata Communications, Puma India, VH1 Supersonic, and Apollo Marketing."),
                ("Procurement Cost Optimization: ", "Secured consistent 15% to 20% cost savings per project by establishing direct supplier rate cards, eliminating contractor markups, and enforcing milestone-driven SLAs across 20+ vendors."),
                ("Client Retention & Account Growth: ", "Maintained an exceptional 30%+ repeat-client rate through transparent commercial accounting, rigorous quality standards, and high-velocity communication under compressed deadlines."),
                ("Crisis Leadership & Rapid Triage: ", "Executed zero-downtime crisis management during unexpected operational disruptions (including weather-induced venue relocation and high-voltage technical transfers), safeguarding multi-lakh equipment and client SLA commitments.")
            ]
        },
        {
            "role": "AI Data Operations Intern",
            "company": "Instawork Services India Pvt. Ltd. | Bengaluru, India",
            "dates": "Dec 2025",
            "bullets": [
                ("Robotics & Vision Data Collection: ", "Supported machine learning pipelines training enterprise robotics and human activity recognition systems, managing high-density multi-modal data streams under strict capture protocols."),
                ("High-Fidelity QA Benchmark: ", "Maintained a verified 99%+ quality assurance benchmark, conducting rigorous schema validation, annotation audits, and edge-case classification to eliminate training distribution drift."),
                ("Cross-Functional Workflow Sync: ", "Harmonized on-ground data collection procedures with ML engineering priorities, authoring standardized operational capture SOPs.")
            ]
        },
        {
            "role": "Business Development Intern",
            "company": "Pencil Mark Interior Solutions LLP | Bengaluru, India",
            "dates": "Jul 2025 – Aug 2025",
            "bullets": [
                ("Pipeline Expansion & Revenue Closing: ", "Spearheaded B2B outbound acquisition across Bengaluru commercial interior and enterprise workspace projects, closing ₹1.5L+ in initial contract revenue."),
                ("High-Velocity Deal Management: ", "Managed 15+ concurrent corporate prospect threads, optimizing follow-up turnaround times and achieving an 18% lead-to-opportunity conversion rate."),
                ("Executive Commendation: ", "Received an official written management commendation from executive leadership for exemplary lead pipeline expansion, structured sales drive, and commercial negotiation results.")
            ]
        },
        {
            "role": "Exhibition Operations & Commercial Lead",
            "company": "Salt in My Coca — AERO India 2025 | Yelahanka Air Force Station, Bengaluru",
            "dates": "Feb 2025",
            "bullets": [
                ("Mission-Critical Expo Execution: ", "Commanded 7 full days of on-ground exhibition stall operations, product display logistics, and high-density visitor flows at Asia's premier defense exposition (100,000+ attendees)."),
                ("Strict Protocol & Security Governance: ", "Managed inventory movements and contractor logistics under stringent Ministry of Defence clearance schedules and flight display windows, maintaining zero inventory shrinkage or asset loss."),
                ("High-Value Stakeholder Interaction: ", "Engaged international military delegations, corporate CXOs, and institutional buyers, capturing qualified commercial leads and reinforcing brand positioning.")
            ]
        },
        {
            "role": "Event Operations Coordinator",
            "company": "TRILOGY: Indo-Jazz Instrumental Fusion | Bangalore Club, Bengaluru",
            "dates": "Jan 2026",
            "bullets": [
                ("End-to-End Production Coordination: ", "Coordinated logistics for an elite instrumental live concert, overseeing artist travel, hospitality, AV technical riders, and stage acoustic timelines under tight schedules."),
                ("Zero-Disruption Execution: ", "Resolved live sound balance and schedule dependencies in real time, delivering a seamless performance experience for high-profile patrons.")
            ]
        },
        {
            "role": "Co-Founder & Operations Manager",
            "company": "Mehra's Kitchen | Bengaluru, India",
            "dates": "2025",
            "bullets": [
                ("Unit Economics & P&L Oversight: ", "Designed operational and financial workflows for food-stall and cloud-kitchen operations, managing raw material procurement, vendor negotiation, and daily cash flow accounting."),
                ("Cost & Margin Control: ", "Analyzed landed ingredient costs and optimized menu pricing structures to secure operational profitability.")
            ]
        },
        {
            "role": "Operations & Sales Associate",
            "company": "Family Enterprise | Kolkata & Bengaluru, India",
            "dates": "2018 – 2020",
            "bullets": [
                ("Commercial Foundation: ", "Commenced commercial career at age 17 managing retail customer relations, inventory warehousing, order fulfillment, and cash reconciliation in a family trading business.")
            ]
        }
    ]

    for job in jobs:
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(4.5)
        p_head.paragraph_format.space_after = Pt(0)
        p_head.paragraph_format.keep_with_next = True
        
        r_role = p_head.add_run(job["role"])
        r_role.font.name = "Calibri"
        r_role.font.size = Pt(9.3)
        r_role.font.bold = True
        r_role.font.color.rgb = C_NAVY

        r_dates = p_head.add_run(f"  |  {job['dates']}")
        r_dates.font.name = "Calibri"
        r_dates.font.size = Pt(8.8)
        r_dates.font.bold = True
        r_dates.font.color.rgb = C_TEAL

        p_comp = doc.add_paragraph()
        p_comp.paragraph_format.space_before = Pt(0)
        p_comp.paragraph_format.space_after = Pt(1.5)
        p_comp.paragraph_format.keep_with_next = True
        r_comp = p_comp.add_run(job["company"])
        r_comp.font.name = "Calibri"
        r_comp.font.size = Pt(8.5)
        r_comp.font.italic = True
        r_comp.font.color.rgb = C_ACCENT

        for b_prefix, b_text in job["bullets"]:
            p_b = doc.add_paragraph()
            p_b.paragraph_format.left_indent = Inches(0.15)
            p_b.paragraph_format.space_before = Pt(0.5)
            p_b.paragraph_format.space_after = Pt(1)
            p_b.paragraph_format.line_spacing = 1.08
            p_b.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            
            r_bullet = p_b.add_run("• ")
            r_bullet.font.name = "Calibri"
            r_bullet.font.size = Pt(8.8)
            r_bullet.font.bold = True
            r_bullet.font.color.rgb = C_TEAL

            r_bold = p_b.add_run(b_prefix)
            r_bold.font.name = "Calibri"
            r_bold.font.size = Pt(8.8)
            r_bold.font.bold = True
            r_bold.font.color.rgb = C_NAVY

            r_text = p_b.add_run(b_text)
            r_text.font.name = "Calibri"
            r_text.font.size = Pt(8.8)
            r_text.font.color.rgb = C_BODY

    # Education
    add_section("Education", space_before=5)
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(3)
    p_edu.paragraph_format.space_after = Pt(0)
    p_edu.paragraph_format.keep_with_next = True
    r_deg = p_edu.add_run("Bachelor of Business Administration (BBA) — International Business")
    r_deg.font.name = "Calibri"
    r_deg.font.size = Pt(9.3)
    r_deg.font.bold = True
    r_deg.font.color.rgb = C_NAVY
    r_edates = p_edu.add_run("  |  2023 – 2026")
    r_edates.font.name = "Calibri"
    r_edates.font.size = Pt(8.8)
    r_edates.font.bold = True
    r_edates.font.color.rgb = C_TEAL

    p_inst = doc.add_paragraph()
    p_inst.paragraph_format.space_before = Pt(0)
    p_inst.paragraph_format.space_after = Pt(1.5)
    p_inst.paragraph_format.keep_with_next = True
    r_inst = p_inst.add_run("Dayananda Sagar University | Bengaluru, India")
    r_inst.font.name = "Calibri"
    r_inst.font.size = Pt(8.5)
    r_inst.font.italic = True
    r_inst.font.color.rgb = C_ACCENT

    p_edub = doc.add_paragraph()
    p_edub.paragraph_format.left_indent = Inches(0.15)
    p_edub.paragraph_format.space_before = Pt(0.5)
    p_edub.paragraph_format.space_after = Pt(2)
    p_edub.paragraph_format.line_spacing = 1.08
    r_b = p_edub.add_run("• ")
    r_b.font.color.rgb = C_TEAL
    r_b.font.bold = True
    r_t = p_edub.add_run(
        "Specialization: International Trade Policy, Incoterms 2020, Global Supply Chain Architecture, Foreign Exchange Management, "
        "Financial Management, Business Statistics. Applied Projects: Landed cost modeling, cross-border freight risk analysis, and automated supplier intelligence."
    )
    r_t.font.name = "Calibri"
    r_t.font.size = Pt(8.8)
    r_t.font.color.rgb = C_BODY

    # Certifications & Honors (Using compact 2-column borderless table)
    add_section("Certifications & Honors", space_before=5)
    
    certs = [
        ("Written Management Commendation (BD)", "Pencil Mark Interior Solutions (2025)"),
        ("Google Digital Marketing Professional", "Google (2024)"),
        ("Service Marketing: A Practical Approach", "NPTEL, IIT Kharagpur (2025)"),
        ("Generative AI Mastermind", "Outskill (2025)"),
        ("AI Tools & ChatGPT for Business Productivity", "be10x (2025)")
    ]

    table = doc.add_table(rows=3, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(3.55)
    table.columns[1].width = Inches(3.55)

    cert_idx = 0
    for r in range(3):
        for c in range(2):
            cell = table.cell(r, c)
            set_cell_margins(cell, top=20, bottom=20, left=40, right=40)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if cert_idx < len(certs):
                name, org = certs[cert_idx]
                r_sym = p.add_run("• ")
                r_sym.font.name = "Calibri"
                r_sym.font.size = Pt(8.5)
                r_sym.font.bold = True
                r_sym.font.color.rgb = C_TEAL

                r_n = p.add_run(f"{name} ")
                r_n.font.name = "Calibri"
                r_n.font.size = Pt(8.5)
                r_n.font.bold = True
                r_n.font.color.rgb = C_NAVY

                r_o = p.add_run(f"({org})")
                r_o.font.name = "Calibri"
                r_o.font.size = Pt(8.2)
                r_o.font.color.rgb = C_MUTED
                cert_idx += 1

    out_docx = BASE_DIR / "ADITYA_MEHRA_UNIVERSAL_MASTER_CV.docx"
    doc.save(out_docx)
    return out_docx

# ==========================================
# 2. BUILD HIGH-IMPACT 1-PAGE EXECUTIVE CV
# ==========================================
def build_1page_cv():
    doc = Document()
    
    # Configure A4 dimensions & compact margins
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.35)
        section.bottom_margin = Inches(0.35)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

    # Header
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    run_name = p_name.add_run("ADITYA MEHRA")
    run_name.font.name = "Calibri"
    run_name.font.size = Pt(18)
    run_name.font.bold = True
    run_name.font.color.rgb = C_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2)
    run_sub = p_sub.add_run("BUSINESS OPERATIONS  |  COMMERCIAL GROWTH  |  AI AUTOMATION")
    run_sub.font.name = "Calibri"
    run_sub.font.size = Pt(9)
    run_sub.font.bold = True
    run_sub.font.color.rgb = C_TEAL

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(4)
    run_contact = p_contact.add_run(
        "Bengaluru, India  |  +91 7003456624  |  adityamehra799@gmail.com  |  "
        "linkedin.com/in/aditya-mehra-b8644b326  |  github.com/AdityaMehra007"
    )
    run_contact.font.name = "Calibri"
    run_contact.font.size = Pt(8.2)
    run_contact.font.color.rgb = C_MUTED

    def add_section(title, space_before=4):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.keep_with_next = True
        run = p.add_run(title.upper())
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.bold = True
        run.font.color.rgb = C_NAVY
        add_clean_border(p)

    # Summary
    add_section("Professional Summary", space_before=2)
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(1)
    p_sum.paragraph_format.space_after = Pt(3)
    p_sum.paragraph_format.line_spacing = 1.08
    p_sum.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_sum = p_sum.add_run(
        "Results-driven Business Operations & Commercial Growth Specialist with extensive hands-on experience directing 300+ commercial activations, "
        "governing high-stakes operations at AERO India 2025 (100k+ attendees, 0 shrinkage), and training robotics AI data pipelines at Instawork (99%+ QA benchmark). "
        "Combines BBA International Business credentials (DSU '26) with proven execution across vendor SLA governance, 15–20% procurement cost reduction, "
        "B2B pipeline acquisition, and modern AI automation. Commended in writing for outstanding commercial drive."
    )
    r_sum.font.name = "Calibri"
    r_sum.font.size = Pt(8.3)
    r_sum.font.color.rgb = C_BODY

    # Competencies
    add_section("Core Competencies", space_before=3)
    p_c1 = doc.add_paragraph()
    p_c1.paragraph_format.left_indent = Inches(0.12)
    p_c1.paragraph_format.space_before = Pt(0)
    p_c1.paragraph_format.space_after = Pt(1)
    p_c1.paragraph_format.line_spacing = 1.05
    r = p_c1.add_run("• Operations & Cost Control: ")
    r.bold = True; r.font.color.rgb = C_NAVY; r.font.size = Pt(8.2)
    r2 = p_c1.add_run("Live Operations Delivery (300+ Events), Direct Supplier SLA Governance, 15–20% Landed Cost Reduction, Crisis Triage.")
    r2.font.size = Pt(8.2); r2.font.color.rgb = C_BODY

    p_c2 = doc.add_paragraph()
    p_c2.paragraph_format.left_indent = Inches(0.12)
    p_c2.paragraph_format.space_before = Pt(0)
    p_c2.paragraph_format.space_after = Pt(1)
    p_c2.paragraph_format.line_spacing = 1.05
    r = p_c2.add_run("• Commercial & B2B Growth: ")
    r.bold = True; r.font.color.rgb = C_NAVY; r.font.size = Pt(8.2)
    r2 = p_c2.add_run("Corporate Client Outreach (Tata Communications, Puma India), Pipeline Management (15+ Threads), High-Ticket Deal Negotiation.")
    r2.font.size = Pt(8.2); r2.font.color.rgb = C_BODY

    p_c3 = doc.add_paragraph()
    p_c3.paragraph_format.left_indent = Inches(0.12)
    p_c3.paragraph_format.space_before = Pt(0)
    p_c3.paragraph_format.space_after = Pt(2)
    p_c3.paragraph_format.line_spacing = 1.05
    r = p_c3.add_run("• AI & Trade Systems: ")
    r.bold = True; r.font.color.rgb = C_NAVY; r.font.size = Pt(8.2)
    r2 = p_c3.add_run("Robotics Data Curation (99%+ QA), GenAI Workflows, Incoterms 2020, Advanced Excel (XLOOKUP, Power Query), CRM (Salesforce/Zoho).")
    r2.font.size = Pt(8.2); r2.font.color.rgb = C_BODY

    # Experience
    add_section("Key Professional Experience", space_before=4)

    one_page_jobs = [
        {
            "role": "Independent Event Director & Commercial Operations Lead",
            "company": "Independent Commercial Operations | Bengaluru, India",
            "dates": "2019 – Present",
            "bullets": [
                ("Scaled Commercial Execution: ", "Directed end-to-end production for 300+ commercial events and brand activations for marquee clients including Tata Communications, Puma India, VH1 Supersonic, and Apollo Marketing."),
                ("Procurement & Cost Savings: ", "Delivered consistent 15% to 20% cost savings per engagement by eliminating intermediary contractor markups, establishing direct supplier rate cards, and enforcing strict vendor SLAs."),
                ("Client Retention & Resilience: ", "Maintained a 30%+ repeat-client rate and 100% on-time execution through rapid crisis triage (venue re-allocations and technical transfers under zero client SLA breach).")
            ]
        },
        {
            "role": "AI Data Operations Intern",
            "company": "Instawork Services India Pvt. Ltd. | Bengaluru, India",
            "dates": "Dec 2025",
            "bullets": [
                ("Robotics & Vision Datasets: ", "Supported ML training pipelines for enterprise robotics and human activity recognition, curating multi-modal sensor and video data streams."),
                ("Quality Assurance Benchmark: ", "Maintained a verified 99%+ quality assurance benchmark, performing granular schema validation and edge-case classification to eliminate training drift.")
            ]
        },
        {
            "role": "Business Development Intern",
            "company": "Pencil Mark Interior Solutions LLP | Bengaluru, India",
            "dates": "Jul 2025 – Aug 2025",
            "bullets": [
                ("Revenue & Pipeline Drive: ", "Accelerated B2B client acquisition across Bengaluru commercial interior projects, managing 15+ concurrent prospect threads and closing ₹1.5L+ in initial contract revenue (18% conversion rate)."),
                ("Executive Commendation: ", "Awarded a formal written management commendation for exemplary lead pipeline expansion, proactive client negotiation, and sales drive.")
            ]
        },
        {
            "role": "Exhibition Operations & Commercial Lead",
            "company": "Salt in My Coca — AERO India 2025 | Yelahanka Air Force Station, Bengaluru",
            "dates": "Feb 2025",
            "bullets": [
                ("High-Density Expo Logistics: ", "Commanded 7 full days of on-ground stall operations, contractor logistics, and visitor flows at Asia's premier defense expo (100,000+ attendees) with zero inventory shrinkage or asset loss under strict military clearances."),
                ("Stakeholder Lead Capture: ", "Interfaced with international military delegations, corporate CXOs, and institutional buyers, qualifying high-value commercial leads.")
            ]
        }
    ]

    for job in one_page_jobs:
        p_head = doc.add_paragraph()
        p_head.paragraph_format.space_before = Pt(3)
        p_head.paragraph_format.space_after = Pt(0)
        p_head.paragraph_format.keep_with_next = True
        
        r_role = p_head.add_run(job["role"])
        r_role.font.name = "Calibri"; r_role.font.size = Pt(8.8); r_role.bold = True; r_role.font.color.rgb = C_NAVY

        r_dates = p_head.add_run(f"  |  {job['dates']}")
        r_dates.font.name = "Calibri"; r_dates.font.size = Pt(8.2); r_dates.bold = True; r_dates.font.color.rgb = C_TEAL

        p_comp = doc.add_paragraph()
        p_comp.paragraph_format.space_before = Pt(0); p_comp.paragraph_format.space_after = Pt(1)
        p_comp.paragraph_format.keep_with_next = True
        r_comp = p_comp.add_run(job["company"])
        r_comp.font.name = "Calibri"; r_comp.font.size = Pt(8); r_comp.italic = True; r_comp.font.color.rgb = C_ACCENT

        for b_pre, b_txt in job["bullets"]:
            p_b = doc.add_paragraph()
            p_b.paragraph_format.left_indent = Inches(0.12)
            p_b.paragraph_format.space_before = Pt(0.5); p_b.paragraph_format.space_after = Pt(0.5)
            p_b.paragraph_format.line_spacing = 1.05
            p_b.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            rb = p_b.add_run("• ")
            rb.font.name = "Calibri"; rb.font.size = Pt(8.2); rb.bold = True; rb.font.color.rgb = C_TEAL

            rp = p_b.add_run(b_pre)
            rp.font.name = "Calibri"; rp.font.size = Pt(8.2); rp.bold = True; rp.font.color.rgb = C_NAVY

            rt = p_b.add_run(b_txt)
            rt.font.name = "Calibri"; rt.font.size = Pt(8.2); rt.font.color.rgb = C_BODY

    # Education & Credentials in a unified bottom grid
    add_section("Education & Credentials", space_before=4)
    
    p_eb = doc.add_paragraph()
    p_eb.paragraph_format.space_before = Pt(2); p_eb.paragraph_format.space_after = Pt(1)
    p_eb.paragraph_format.line_spacing = 1.05
    r_deg = p_eb.add_run("Bachelor of Business Administration (BBA) — International Business  |  2023 – 2026\n")
    r_deg.bold = True; r_deg.font.color.rgb = C_NAVY; r_deg.font.size = Pt(8.5)
    r_uni = p_eb.add_run("Dayananda Sagar University, Bengaluru — Focus: International Trade & EXIM Policies, Incoterms 2020, Supply Chain, Business Statistics.")
    r_uni.font.color.rgb = C_BODY; r_uni.font.size = Pt(8.2)

    p_cr = doc.add_paragraph()
    p_cr.paragraph_format.space_before = Pt(2); p_cr.paragraph_format.space_after = Pt(0)
    p_cr.paragraph_format.line_spacing = 1.05
    r_cb = p_cr.add_run("Certifications: ")
    r_cb.bold = True; r_cb.font.color.rgb = C_NAVY; r_cb.font.size = Pt(8.2)
    r_cl = p_cr.add_run("Written Commendation (Pencil Mark, 2025) • Google Digital Marketing (2024) • Service Marketing (IIT Kharagpur NPTEL, 2025) • GenAI Mastermind (Outskill, 2025)")
    r_cl.font.color.rgb = C_MUTED; r_cl.font.size = Pt(8.2)

    out_docx = BASE_DIR / "ADITYA_MEHRA_EXECUTIVE_1PAGE_CV.docx"
    doc.save(out_docx)
    return out_docx

# EXECUTE AND VERIFY IN WORD COM
def verify_and_export():
    docx_2p = build_2page_cv()
    docx_1p = build_1page_cv()
    print(f"Generated DOCX 2-Page: {docx_2p}")
    print(f"Generated DOCX 1-Page: {docx_1p}")

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False

    # Check 2-Page DOCX
    doc2 = word.Documents.Open(str(docx_2p))
    pages_2p = doc2.ComputeStatistics(2)
    print(f"Verified 2-Page DOCX Page Count: {pages_2p}")
    pdf_2p = BASE_DIR / "ADITYA_MEHRA_UNIVERSAL_MASTER_CV.pdf"
    doc2.ExportAsFixedFormat(str(pdf_2p), 17) # 17 = wdExportFormatPDF
    doc2.Close()
    print(f"Exported Native Word PDF (2-Page): {pdf_2p}")

    # Check 1-Page DOCX
    doc1 = word.Documents.Open(str(docx_1p))
    pages_1p = doc1.ComputeStatistics(2)
    print(f"Verified 1-Page DOCX Page Count: {pages_1p}")
    pdf_1p = BASE_DIR / "ADITYA_MEHRA_EXECUTIVE_1PAGE_CV.pdf"
    doc1.ExportAsFixedFormat(str(pdf_1p), 17)
    doc1.Close()
    print(f"Exported Native Word PDF (1-Page): {pdf_1p}")

    word.Quit()
    return pages_2p, pages_1p

if __name__ == "__main__":
    p2, p1 = verify_and_export()
    print(f"FINAL AUDIT: 2-Page Count = {p2}, 1-Page Count = {p1}")
