import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client
import pypdf

BASE_DIR = Path(r"e:\anti")

# Palette: Executive World-Class Typography
C_NAVY = RGBColor(15, 23, 42)       # #0f172a Deep Slate Navy
C_SLATE = RGBColor(51, 65, 85)      # #334155 Professional Slate
C_BODY = RGBColor(30, 41, 59)       # #1e293b High contrast body charcoal
C_MUTED = RGBColor(100, 116, 139)   # #64748b Metadata slate
C_LINK = "1E3A8A"                   # #1e3a8a Deep Royal Blue for hyperlinks

def add_clean_divider(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')       # 0.5 pt clean hairline
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CBD5E1')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_hyperlink_run(paragraph, url, display_text, font_size_pt=9.0, color_hex="1E3A8A", bold=False, underline=True):
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'<w:color w:val="{color_hex}"/>'
        f'{"<w:b/>" if bold else ""}'
        f'{"<w:u w:val=\"single\"/>" if underline else ""}'
        f'<w:sz w:val="{int(font_size_pt * 2)}"/>'
        f'</w:rPr>'
        f'<w:t>{display_text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    paragraph._p.append(hyperlink)

def build_full_cv(top_margin=0.42, bottom_margin=0.42, left_margin=0.52, right_margin=0.52,
                  sec_before=6.0, bullet_after=2.0, body_size=8.8, heading_size=9.8):
    doc = Document()

    for section in doc.sections:
        section.page_width = Inches(8.27)    # 210 mm
        section.page_height = Inches(11.69)  # 297 mm
        section.top_margin = Inches(top_margin)
        section.bottom_margin = Inches(bottom_margin)
        section.left_margin = Inches(left_margin)
        section.right_margin = Inches(right_margin)

    # 1. HEADER
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1.5)
    r_name = p_name.add_run("ADITYA MEHRA")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(21)
    r_name.font.bold = True
    r_name.font.color.rgb = C_NAVY

    # SUBTITLE
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2.5)
    r_sub = p_sub.add_run("BBA Graduate  •  Operations, Client Support & AI Productivity  •  Entry-Level")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.5)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_SLATE

    # CONTACT ROW 1: Location • Phone • Email
    p_con1 = doc.add_paragraph()
    p_con1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con1.paragraph_format.space_before = Pt(0)
    p_con1.paragraph_format.space_after = Pt(1.5)
    
    r_loc = p_con1.add_run("Bengaluru, Karnataka, India  •  ")
    r_loc.font.name = "Calibri"; r_loc.font.size = Pt(8.8); r_loc.font.color.rgb = C_MUTED
    
    r_ph = p_con1.add_run("Phone: ")
    r_ph.font.name = "Calibri"; r_ph.font.size = Pt(8.8); r_ph.font.bold = True; r_ph.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "tel:+917003456624", "+91 7003456624", font_size_pt=8.8, color_hex=C_LINK, underline=False)
    
    r_sep1 = p_con1.add_run("  •  Email: ")
    r_sep1.font.name = "Calibri"; r_sep1.font.size = Pt(8.8); r_sep1.font.bold = True; r_sep1.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "mailto:adityamehra799@gmail.com", "adityamehra799@gmail.com", font_size_pt=8.8, color_hex=C_LINK, underline=False)

    # CONTACT ROW 2: Portfolio • LinkedIn • GitHub
    p_con2 = doc.add_paragraph()
    p_con2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con2.paragraph_format.space_before = Pt(0)
    p_con2.paragraph_format.space_after = Pt(3.0)

    r_web_lbl = p_con2.add_run("Portfolio: ")
    r_web_lbl.font.name = "Calibri"; r_web_lbl.font.size = Pt(8.8); r_web_lbl.font.bold = True; r_web_lbl.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://adi-digital-universe.ai.studio/", "adi-digital-universe.ai.studio", font_size_pt=8.8, color_hex=C_LINK, underline=True)

    r_sep2 = p_con2.add_run("  •  LinkedIn: ")
    r_sep2.font.name = "Calibri"; r_sep2.font.size = Pt(8.8); r_sep2.font.bold = True; r_sep2.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://www.linkedin.com/in/aditya-mehra-b8644b326", "linkedin.com/in/aditya-mehra-b8644b326", font_size_pt=8.8, color_hex=C_LINK, underline=True)

    r_sep3 = p_con2.add_run("  •  GitHub: ")
    r_sep3.font.name = "Calibri"; r_sep3.font.size = Pt(8.8); r_sep3.font.bold = True; r_sep3.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://github.com/AdityaMehra007", "github.com/AdityaMehra007", font_size_pt=8.8, color_hex=C_LINK, underline=True)

    def add_sec(title, before=sec_before):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title.upper())
        r.font.name = "Calibri"
        r.font.size = Pt(heading_size)
        r.font.bold = True
        r.font.color.rgb = C_NAVY
        add_clean_divider(p)

    def add_bullet(p, title_bold, text_body, space_after=bullet_after):
        p.paragraph_format.left_indent = Inches(0.12)
        p.paragraph_format.space_before = Pt(0.5)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.08
        
        rb = p.add_run("• ")
        rb.font.name = "Calibri"; rb.font.size = Pt(body_size); rb.bold = True; rb.font.color.rgb = C_SLATE

        if title_bold:
            rt = p.add_run(title_bold)
            rt.font.name = "Calibri"; rt.font.size = Pt(body_size); rt.bold = True; rt.font.color.rgb = C_NAVY

        rc = p.add_run(text_body)
        rc.font.name = "Calibri"; rc.font.size = Pt(body_size); rc.font.color.rgb = C_BODY

    # 2. CAREER OBJECTIVE
    add_sec("Career Objective", before=2.0)
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(1.0)
    p_obj.paragraph_format.space_after = Pt(2.0)
    p_obj.paragraph_format.line_spacing = 1.10
    r_obj = p_obj.add_run(
        "Dedicated BBA graduate in International Business with hands-on experience across client outreach, operational coordination, and modern AI productivity workflows. "
        "Quick learner with strong computer proficiency, eager to support cross-functional teams in daily operations, customer service, vendor communication, and business administration."
    )
    r_obj.font.name = "Calibri"
    r_obj.font.size = Pt(body_size)
    r_obj.font.color.rgb = C_BODY

    # 3. WORK EXPERIENCE & INTERNSHIPS (Detailed 2-line entries for each role)
    add_sec("Work Experience & Internships", before=sec_before)
    
    # Role 1
    p_exp1_a = doc.add_paragraph()
    add_bullet(p_exp1_a, "Business Development Intern | Pencil Mark Solutions: ",
               "Conducted targeted outreach to prospective corporate clients via phone and email to introduce commercial interior design services.")
    p_exp1_b = doc.add_paragraph()
    add_bullet(p_exp1_b, "",
               "Assisted senior management with preparing service proposals, cost estimates, client meeting presentations, and follow-up schedules.")

    # Role 2
    p_exp2_a = doc.add_paragraph()
    add_bullet(p_exp2_a, "Operations Intern | Instawork Services India: ",
               "Organized, structured, and verified computer datasets to support engineering teams with robotics and AI operations.")
    p_exp2_b = doc.add_paragraph()
    add_bullet(p_exp2_b, "",
               "Maintained strict data accuracy standards by cross-checking spreadsheets and reporting workflow discrepancies promptly.")

    # Role 3
    p_exp3_a = doc.add_paragraph()
    add_bullet(p_exp3_a, "Stall Coordinator | AERO India Exhibition: ",
               "Managed exhibition booth setup, welcomed trade delegates and international visitors, answered inquiries, and maintained inventory control.")
    p_exp3_b = doc.add_paragraph()
    add_bullet(p_exp3_b, "",
               "Monitored promotional merchandise and product catalogues to ensure continuous availability throughout the multi-day event.")

    # Role 4
    p_exp4_a = doc.add_paragraph()
    add_bullet(p_exp4_a, "Event Coordinator | College & Music Events: ",
               "Coordinated with technical vendors for stage, sound, and lighting setups while managing hospitality and travel logistics for performing artists.")
    p_exp4_b = doc.add_paragraph()
    add_bullet(p_exp4_b, "",
               "Supervised on-ground event volunteers and maintained strict schedules to ensure seamless show operations and crowd management.")

    # Role 5
    p_exp5_a = doc.add_paragraph()
    add_bullet(p_exp5_a, "Store Assistant | Family Retail Store: ",
               "Managed daily counter sales, customer billing, cash registers, and supplier restocking to maintain seamless store operations.")
    p_exp5_b = doc.add_paragraph()
    add_bullet(p_exp5_b, "",
               "Communicated regularly with wholesale distributors to place repeat merchandise orders and verify incoming delivery receipts.")

    # 4. PROJECTS & AI PROOF OF WORK
    add_sec("Projects & AI Proof of Work", before=sec_before)
    p_p1_a = doc.add_paragraph()
    add_bullet(p_p1_a, "Web Apps & Digital Portfolio (adi-digital-universe.ai.studio): ",
               "Built and deployed interactive web applications and a personal portfolio website using modern AI coding workflows and cloud hosting.")
    p_p1_b = doc.add_paragraph()
    add_bullet(p_p1_b, "",
               "Maintained active public GitHub repositories showcasing code implementations, interface prototypes, and practical automation scripts.")

    p_p2_a = doc.add_paragraph()
    add_bullet(p_p2_a, "AI Multimedia & Content Production: ",
               "Produced promotional video assets, visual marketing collateral, and digital presentation decks using modern generative AI tools.")
    p_p2_b = doc.add_paragraph()
    add_bullet(p_p2_b, "",
               "Accelerated creative content delivery by combining AI image generation workflows, video storyboard drafting, and prompt refinement.")

    # 5. EDUCATION
    add_sec("Education", before=sec_before)
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(1.0)
    p_edu.paragraph_format.space_after = Pt(2.0)
    p_edu.paragraph_format.line_spacing = 1.08
    p_edu.paragraph_format.keep_with_next = True

    r_deg = p_edu.add_run("Bachelor of Business Administration (BBA) — International Business  •  ")
    r_deg.font.name = "Calibri"; r_deg.font.size = Pt(body_size + 0.2); r_deg.bold = True; r_deg.font.color.rgb = C_NAVY

    r_uni = p_edu.add_run("Dayananda Sagar University, Bengaluru\n")
    r_uni.font.name = "Calibri"; r_uni.font.size = Pt(body_size); r_uni.font.italic = True; r_uni.font.color.rgb = C_SLATE

    r_sub = p_edu.add_run("Relevant Coursework: International Trade, Marketing Management, Supply Chain Logistics, Business Communication, Accounting.")
    r_sub.font.name = "Calibri"; r_sub.font.size = Pt(body_size - 0.4); r_sub.font.color.rgb = C_BODY

    # 6. KEY SKILLS & COMPETENCIES
    add_sec("Key Skills & Competencies", before=sec_before)
    skills = [
        ("Computer & AI Tools: ", "ChatGPT, Generative AI Platforms, MS Excel, MS Word, MS PowerPoint, Google Workspace, GitHub, Web Research."),
        ("Workplace Competencies: ", "Client Communication, Customer Relations, Operations Support, Vendor Coordination, Team Collaboration, Problem Solving."),
        ("Languages: ", "English, Hindi.")
    ]
    for s_title, s_desc in skills:
        p_s = doc.add_paragraph()
        add_bullet(p_s, s_title, s_desc)

    # 7. CERTIFICATES & HONORS
    add_sec("Certificates & Honors", before=sec_before)
    certs = [
        ("Internship Certificate: ", "Pencil Mark Solutions (Client Outreach & Business Development)."),
        ("Fundamentals of Digital Marketing: ", "Google (Online Marketing, SEO & Web Analytics)."),
        ("Marketing Management Certification: ", "NPTEL, IIT Kharagpur (Marketing Principles & Consumer Behavior)."),
        ("AI Productivity Workshop: ", "Outskill (Generative AI Workflows & Modern Office Automation).")
    ]
    for c_title, c_desc in certs:
        p_c = doc.add_paragraph()
        add_bullet(p_c, c_title, c_desc, space_after=1.5)

    test_docx = BASE_DIR / "test_full_cv.docx"
    doc.save(test_docx)
    return test_docx

if __name__ == "__main__":
    docx_path = build_full_cv()
    print("Saved test docx:", docx_path)
