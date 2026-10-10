import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client
import pypdf

BASE_DIR = Path(r"e:\anti")

# Corporate Palette (Executive Slate Navy)
C_NAVY = RGBColor(15, 23, 42)       # #0F172A Deep Navy
C_SLATE = RGBColor(51, 65, 85)      # #334155 Slate
C_BODY = RGBColor(30, 41, 59)       # #1E293B Body
C_MUTED = RGBColor(100, 116, 139)   # #64748B Secondary
C_LINK = "1E3A8A"                   # #1E3A8A Classic Blue

def add_clean_divider(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')       # 0.5 pt
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CBD5E1')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_hyperlink_run(paragraph, url, display_text, font_size_pt=8.5, color_hex="1E3A8A", bold=False, underline=True):
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'<w:color w:val="{color_hex}"/>'
        f'{"<w:b/>" if bold else ""}'
        f'{"<w:u w:val=\"single\"/>" if underline else ""}'
        f'<w:sz w:val="{int(font_size_pt * 2)}"/>'
        f'</w:rPr>'
        f'<w:t>{display_text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    paragraph._p.append(hyperlink)

def add_bullet_item(doc, prefix, body, space_after=1.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.line_spacing = 1.09
    
    r_sym = p.add_run("•  ")
    r_sym.font.name = "Calibri"
    r_sym.font.size = Pt(8.5)
    r_sym.font.bold = True
    r_sym.font.color.rgb = C_SLATE
    
    if prefix:
        r_pre = p.add_run(prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(8.5)
        r_pre.font.bold = True
        r_pre.font.color.rgb = C_NAVY
        
    r_body = p.add_run(body)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(8.5)
    r_body.font.color.rgb = C_BODY

def build_meta_sc_docx(filename):
    doc = Document()
    
    # Standard A4 Margins calibrated precisely for 1-page vertical density (94-96% coverage)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.36)
        section.bottom_margin = Inches(0.36)
        section.left_margin = Inches(0.48)
        section.right_margin = Inches(0.48)
        
    # 1. HEADER - NAME
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1.0)
    r_name = p_name.add_run("ADITYA MEHRA")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(19.0)
    r_name.font.bold = True
    r_name.font.color.rgb = C_NAVY
    
    # SUBTITLE (Exact Requisition Match: Allegis Global Solutions / Meta SCM)
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2.0)
    r_sub = p_sub.add_run("Supply Chain Onboarding Associate (SCOA)  |  Meta Global SCM Program Track")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.3)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_SLATE
    
    # CONTACT ROW 1: Location | Phone | Email
    p_con1 = doc.add_paragraph()
    p_con1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con1.paragraph_format.space_before = Pt(0)
    p_con1.paragraph_format.space_after = Pt(1.5)
    
    r_loc = p_con1.add_run("Bengaluru, Karnataka, India (Commerce @ Mantri Proximity)  |  ")
    r_loc.font.name = "Calibri"; r_loc.font.size = Pt(8.5); r_loc.font.color.rgb = C_MUTED
    
    r_ph = p_con1.add_run("Phone: ")
    r_ph.font.name = "Calibri"; r_ph.font.size = Pt(8.5); r_ph.font.bold = True; r_ph.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "tel:+917003456624", "+91 7003456624", font_size_pt=8.5, color_hex=C_LINK, underline=False)
    
    r_sep1 = p_con1.add_run("  |  Email: ")
    r_sep1.font.name = "Calibri"; r_sep1.font.size = Pt(8.5); r_sep1.font.bold = True; r_sep1.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "mailto:adityamehra799@gmail.com", "adityamehra799@gmail.com", font_size_pt=8.5, color_hex=C_LINK, underline=False)
    
    # CONTACT ROW 2: Portfolio | LinkedIn | GitHub
    p_con2 = doc.add_paragraph()
    p_con2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con2.paragraph_format.space_before = Pt(0)
    p_con2.paragraph_format.space_after = Pt(3.0)
    
    r_w_lbl = p_con2.add_run("Portfolio: ")
    r_w_lbl.font.name = "Calibri"; r_w_lbl.font.size = Pt(8.5); r_w_lbl.font.bold = True; r_w_lbl.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://adi-digital-universe.ai.studio/", "adi-digital-universe.ai.studio", font_size_pt=8.5, color_hex=C_LINK, underline=True)
    
    r_sep2 = p_con2.add_run("  |  LinkedIn: ")
    r_sep2.font.name = "Calibri"; r_sep2.font.size = Pt(8.5); r_sep2.font.bold = True; r_sep2.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://www.linkedin.com/in/aditya-mehra-b8644b326", "linkedin.com/in/aditya-mehra-b8644b326", font_size_pt=8.5, color_hex=C_LINK, underline=True)
    
    r_sep3 = p_con2.add_run("  |  GitHub: ")
    r_sep3.font.name = "Calibri"; r_sep3.font.size = Pt(8.5); r_sep3.font.bold = True; r_sep3.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://github.com/AdityaMehra007", "github.com/AdityaMehra007", font_size_pt=8.5, color_hex=C_LINK, underline=True)
    
    def add_sec_header(title, before=3.2):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(1.8)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title.upper())
        r.font.name = "Calibri"
        r.font.size = Pt(9.3)
        r.font.bold = True
        r.font.color.rgb = C_NAVY
        add_clean_divider(p)

    # 2. PROFESSIONAL PROFILE & CAREER OBJECTIVE
    add_sec_header("Professional Profile & Career Objective", before=1.8)
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(1.0)
    p_obj.paragraph_format.space_after = Pt(2.0)
    p_obj.paragraph_format.line_spacing = 1.10
    r_obj = p_obj.add_run(
        "Detail-oriented BBA graduate (International Business) specifically targeting the Onboarding Associate — Meta (Global) "
        "Supply Chain requisition at Allegis Global Solutions (Req ID: 744000152372959). Brings direct practical experience in end-to-end "
        "supplier coordination, document vetting, compliance record verification (contracts, NDAs, certificates), and strict SOP adherence. "
        "Proven ability to manage high-volume trackers in MS Excel, drive timely follow-ups across cross-functional stakeholders "
        "(Finance, Risk & Compliance, Operations), and resolve onboarding bottlenecks within established SLAs. Fully available for the "
        "3:00 PM – 12:00 AM shift schedule and hybrid operations at Commerce @ Mantri, Bengaluru."
    )
    r_obj.font.name = "Calibri"
    r_obj.font.size = Pt(8.5)
    r_obj.font.color.rgb = C_BODY

    # 3. EDUCATION
    add_sec_header("Education", before=2.8)
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(1.0)
    p_edu.paragraph_format.space_after = Pt(1.0)
    p_edu.paragraph_format.line_spacing = 1.09
    p_edu.paragraph_format.keep_with_next = True
    
    r_deg = p_edu.add_run("Bachelor of Business Administration (BBA) — International Business (2023 – 2026)\n")
    r_deg.font.name = "Calibri"; r_deg.font.size = Pt(8.8); r_deg.bold = True; r_deg.font.color.rgb = C_NAVY
    
    r_uni = p_edu.add_run("Dayananda Sagar University (DSU), Bengaluru\n")
    r_uni.font.name = "Calibri"; r_uni.font.size = Pt(8.5); r_uni.font.italic = True; r_uni.font.color.rgb = C_SLATE
    
    r_crs = p_edu.add_run("Relevant SCM Coursework: Supply Chain Management, Global Logistics & Trade, Business Law & Corporate Compliance, Operations Management, Business Communication.")
    r_crs.font.name = "Calibri"; r_crs.font.size = Pt(8.2); r_crs.font.color.rgb = C_BODY

    # 4. WORK EXPERIENCE & OPERATIONAL INTERNSHIPS
    add_sec_header("Work Experience & Operational Internships", before=3.0)
    
    experiences = [
        {
            "role": "Operations Intern (Data Verification, Auditing & QA)",
            "company": "Instawork Services India, Bengaluru",
            "bullets": [
                "Managed high-volume data verification and document vetting workflows for robotics and AI enterprise datasets, maintaining 98%+ precision.",
                "Audited incoming partner submissions against rigorous SOP guidelines, identifying data inconsistencies, missing fields, and compliance exceptions.",
                "Maintained structured verification logs and escalated operational bottlenecks to team leads to prevent pipeline handoff delays and preserve turnaround SLAs."
            ]
        },
        {
            "role": "Operations & Supplier Coordination Intern",
            "company": "Pencil Mark Interior Solutions, Bengaluru",
            "bullets": [
                "Coordinated end-to-end commercial supplier vetting for turnkey interior projects, gathering vendor registration profiles, compliance documents, and NDAs.",
                "Reviewed supplier material quotations against contract specifications, proactively following up to resolve billing discrepancies and delivery delays.",
                "Maintained centralized vendor tracking sheets in MS Excel; awarded a written Letter of Commendation from executive leadership for exceptional diligence."
            ]
        },
        {
            "role": "Vendor Logistics & Compliance Coordinator",
            "company": "AERO India Exhibition (Salt in My Coca), Bengaluru",
            "bullets": [
                "Facilitated on-site vendor onboarding, gate pass compliance, and stall logistics under high-security defense exhibition protocols.",
                "Inspected incoming vendor consignments, reconciled delivery manifests against physical inventory, and ensured zero loss or compliance breach."
            ]
        },
        {
            "role": "Vendor Operations & Event Logistics Coordinator",
            "company": "Commercial & College Events, Bengaluru",
            "bullets": [
                "Managed multi-vendor coordination across staging, sound, lighting, and venue infrastructure suppliers for large-scale live productions.",
                "Enforced strict vendor onboarding schedules, verified SLA deliverables against signed service contracts, and resolved on-ground supplier issues."
            ]
        },
        {
            "role": "Supplier Procurement & Inventory Assistant",
            "company": "Family Retail Store & Food Business, Bengaluru",
            "bullets": [
                "Handled daily supplier order placements, invoice reconciliation, vendor payment registers, and physical inventory auditing.",
                "Liaised with wholesale distributors to monitor stock replenishment cycles, resolving delivery discrepancies and ensuring continuous supply continuity."
            ]
        }
    ]
    
    for exp in experiences:
        p_hdr = doc.add_paragraph()
        p_hdr.paragraph_format.space_before = Pt(1.8)
        p_hdr.paragraph_format.space_after = Pt(0.6)
        p_hdr.paragraph_format.keep_with_next = True
        
        r_r = p_hdr.add_run(exp["role"])
        r_r.font.name = "Calibri"; r_r.font.size = Pt(8.7); r_r.font.bold = True; r_r.font.color.rgb = C_NAVY
        
        r_div = p_hdr.add_run("  |  ")
        r_div.font.name = "Calibri"; r_div.font.size = Pt(8.7); r_div.font.color.rgb = C_MUTED
        
        r_c = p_hdr.add_run(exp["company"])
        r_c.font.name = "Calibri"; r_c.font.size = Pt(8.5); r_c.font.italic = True; r_c.font.color.rgb = C_SLATE
        
        for b in exp["bullets"]:
            add_bullet_item(doc, "", b, space_after=0.8)

    # 5. CORE COMPETENCIES & SCM REQUISITION KEYWORDS
    add_sec_header("Core Competencies & SCM Requisition Keywords", before=2.8)
    skills = [
        ("Supplier Onboarding & Vetting", "End-to-End Supplier Onboarding, Document Vetting, Compliance Verification (NDAs, Contracts, Insurance, Registrations), SLA Tracking, SOP Adherence, Exception Escalation."),
        ("Systems, Data & VMS Platforms", "MS Excel (VLOOKUP, Pivot Tables, Data Formatting, Reconciliation), MS Office Suite, Google Workspace, Supplier Portals & VMS Navigation Concepts, Dashboard Maintenance."),
        ("Operational & Communication Skills", "Cross-Functional Collaboration (Finance, Risk & Compliance, Client Delivery), Stakeholder Follow-ups, Audit & Profile Governance, 3 PM - 12 AM Shift Adaptability, Fluent English & Hindi.")
    ]
    for cat_name, skill_str in skills:
        add_bullet_item(doc, cat_name + ":", skill_str, space_after=0.8)

    # 6. PROFESSIONAL CERTIFICATIONS & HONORS
    add_sec_header("Professional Certifications & Honors", before=2.6)
    certs = [
        ("Internship Certificate & Letter of Commendation", "Pencil Mark Solutions (Commercial Supplier Coordination & Operations Diligence)."),
        ("Fundamentals of Digital Marketing Certification", "Google (E-Commerce Operations, Web Tracking & Digital Data Verification)."),
        ("Operations & Marketing Management", "NPTEL, IIT Kharagpur (Supply Chain Principles, Service Operations & Logistics)."),
        ("AI Productivity & Modern Office Automation", "Outskill (Modern Workflow Streamlining, Automated Data Verification & Office Tech).")
    ]
    for c_title, c_issuer in certs:
        add_bullet_item(doc, c_title + ":", c_issuer, space_after=0.6)
        
    doc.save(filename)
    print(f"Saved DOCX: {filename}")
    return filename

def build_meta_sc_html(filename):
    html = """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Aditya Mehra - Supply Chain Onboarding Associate (Meta Global SCM)</title>
    <style>
        @page { size: A4; margin: 8mm 12mm; }
        * { box-sizing: border-box; margin: 0; padding: 0; }
        body {
            font-family: 'Calibri', 'Segoe UI', Arial, sans-serif;
            color: #1E293B;
            background: #FFFFFF;
            line-height: 1.25;
            font-size: 8.8pt;
            padding: 6mm 10mm;
            max-width: 210mm;
            margin: 0 auto;
        }
        h1 {
            font-size: 19pt;
            font-weight: 700;
            color: #0F172A;
            text-align: center;
            letter-spacing: 0.5px;
            margin-bottom: 2px;
        }
        .subtitle {
            font-size: 9.3pt;
            font-weight: 700;
            color: #334155;
            text-align: center;
            margin-bottom: 3px;
        }
        .contact-bar {
            text-align: center;
            font-size: 8.5pt;
            color: #64748B;
            margin-bottom: 2px;
        }
        .contact-bar a {
            color: #1E3A8A;
            text-decoration: underline;
        }
        .section-title {
            font-size: 9.3pt;
            font-weight: 700;
            color: #0F172A;
            text-transform: uppercase;
            letter-spacing: 0.6px;
            border-bottom: 1px solid #CBD5E1;
            padding-bottom: 1px;
            margin-top: 6px;
            margin-bottom: 3px;
        }
        p.objective {
            font-size: 8.5pt;
            color: #1E293B;
            text-align: justify;
            margin-bottom: 3px;
            line-height: 1.24;
        }
        .edu-degree {
            font-size: 8.8pt;
            font-weight: 700;
            color: #0F172A;
        }
        .edu-school {
            font-size: 8.5pt;
            font-style: italic;
            color: #334155;
            margin-bottom: 1px;
        }
        .edu-coursework {
            font-size: 8.2pt;
            color: #1E293B;
        }
        .exp-block {
            margin-bottom: 3px;
        }
        .exp-header {
            font-size: 8.7pt;
            margin-bottom: 1px;
        }
        .exp-header .role {
            font-weight: 700;
            color: #0F172A;
        }
        .exp-header .sep {
            color: #94A3B8;
            margin: 0 4px;
        }
        .exp-header .company {
            font-style: italic;
            color: #334155;
        }
        ul {
            list-style: none;
            padding-left: 0;
        }
        li {
            position: relative;
            padding-left: 14px;
            font-size: 8.5pt;
            color: #1E293B;
            margin-bottom: 1.5px;
            line-height: 1.22;
        }
        li::before {
            content: "•";
            position: absolute;
            left: 2px;
            color: #334155;
            font-weight: bold;
        }
    </style>
</head>
<body>
    <h1>ADITYA MEHRA</h1>
    <div class="subtitle">Supply Chain Onboarding Associate (SCOA) | Meta Global SCM Program Track</div>
    <div class="contact-bar">
        Bengaluru, Karnataka, India (Commerce @ Mantri Proximity) &nbsp;|&nbsp; 
        <strong>Phone:</strong> <a href="tel:+917003456624">+91 7003456624</a> &nbsp;|&nbsp; 
        <strong>Email:</strong> <a href="mailto:adityamehra799@gmail.com">adityamehra799@gmail.com</a>
    </div>
    <div class="contact-bar">
        <strong>Portfolio:</strong> <a href="https://adi-digital-universe.ai.studio/">adi-digital-universe.ai.studio</a> &nbsp;|&nbsp; 
        <strong>LinkedIn:</strong> <a href="https://www.linkedin.com/in/aditya-mehra-b8644b326">linkedin.com/in/aditya-mehra-b8644b326</a> &nbsp;|&nbsp; 
        <strong>GitHub:</strong> <a href="https://github.com/AdityaMehra007">github.com/AdityaMehra007</a>
    </div>

    <div class="section-title">Professional Profile & Career Objective</div>
    <p class="objective">
        Detail-oriented BBA graduate (International Business) specifically targeting the Onboarding Associate — Meta (Global) Supply Chain requisition at Allegis Global Solutions (Req ID: 744000152372959). Brings direct practical experience in end-to-end supplier coordination, document vetting, compliance record verification (contracts, NDAs, certificates), and strict SOP adherence. Proven ability to manage high-volume trackers in MS Excel, drive timely follow-ups across cross-functional stakeholders (Finance, Risk & Compliance, Operations), and resolve onboarding bottlenecks within established SLAs. Fully available for the 3:00 PM – 12:00 AM shift schedule and hybrid operations at Commerce @ Mantri, Bengaluru.
    </p>

    <div class="section-title">Education</div>
    <div class="edu-degree">Bachelor of Business Administration (BBA) — International Business (2023 – 2026)</div>
    <div class="edu-school">Dayananda Sagar University (DSU), Bengaluru</div>
    <div class="edu-coursework"><strong>Relevant SCM Coursework:</strong> Supply Chain Management, Global Logistics & Trade, Business Law & Corporate Compliance, Operations Management, Business Communication.</div>

    <div class="section-title">Work Experience & Operational Internships</div>
    
    <div class="exp-block">
        <div class="exp-header">
            <span class="role">Operations Intern (Data Verification, Auditing & QA)</span>
            <span class="sep">|</span>
            <span class="company">Instawork Services India, Bengaluru</span>
        </div>
        <ul>
            <li>Managed high-volume data verification and document vetting workflows for robotics and AI enterprise datasets, maintaining 98%+ precision.</li>
            <li>Audited incoming partner submissions against rigorous SOP guidelines, identifying data inconsistencies, missing fields, and compliance exceptions.</li>
            <li>Maintained structured verification logs and escalated operational bottlenecks to team leads to prevent pipeline handoff delays and preserve turnaround SLAs.</li>
        </ul>
    </div>

    <div class="exp-block">
        <div class="exp-header">
            <span class="role">Operations & Supplier Coordination Intern</span>
            <span class="sep">|</span>
            <span class="company">Pencil Mark Interior Solutions, Bengaluru</span>
        </div>
        <ul>
            <li>Coordinated end-to-end commercial supplier vetting for turnkey interior projects, gathering vendor registration profiles, compliance documents, and NDAs.</li>
            <li>Reviewed supplier material quotations against contract specifications, proactively following up to resolve billing discrepancies and delivery delays.</li>
            <li>Maintained centralized vendor tracking sheets in MS Excel; awarded a written Letter of Commendation from executive leadership for exceptional diligence.</li>
        </ul>
    </div>

    <div class="exp-block">
        <div class="exp-header">
            <span class="role">Vendor Logistics & Compliance Coordinator</span>
            <span class="sep">|</span>
            <span class="company">AERO India Exhibition (Salt in My Coca), Bengaluru</span>
        </div>
        <ul>
            <li>Facilitated on-site vendor onboarding, gate pass compliance, and stall logistics under high-security defense exhibition protocols.</li>
            <li>Inspected incoming vendor consignments, reconciled delivery manifests against physical inventory, and ensured zero loss or compliance breach.</li>
        </ul>
    </div>

    <div class="exp-block">
        <div class="exp-header">
            <span class="role">Vendor Operations & Event Logistics Coordinator</span>
            <span class="sep">|</span>
            <span class="company">Commercial & College Events, Bengaluru</span>
        </div>
        <ul>
            <li>Managed multi-vendor coordination across staging, sound, lighting, and venue infrastructure suppliers for large-scale live productions.</li>
            <li>Enforced strict vendor onboarding schedules, verified SLA deliverables against signed service contracts, and resolved on-ground supplier issues.</li>
        </ul>
    </div>

    <div class="exp-block">
        <div class="exp-header">
            <span class="role">Supplier Procurement & Inventory Assistant</span>
            <span class="sep">|</span>
            <span class="company">Family Retail Store & Food Business, Bengaluru</span>
        </div>
        <ul>
            <li>Handled daily supplier order placements, invoice reconciliation, vendor payment registers, and physical inventory auditing.</li>
            <li>Liaised with wholesale distributors to monitor stock replenishment cycles, resolving delivery discrepancies and ensuring continuous supply continuity.</li>
        </ul>
    </div>

    <div class="section-title">Core Competencies & SCM Requisition Keywords</div>
    <ul>
        <li><strong>Supplier Onboarding & Vetting:</strong> End-to-End Supplier Onboarding, Document Vetting, Compliance Verification (NDAs, Contracts, Insurance, Registrations), SLA Tracking, SOP Adherence, Exception Escalation.</li>
        <li><strong>Systems, Data & VMS Platforms:</strong> MS Excel (VLOOKUP, Pivot Tables, Data Formatting, Reconciliation), MS Office Suite, Google Workspace, Supplier Portals & VMS Navigation Concepts, Dashboard Maintenance.</li>
        <li><strong>Operational & Communication Skills:</strong> Cross-Functional Collaboration (Finance, Risk & Compliance, Client Delivery), Stakeholder Follow-ups, Audit & Profile Governance, 3 PM - 12 AM Shift Adaptability, Fluent English & Hindi.</li>
    </ul>

    <div class="section-title">Professional Certifications & Honors</div>
    <ul>
        <li><strong>Internship Certificate & Letter of Commendation:</strong> Pencil Mark Solutions (Commercial Supplier Coordination & Operations Diligence).</li>
        <li><strong>Fundamentals of Digital Marketing Certification:</strong> Google (E-Commerce Operations, Web Tracking & Digital Data Verification).</li>
        <li><strong>Operations & Marketing Management:</strong> NPTEL, IIT Kharagpur (Supply Chain Principles, Service Operations & Logistics).</li>
        <li><strong>AI Productivity & Modern Office Automation:</strong> Outskill (Modern Workflow Streamlining, Automated Data Verification & Office Tech).</li>
    </ul>
</body>
</html>
"""
    out_path = BASE_DIR / filename
    out_path.write_text(html, encoding="utf-8")
    print(f"Saved HTML: {filename}")
    return str(out_path)

def build_meta_sc_md(filename):
    md = """# ADITYA MEHRA
**Supply Chain Onboarding Associate (SCOA) | Meta Global SCM Program Track**  
Bengaluru, Karnataka, India (Commerce @ Mantri Proximity) | Phone: [+91 7003456624](tel:+917003456624) | Email: [adityamehra799@gmail.com](mailto:adityamehra799@gmail.com)  
Portfolio: [adi-digital-universe.ai.studio](https://adi-digital-universe.ai.studio/) | LinkedIn: [aditya-mehra-b8644b326](https://www.linkedin.com/in/aditya-mehra-b8644b326) | GitHub: [AdityaMehra007](https://github.com/AdityaMehra007)

---

## PROFESSIONAL PROFILE & CAREER OBJECTIVE
Detail-oriented BBA graduate (International Business) specifically targeting the Onboarding Associate — Meta (Global) Supply Chain requisition at Allegis Global Solutions (Req ID: 744000152372959). Brings direct practical experience in end-to-end supplier coordination, document vetting, compliance record verification (contracts, NDAs, certificates), and strict SOP adherence. Proven ability to manage high-volume trackers in MS Excel, drive timely follow-ups across cross-functional stakeholders (Finance, Risk & Compliance, Operations), and resolve onboarding bottlenecks within established SLAs. Fully available for the 3:00 PM – 12:00 AM shift schedule and hybrid operations at Commerce @ Mantri, Bengaluru.

---

## EDUCATION
**Bachelor of Business Administration (BBA) — International Business (2023 – 2026)**  
*Dayananda Sagar University (DSU), Bengaluru*  
- **Relevant SCM Coursework:** Supply Chain Management, Global Logistics & Trade, Business Law & Corporate Compliance, Operations Management, Business Communication.

---

## WORK EXPERIENCE & OPERATIONAL INTERNSHIPS

### Operations Intern (Data Verification, Auditing & QA) | *Instawork Services India, Bengaluru*
- Managed high-volume data verification and document vetting workflows for robotics and AI enterprise datasets, maintaining 98%+ precision.
- Audited incoming partner submissions against rigorous SOP guidelines, identifying data inconsistencies, missing fields, and compliance exceptions.
- Maintained structured verification logs and escalated operational bottlenecks to team leads to prevent pipeline handoff delays and preserve turnaround SLAs.

### Operations & Supplier Coordination Intern | *Pencil Mark Interior Solutions, Bengaluru*
- Coordinated end-to-end commercial supplier vetting for turnkey interior projects, gathering vendor registration profiles, compliance documents, and NDAs.
- Reviewed supplier material quotations against contract specifications, proactively following up to resolve billing discrepancies and delivery delays.
- Maintained centralized vendor tracking sheets in MS Excel; awarded a written Letter of Commendation from executive leadership for exceptional diligence.

### Vendor Logistics & Compliance Coordinator | *AERO India Exhibition (Salt in My Coca), Bengaluru*
- Facilitated on-site vendor onboarding, gate pass compliance, and stall logistics under high-security defense exhibition protocols.
- Inspected incoming vendor consignments, reconciled delivery manifests against physical inventory, and ensured zero loss or compliance breach.

### Vendor Operations & Event Logistics Coordinator | *Commercial & College Events, Bengaluru*
- Managed multi-vendor coordination across staging, sound, lighting, and venue infrastructure suppliers for large-scale live productions.
- Enforced strict vendor onboarding schedules, verified SLA deliverables against signed service contracts, and resolved on-ground supplier issues.

### Supplier Procurement & Inventory Assistant | *Family Retail Store & Food Business, Bengaluru*
- Handled daily supplier order placements, invoice reconciliation, vendor payment registers, and physical inventory auditing.
- Liaised with wholesale distributors to monitor stock replenishment cycles, resolving delivery discrepancies and ensuring continuous supply continuity.

---

## CORE COMPETENCIES & SCM REQUISITION KEYWORDS
- **Supplier Onboarding & Vetting:** End-to-End Supplier Onboarding, Document Vetting, Compliance Verification (NDAs, Contracts, Insurance, Registrations), SLA Tracking, SOP Adherence, Exception Escalation.
- **Systems, Data & VMS Platforms:** MS Excel (VLOOKUP, Pivot Tables, Data Formatting, Reconciliation), MS Office Suite, Google Workspace, Supplier Portals & VMS Navigation Concepts, Dashboard Maintenance.
- **Operational & Communication Skills:** Cross-Functional Collaboration (Finance, Risk & Compliance, Client Delivery), Stakeholder Follow-ups, Audit & Profile Governance, 3 PM - 12 AM Shift Adaptability, Fluent English & Hindi.

---

## PROFESSIONAL CERTIFICATIONS & HONORS
- **Internship Certificate & Letter of Commendation:** Pencil Mark Solutions (Commercial Supplier Coordination & Operations Diligence).
- **Fundamentals of Digital Marketing Certification:** Google (E-Commerce Operations, Web Tracking & Digital Data Verification).
- **Operations & Marketing Management:** NPTEL, IIT Kharagpur (Supply Chain Principles, Service Operations & Logistics).
- **AI Productivity & Modern Office Automation:** Outskill (Modern Workflow Streamlining, Automated Data Verification & Office Tech).
"""
    out_path = BASE_DIR / filename
    out_path.write_text(md, encoding="utf-8")
    print(f"Saved MD: {filename}")
    return str(out_path)

def main():
    print("=" * 75)
    print("BUILDING ULTIMATE META SCM ONBOARDING CV SUITE (AGS REQUISITION MATCH)")
    print("=" * 75)
    
    file_prefix = "ADITYA_MEHRA_META_SCM_ONBOARDING_CV"
    docx_file = BASE_DIR / f"{file_prefix}.docx"
    html_file = f"{file_prefix}.html"
    md_file = f"{file_prefix}.md"
    pdf_file = BASE_DIR / f"{file_prefix}.pdf"
    
    build_meta_sc_docx(str(docx_file))
    build_meta_sc_html(html_file)
    build_meta_sc_md(md_file)
    
    # Word COM Validation & PDF Export
    print("\nValidating with Microsoft Word COM Engine...")
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    
    doc = word.Documents.Open(str(docx_file))
    pages = doc.ComputeStatistics(2)
    print(f"Word COM Computed Pages: {pages}")
    
    doc.ExportAsFixedFormat(str(pdf_file), 17)
    doc.Close()
    word.Quit()
    print(f"Exported PDF: {pdf_file}")
    
    # PyPDF Validation
    reader = pypdf.PdfReader(str(pdf_file))
    pdf_pages = len(reader.pages)
    print(f"PyPDF Verified Pages: {pdf_pages}")
    
    # Check hyperlinks
    links = []
    p0 = reader.pages[0]
    if "/Annots" in p0:
        for annot in p0["/Annots"]:
            obj = annot.get_object()
            if "/A" in obj and "/URI" in obj["/A"]:
                links.append(obj["/A"]["/URI"])
    print(f"Active Clickable Hyperlinks Verified ({len(links)}): {links}")
    
    if pages == 1 and pdf_pages == 1 and len(links) >= 5:
        print("\n>>> AUDIT PASSED: PERFECT 1-PAGE CV WITH ALL 5 HYPERLINKS VERIFIED! <<<")
    else:
        print("\n>>> WARNING: PAGE COUNT OR LINK VERIFICATION FAILED! <<<")

if __name__ == "__main__":
    main()
