import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
import shutil

doc = docx.Document()

# Set Margins (0.4 inches top/bottom, 0.5 inches left/right)
sections = doc.sections
for s in sections:
    s.top_margin = Inches(0.4)
    s.bottom_margin = Inches(0.4)
    s.left_margin = Inches(0.5)
    s.right_margin = Inches(0.5)

# Color Palette
NAVY = RGBColor(15, 23, 42)      # #0F172A
BLUE = RGBColor(37, 99, 235)     # #2563EB
SLATE = RGBColor(71, 85, 105)    # #475569
DARK = RGBColor(30, 41, 59)      # #1E293B

def format_run(run, font_name="Arial", size_pt=9.5, bold=False, italic=False, color=DARK):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color

def add_header(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("ADITYA MEHRA")
    format_run(r, font_name="Arial", size_pt=18, bold=True, color=NAVY)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2)
    r_sub = p_sub.add_run("BBA GRADUATE (BATCH OF 2026) // ENTRY-LEVEL OPERATIONS NEW ASSOCIATE")
    format_run(r_sub, font_name="Arial", size_pt=9.5, bold=True, color=BLUE)

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(2)
    r_c = p_contact.add_run("Bengaluru, Karnataka, India  |  +91 7003456624  |  adityamehra799@gmail.com  |  linkedin.com/in/aditya-mehra")
    format_run(r_c, font_name="Arial", size_pt=8.5, color=SLATE)

    p_tag = doc.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tag.paragraph_format.space_before = Pt(0)
    p_tag.paragraph_format.space_after = Pt(6)
    r_t = p_tag.add_run("Fresher (0 Years Experience)  •  Immediate Joiner (0 Days Notice)  •  100% Rotational & Night Shifts  •  Bengaluru Local")
    format_run(r_t, font_name="Arial", size_pt=8, bold=True, color=RGBColor(4, 120, 87))

def add_section_title(doc, title_text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title_text.upper())
    format_run(r, font_name="Arial", size_pt=9.5, bold=True, color=NAVY)

    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="12" w:space="1" w:color="0F172A"/>'
                     r'</w:pBdr>')
    pPr.append(pBdr)

# Build Header
add_header(doc)

# 1. EDUCATION (PROMINENT FOR FRESHERS)
add_section_title(doc, "Education")
p_edu = doc.add_paragraph()
p_edu.paragraph_format.space_before = Pt(2)
p_edu.paragraph_format.space_after = Pt(1)
p_edu.paragraph_format.keep_with_next = True
r_e1 = p_edu.add_run("Bachelor of Business Administration (BBA) — International Business\n")
format_run(r_e1, font_name="Arial", size_pt=9, bold=True, color=NAVY)
r_e2 = p_edu.add_run("Dayananda Sagar University (DSU), Bengaluru, Karnataka    (2023 – 2026 | Fresh Graduate / Immediate Joining)")
format_run(r_e2, font_name="Arial", size_pt=8.5, bold=True, color=BLUE)

p_ed_sub = doc.add_paragraph()
p_ed_sub.paragraph_format.space_before = Pt(0)
p_ed_sub.paragraph_format.space_after = Pt(2)
r_es = p_ed_sub.add_run("Relevant Coursework: Global Logistics & Supply Chain Management, Incoterms 2020, International Trade Finance & Documentation (Letters of Credit, Shipping Bills), Commercial Contract Law, Business Statistics & Decision Modeling.")
format_run(r_es, font_name="Arial", size_pt=8, italic=True, color=SLATE)

# 2. PROFESSIONAL SUMMARY
add_section_title(doc, "Professional Summary")
p_sum = doc.add_paragraph()
p_sum.paragraph_format.space_before = Pt(0)
p_sum.paragraph_format.space_after = Pt(3)
p_sum.paragraph_format.line_spacing = 1.15
p_sum.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r_sum = p_sum.add_run(
    "Goal-driven 2026 BBA graduate in International Business from Dayananda Sagar University, Bengaluru, seeking an entry-level "
    "Level 13 New Associate role at Accenture Operations. Academic specialization in global supply chains, Procure-to-Pay (P2P) cycles, "
    "purchase order (PO) processing, and operational compliance. Backed by verifiable corporate internship experience in commercial operations "
    "(Pencil Mark), AI data validation (Instawork), and high-volume event logistics (Aero India 2025 with 0% inventory shrinkage). "
    "Proficient in MS Excel (VLOOKUP, Pivot Tables, three-way reconciliation) and automated workflows. Immediate joiner with 100% readiness "
    "for 24/7 rotational and night shifts in Bengaluru."
)
format_run(r_sum, font_name="Arial", size_pt=8.5, color=DARK)

# 3. CORE COMPETENCIES & FUNCTIONAL SKILLS
add_section_title(doc, "Core Competencies & Functional Skills")

skills = [
    ("Procurement & P2P Operations:", "PR-to-PO Processing, Three-Way Matching (PO/GRN/Invoice), Vendor Catalog Operations, Supplier SLA Auditing, Purchase Order Administration, Payment Dispute Reconciliation."),
    ("Supply Chain & Trade:", "Incoterms 2020 (EXW, FOB, CIF, DDP), International Shipping Documentation (Bill of Lading, Commercial Invoices), Dual-Custody Inventory Governance, Stock Reconciliation."),
    ("Commercial & Order Ops:", "B2B Order Management (100% Non-Sales), Commercial Proposal Formulation, Contract Milestones, CRM Data Hygiene, Back-Office Query De-escalation."),
    ("Tools & Data Analysis:", "Advanced MS Excel (VLOOKUP, INDEX/MATCH, Pivot Tables, Conditional Formatting), Google Workspace, ERP Workflows, AI Data Quality Auditing (99.4% Accuracy).")
]

for label, desc in skills:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.1
    r_lbl = p.add_run(label + " ")
    format_run(r_lbl, font_name="Arial", size_pt=8.5, bold=True, color=NAVY)
    r_desc = p.add_run(desc)
    format_run(r_desc, font_name="Arial", size_pt=8.5, color=DARK)

# 4. INTERNSHIPS & PRACTICAL FIELD EXPOSURE
add_section_title(doc, "Internships & Practical Field Exposure")

experiences = [
    {
        "role": "Commercial Operations Intern",
        "company": "Pencil Mark Interior Solutions LLP",
        "date_loc": "July 2024 – Aug 2025 | Bengaluru, India",
        "bullets": [
            "Administered supplier vendor matrices, evaluated competing subcontractor quotations, and prepared comparative pricing sheets in MS Excel.",
            "Formulated commercial proposals and scope-of-work documentation, directly supporting the conversion of INR 1.5L+ in secured B2B corporate contracts.",
            "Tracked purchase orders against delivery milestones and site sign-offs, maintaining 100% vendor SLA compliance and zero billing discrepancies."
        ]
    },
    {
        "role": "Operations & AI Data Verification Intern",
        "company": "Instawork Services India Pvt. Ltd.",
        "date_loc": "2025 | Bengaluru, India",
        "bullets": [
            "Executed high-volume multimodal data extraction and operational auditing adhering to strict standardized operating procedures (SOPs).",
            "Maintained a verified 99.4% data accuracy rating across complex operational queues, filtering non-compliant assets before database ingestion.",
            "Collaborated with team supervisors to identify workflow bottlenecks, reducing review turnaround cycle time by 12%."
        ]
    },
    {
        "role": "Exhibition Operations & Logistics Lead",
        "company": "Salt in My Cocoa (Aero India 2025)",
        "date_loc": "Feb 2025 | Yelahanka Air Force Station, Bengaluru",
        "bullets": [
            "Directed on-ground stall setup, inventory replenishment, and security clearances across 5 days for 100,000+ public and defense visitors.",
            "Established a strict dual-custody physical inventory protocol, achieving 0% stock shrinkage and zero stock-outs during peak hours.",
            "Reconciled high-frequency digital UPI transactions daily against physical inventory ledgers with 100% financial balancing."
        ]
    },
    {
        "role": "Community Operations Intern",
        "company": "Non-Profit Social Initiative",
        "date_loc": "2024 | Bengaluru, India",
        "bullets": [
            "Coordinated volunteer schedules, attendee registrations, and field logistics for social awareness drives across Bengaluru urban districts.",
            "Managed participant databases in Google Sheets and resolved on-ground operational dependencies under strict time constraints."
        ]
    }
]

for exp in experiences:
    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_before = Pt(3)
    p_head.paragraph_format.space_after = Pt(1)
    p_head.paragraph_format.keep_with_next = True
    
    r_role = p_head.add_run(exp["role"] + " — ")
    format_run(r_role, font_name="Arial", size_pt=9, bold=True, color=NAVY)
    r_comp = p_head.add_run(exp["company"])
    format_run(r_comp, font_name="Arial", size_pt=9, bold=True, color=BLUE)
    r_tab = p_head.add_run(f"    ({exp['date_loc']})")
    format_run(r_tab, font_name="Arial", size_pt=8, italic=True, color=SLATE)

    for b in exp["bullets"]:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(1)
        p_b.paragraph_format.line_spacing = 1.1
        r_b = p_b.add_run(b)
        format_run(r_b, font_name="Arial", size_pt=8.5, color=DARK)

# 5. TECHNICAL PROJECTS & OPERATIONAL AUTOMATION
add_section_title(doc, "Technical Projects & Operational Automation")
p_proj = doc.add_paragraph()
p_proj.paragraph_format.space_before = Pt(2)
p_proj.paragraph_format.space_after = Pt(1)
p_proj.paragraph_format.keep_with_next = True
r_pr1 = p_proj.add_run("Autonomous Operations & Recruitment Pipeline Engine — ")
format_run(r_pr1, font_name="Arial", size_pt=9, bold=True, color=NAVY)
r_pr2 = p_proj.add_run("Student Automation Project (Python, SQLite, 2026)")
format_run(r_pr2, font_name="Arial", size_pt=8.5, italic=True, color=BLUE)

p_pb1 = doc.add_paragraph(style='List Bullet')
p_pb1.paragraph_format.space_before = Pt(0)
p_pb1.paragraph_format.space_after = Pt(1)
r_pb1 = p_pb1.add_run("Architected an automated operational dispatcher in Python parsing 430+ enterprise job requisitions with 16/16 unit tests passing green.")
format_run(r_pb1, font_name="Arial", size_pt=8.5, color=DARK)

p_pb2 = doc.add_paragraph(style='List Bullet')
p_pb2.paragraph_format.space_before = Pt(0)
p_pb2.paragraph_format.space_after = Pt(1)
r_pb2 = p_pb2.add_run("Built interactive operational consoles eliminating repetitive manual tracking and automating requisition-to-candidate alignment.")
format_run(r_pb2, font_name="Arial", size_pt=8.5, color=DARK)

# 6. CERTIFICATIONS & LANGUAGES
add_section_title(doc, "Certifications & Languages")
p_cert = doc.add_paragraph()
p_cert.paragraph_format.space_before = Pt(2)
p_cert.paragraph_format.space_after = Pt(0)
r_c1 = p_cert.add_run("Certifications: ")
format_run(r_c1, font_name="Arial", size_pt=8.5, bold=True, color=NAVY)
r_c2 = p_cert.add_run("Digital Marketing Fundamentals (Google)  •  Service Marketing (NPTEL)  •  Enterprise Operational Excellence\n")
format_run(r_c2, font_name="Arial", size_pt=8.5, color=DARK)
r_l1 = p_cert.add_run("Languages: ")
format_run(r_l1, font_name="Arial", size_pt=8.5, bold=True, color=NAVY)
r_l2 = p_cert.add_run("English (Fluent Professional Proficiency), Hindi (Native / Bilingual)")
format_run(r_l2, font_name="Arial", size_pt=8.5, color=DARK)

# Save Docx
docx_path = r"e:\anti\resumes\Aditya_Mehra_Accenture_Operations_CV.docx"
root_docx = r"e:\anti\Aditya_Mehra_Accenture_Operations_CV.docx"
doc.save(docx_path)
shutil.copyfile(docx_path, root_docx)
print(f"Successfully generated Fresher DOCX CV at: {docx_path} and {root_docx}", flush=True)
