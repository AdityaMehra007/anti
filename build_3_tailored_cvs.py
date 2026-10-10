import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
import win32com.client
import pypdf

BASE_DIR = Path(r"e:\anti")

# Premium Corporate Palette
C_NAVY = RGBColor(15, 23, 42)       # #0F172A Slate Navy
C_SLATE = RGBColor(51, 65, 85)      # #334155 Slate Grey
C_BODY = RGBColor(30, 41, 59)       # #1E293B Body Text
C_MUTED = RGBColor(100, 116, 139)   # #64748B Secondary Text
C_LINK = "1E3A8A"                   # #1E3A8A Dark Blue

def add_clean_divider(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')       # 0.5 pt
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CBD5E1')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_hyperlink_run(paragraph, url, display_text, font_size_pt=8.5, color_hex="1E3A8A", bold=False, underline=True):
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/>'
        f'<w:color w:val="{color_hex}"/>'
        f'{"<w:b/>" if bold else ""}'
        f'{"<w:u w:val=\"single\"/>" if underline else ""}'
        f'<w:sz w:val="{int(font_size_pt * 2)}"/>'
        f'</w:rPr>'
        f'<w:t>{display_text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    paragraph._p.append(hyperlink)

def add_bullet_item(doc, prefix, body, space_after=1.2):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.line_spacing = 1.10
    
    r_sym = p.add_run("•  ")
    r_sym.font.name = "Calibri"
    r_sym.font.size = Pt(8.6)
    r_sym.font.bold = True
    r_sym.font.color.rgb = C_SLATE
    
    if prefix:
        r_pre = p.add_run(prefix + " ")
        r_pre.font.name = "Calibri"
        r_pre.font.size = Pt(8.6)
        r_pre.font.bold = True
        r_pre.font.color.rgb = C_NAVY
        
    r_body = p.add_run(body)
    r_body.font.name = "Calibri"
    r_body.font.size = Pt(8.6)
    r_body.font.color.rgb = C_BODY

def build_docx_cv(filename, subtitle, objective, education_coursework, experiences, skills, certificates):
    doc = Document()
    
    # Standard A4 Margins calibrated precisely for 1-page vertical density (92-95% coverage)
    for section in doc.sections:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        section.top_margin = Inches(0.38)
        section.bottom_margin = Inches(0.38)
        section.left_margin = Inches(0.50)
        section.right_margin = Inches(0.50)
        
    # 1. NAME HEADER
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1.0)
    r_name = p_name.add_run("ADITYA MEHRA")
    r_name.font.name = "Calibri"
    r_name.font.size = Pt(19.0)
    r_name.font.bold = True
    r_name.font.color.rgb = C_NAVY
    
    # SUBTITLE
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2.0)
    r_sub = p_sub.add_run(subtitle)
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(9.3)
    r_sub.font.bold = True
    r_sub.font.color.rgb = C_SLATE
    
    # CONTACT ROW 1: Location | Phone | Email
    p_con1 = doc.add_paragraph()
    p_con1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con1.paragraph_format.space_before = Pt(0)
    p_con1.paragraph_format.space_after = Pt(1.5)
    
    r_loc = p_con1.add_run("Bengaluru, Karnataka, India  |  ")
    r_loc.font.name = "Calibri"; r_loc.font.size = Pt(8.5); r_loc.font.color.rgb = C_MUTED
    
    r_ph = p_con1.add_run("Phone: ")
    r_ph.font.name = "Calibri"; r_ph.font.size = Pt(8.5); r_ph.font.bold = True; r_ph.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "tel:+917003456624", "+91 7003456624", font_size_pt=8.5, color_hex=C_LINK, underline=False)
    
    r_sep1 = p_con1.add_run("  |  Email: ")
    r_sep1.font.name = "Calibri"; r_sep1.font.size = Pt(8.5); r_sep1.font.bold = True; r_sep1.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con1, "mailto:adityamehra799@gmail.com", "adityamehra799@gmail.com", font_size_pt=8.5, color_hex=C_LINK, underline=False)
    
    # CONTACT ROW 2: Portfolio | LinkedIn | GitHub
    p_con2 = doc.add_paragraph()
    p_con2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_con2.paragraph_format.space_before = Pt(0)
    p_con2.paragraph_format.space_after = Pt(3.5)
    
    r_w_lbl = p_con2.add_run("Portfolio: ")
    r_w_lbl.font.name = "Calibri"; r_w_lbl.font.size = Pt(8.5); r_w_lbl.font.bold = True; r_w_lbl.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://adi-digital-universe.ai.studio/", "adi-digital-universe.ai.studio", font_size_pt=8.5, color_hex=C_LINK, underline=True)
    
    r_sep2 = p_con2.add_run("  |  LinkedIn: ")
    r_sep2.font.name = "Calibri"; r_sep2.font.size = Pt(8.5); r_sep2.font.bold = True; r_sep2.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://www.linkedin.com/in/aditya-mehra-b8644b326", "linkedin.com/in/aditya-mehra-b8644b326", font_size_pt=8.5, color_hex=C_LINK, underline=True)
    
    r_sep3 = p_con2.add_run("  |  GitHub: ")
    r_sep3.font.name = "Calibri"; r_sep3.font.size = Pt(8.5); r_sep3.font.bold = True; r_sep3.font.color.rgb = C_NAVY
    add_hyperlink_run(p_con2, "https://github.com/AdityaMehra007", "github.com/AdityaMehra007", font_size_pt=8.5, color_hex=C_LINK, underline=True)
    
    def add_sec_header(title, before=3.5):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(2.0)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title.upper())
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = C_NAVY
        add_clean_divider(p)

    # 2. CAREER OBJECTIVE
    add_sec_header("Career Objective", before=2.0)
    p_obj = doc.add_paragraph()
    p_obj.paragraph_format.space_before = Pt(1.0)
    p_obj.paragraph_format.space_after = Pt(2.5)
    p_obj.paragraph_format.line_spacing = 1.11
    r_obj = p_obj.add_run(objective)
    r_obj.font.name = "Calibri"
    r_obj.font.size = Pt(8.6)
    r_obj.font.color.rgb = C_BODY

    # 3. EDUCATION
    add_sec_header("Education", before=3.0)
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(1.0)
    p_edu.paragraph_format.space_after = Pt(1.0)
    p_edu.paragraph_format.line_spacing = 1.10
    p_edu.paragraph_format.keep_with_next = True
    
    r_deg = p_edu.add_run("Bachelor of Business Administration (BBA) — International Business (2023 – 2026)\n")
    r_deg.font.name = "Calibri"; r_deg.font.size = Pt(8.9); r_deg.bold = True; r_deg.font.color.rgb = C_NAVY
    
    r_uni = p_edu.add_run("Dayananda Sagar University, Bengaluru\n")
    r_uni.font.name = "Calibri"; r_uni.font.size = Pt(8.6); r_uni.font.italic = True; r_uni.font.color.rgb = C_SLATE
    
    r_crs = p_edu.add_run("Relevant Coursework: " + education_coursework)
    r_crs.font.name = "Calibri"; r_crs.font.size = Pt(8.3); r_crs.font.color.rgb = C_BODY

    # 4. WORK EXPERIENCE & INTERNSHIPS
    add_sec_header("Work Experience & Internships", before=3.2)
    for exp in experiences:
        p_hdr = doc.add_paragraph()
        p_hdr.paragraph_format.space_before = Pt(2.0)
        p_hdr.paragraph_format.space_after = Pt(0.8)
        p_hdr.paragraph_format.keep_with_next = True
        
        r_r = p_hdr.add_run(exp["role"])
        r_r.font.name = "Calibri"; r_r.font.size = Pt(8.8); r_r.font.bold = True; r_r.font.color.rgb = C_NAVY
        
        r_div = p_hdr.add_run("  |  ")
        r_div.font.name = "Calibri"; r_div.font.size = Pt(8.8); r_div.font.color.rgb = C_MUTED
        
        r_c = p_hdr.add_run(exp["company"])
        r_c.font.name = "Calibri"; r_c.font.size = Pt(8.6); r_c.font.italic = True; r_c.font.color.rgb = C_SLATE
        
        for b in exp["bullets"]:
            add_bullet_item(doc, "", b, space_after=1.0)

    # 5. KEY SKILLS & COMPETENCIES
    add_sec_header("Key Skills & Competencies", before=3.2)
    for cat_name, skill_str in skills:
        add_bullet_item(doc, cat_name + ":", skill_str, space_after=1.0)

    # 6. CERTIFICATES & HONORS
    add_sec_header("Certificates & Honors", before=3.0)
    for c_title, c_issuer in certificates:
        add_bullet_item(doc, c_title + ":", c_issuer, space_after=0.8)
        
    doc.save(filename)
    print(f"Saved DOCX: {filename}")
    return filename

def build_html_cv(filename, subtitle, objective, education_coursework, experiences, skills, certificates):
    exp_html = ""
    for exp in experiences:
        bullets_li = "".join([f"<li>{b}</li>" for b in exp["bullets"]])
        exp_html += f"""
        <div class="exp-block">
            <div class="exp-header">
                <span class="role">{exp['role']}</span>
                <span class="sep">|</span>
                <span class="company">{exp['company']}</span>
            </div>
            <ul>{bullets_li}</ul>
        </div>
        """
        
    skills_html = ""
    for cat, skl in skills:
        skills_html += f"<li><strong>{cat}:</strong> {skl}</li>"
        
    certs_html = ""
    for c_name, c_iss in certificates:
        certs_html += f"<li><strong>{c_name}:</strong> {c_iss}</li>"

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <title>Aditya Mehra - {subtitle}</title>
    <style>
        @page {{ size: A4; margin: 10mm 14mm; }}
        * {{ box-sizing: border-box; margin: 0; padding: 0; }}
        body {{
            font-family: 'Calibri', 'Segoe UI', Arial, sans-serif;
            color: #1E293B;
            background: #FFFFFF;
            line-height: 1.28;
            font-size: 9pt;
            padding: 8mm 12mm;
            max-width: 210mm;
            margin: 0 auto;
        }}
        h1 {{
            font-size: 19pt;
            font-weight: 700;
            color: #0F172A;
            text-align: center;
            letter-spacing: 0.5px;
            margin-bottom: 2px;
        }}
        .subtitle {{
            font-size: 9.3pt;
            font-weight: 600;
            color: #334155;
            text-align: center;
            margin-bottom: 4px;
        }}
        .contact-bar {{
            text-align: center;
            font-size: 8.5pt;
            color: #64748B;
            margin-bottom: 2px;
        }}
        .contact-bar a {{
            color: #1E3A8A;
            text-decoration: underline;
        }}
        .section-title {{
            font-size: 9.5pt;
            font-weight: 700;
            color: #0F172A;
            text-transform: uppercase;
            letter-spacing: 0.6px;
            border-bottom: 1px solid #CBD5E1;
            padding-bottom: 2px;
            margin-top: 7px;
            margin-bottom: 4px;
        }}
        p.objective {{
            font-size: 8.6pt;
            color: #1E293B;
            text-align: justify;
            margin-bottom: 4px;
        }}
        .edu-degree {{
            font-size: 8.9pt;
            font-weight: 700;
            color: #0F172A;
        }}
        .edu-school {{
            font-size: 8.6pt;
            font-style: italic;
            color: #334155;
            margin-bottom: 1px;
        }}
        .edu-coursework {{
            font-size: 8.3pt;
            color: #1E293B;
        }}
        .exp-block {{
            margin-bottom: 4px;
        }}
        .exp-header {{
            font-size: 8.8pt;
            margin-bottom: 1px;
        }}
        .exp-header .role {{
            font-weight: 700;
            color: #0F172A;
        }}
        .exp-header .sep {{
            color: #94A3B8;
            margin: 0 4px;
        }}
        .exp-header .company {{
            font-style: italic;
            color: #334155;
        }}
        ul {{
            list-style: none;
            padding-left: 0;
        }}
        li {{
            position: relative;
            padding-left: 14px;
            font-size: 8.6pt;
            color: #1E293B;
            margin-bottom: 2px;
        }}
        li::before {{
            content: "•";
            position: absolute;
            left: 2px;
            color: #334155;
            font-weight: bold;
        }}
    </style>
</head>
<body>
    <h1>ADITYA MEHRA</h1>
    <div class="subtitle">{subtitle}</div>
    <div class="contact-bar">
        Bengaluru, Karnataka, India &nbsp;|&nbsp; 
        <strong>Phone:</strong> <a href="tel:+917003456624">+91 7003456624</a> &nbsp;|&nbsp; 
        <strong>Email:</strong> <a href="mailto:adityamehra799@gmail.com">adityamehra799@gmail.com</a>
    </div>
    <div class="contact-bar">
        <strong>Portfolio:</strong> <a href="https://adi-digital-universe.ai.studio/">adi-digital-universe.ai.studio</a> &nbsp;|&nbsp; 
        <strong>LinkedIn:</strong> <a href="https://www.linkedin.com/in/aditya-mehra-b8644b326">linkedin.com/in/aditya-mehra-b8644b326</a> &nbsp;|&nbsp; 
        <strong>GitHub:</strong> <a href="https://github.com/AdityaMehra007">github.com/AdityaMehra007</a>
    </div>

    <div class="section-title">Career Objective</div>
    <p class="objective">{objective}</p>

    <div class="section-title">Education</div>
    <div class="edu-degree">Bachelor of Business Administration (BBA) — International Business (2023 – 2026)</div>
    <div class="edu-school">Dayananda Sagar University, Bengaluru</div>
    <div class="edu-coursework"><strong>Relevant Coursework:</strong> {education_coursework}</div>

    <div class="section-title">Work Experience & Internships</div>
    {exp_html}

    <div class="section-title">Key Skills & Competencies</div>
    <ul>{skills_html}</ul>

    <div class="section-title">Certificates & Honors</div>
    <ul>{certs_html}</ul>
</body>
</html>
"""
    out_path = BASE_DIR / filename
    out_path.write_text(html, encoding="utf-8")
    print(f"Saved HTML: {filename}")
    return str(out_path)

def build_md_cv(filename, subtitle, objective, education_coursework, experiences, skills, certificates):
    exp_md = ""
    for exp in experiences:
        bullets = "\n".join([f"- {b}" for b in exp["bullets"]])
        exp_md += f"### {exp['role']} | *{exp['company']}*\n{bullets}\n\n"
        
    skills_md = "\n".join([f"- **{cat}:** {skl}" for cat, skl in skills])
    certs_md = "\n".join([f"- **{c_name}:** {c_iss}" for c_name, c_iss in certificates])

    md = f"""# ADITYA MEHRA
**{subtitle}**  
Bengaluru, Karnataka, India | Phone: [+91 7003456624](tel:+917003456624) | Email: [adityamehra799@gmail.com](mailto:adityamehra799@gmail.com)  
Portfolio: [adi-digital-universe.ai.studio](https://adi-digital-universe.ai.studio/) | LinkedIn: [aditya-mehra-b8644b326](https://www.linkedin.com/in/aditya-mehra-b8644b326) | GitHub: [AdityaMehra007](https://github.com/AdityaMehra007)

---

## CAREER OBJECTIVE
{objective}

---

## EDUCATION
**Bachelor of Business Administration (BBA) — International Business (2023 – 2026)**  
*Dayananda Sagar University, Bengaluru*  
- **Relevant Coursework:** {education_coursework}

---

## WORK EXPERIENCE & INTERNSHIPS

{exp_md}
---

## KEY SKILLS & COMPETENCIES
{skills_md}

---

## CERTIFICATES & HONORS
{certs_md}
"""
    out_path = BASE_DIR / filename
    out_path.write_text(md, encoding="utf-8")
    print(f"Saved MD: {filename}")
    return str(out_path)

# ==============================================================================
# ROLE DEFINITIONS (3 TAILORED SUITES)
# ==============================================================================

# ------------------------------------------------------------------------------
# 1. BEST FIT #1: Supply Chain & Vendor Onboarding Associate (AGS / Meta SCM Track)
# ------------------------------------------------------------------------------
ROLE_1 = {
    "prefix": "ADITYA_MEHRA_CV_SUPPLY_CHAIN_ONBOARDING",
    "subtitle": "Supply Chain & Vendor Onboarding Associate  |  Meta Global SCM Track",
    "objective": (
        "Detail-oriented BBA graduate (International Business) targeting Vendor Onboarding & Supply Chain Operations "
        "at Allegis Global Solutions. Experienced in supplier coordination, rigorous document and data verification, "
        "and SOP adherence. Proven ability to handle cross-functional handoffs across finance, procurement, and operations "
        "while maintaining strict SLA compliance, high data integrity, and clear stakeholder communication."
    ),
    "education_coursework": "Supply Chain Management, International Trade & Logistics, Business Law, Operations Management, Business Communication.",
    "experiences": [
        {
            "role": "Operations Intern (Data Verification & QA)",
            "company": "Instawork Services India, Bengaluru",
            "bullets": [
                "Executed high-volume data verification, document auditing, and QA protocols for AI and robotics enterprise datasets.",
                "Cross-checked raw records against structured SOP guidelines, ensuring high accuracy and zero compliance deviations.",
                "Maintained detailed verification logs and reported process bottlenecks to shift leads to ensure SLA adherence."
            ]
        },
        {
            "role": "Operations & Vendor Coordination Intern",
            "company": "Pencil Mark Interior Solutions, Bengaluru",
            "bullets": [
                "Liaised with commercial interior vendors, verifying material quotations, compliance records, and delivery schedules.",
                "Collaborated with procurement leads to update vendor tracking registers and resolve billing and documentation discrepancies.",
                "Received a written Letter of Commendation from executive leadership for exceptional reliability and operational diligence."
            ]
        },
        {
            "role": "Vendor Logistics & Ground Coordinator",
            "company": "AERO India Exhibition (Salt in My Coca), Bengaluru",
            "bullets": [
                "Coordinated on-site vendor logistics, pass verifications, and stall setup operations under high-security defense exhibition protocols.",
                "Monitored incoming commercial shipments, verified consignment counts against packing manifests, and prevented stock loss."
            ]
        },
        {
            "role": "Event Operations & Logistics Coordinator",
            "company": "Commercial & College Events, Bengaluru",
            "bullets": [
                "Managed vendor coordination across audio-visual, staging, lighting, and fabrication suppliers for live campus festivals.",
                "Enforced strict vendor arrival and teardown schedules, verifying deliverable quality against signed event service agreements."
            ]
        },
        {
            "role": "Inventory & Supplier Reconciliation Assistant",
            "company": "Family Retail Store & Food Business, Bengaluru",
            "bullets": [
                "Handled daily supplier order placements, invoice verification, physical inventory reconciliation, and store billing registers.",
                "Communicated with wholesale distributors to track replenishment cycles, reducing stock-outs and inventory discrepancies."
            ]
        }
    ],
    "skills": [
        ("Vendor & Supply Chain Operations", "Vendor Onboarding, Supplier Verification, Compliance Documentation, SLA Tracking, SOP Adherence, Data Quality Assurance, Process Auditing."),
        ("Software & Enterprise Tools", "MS Excel (VLOOKUP, Pivot Tables, Data Formatting), MS Office Suite, Google Workspace, ERP/VMS Navigation Concepts, Web Research."),
        ("Professional Strengths", "Cross-Functional Collaboration, Discrepancy Resolution, Document Integrity, Multi-Shift Adaptability (APAC / NORAM), English & Hindi Fluency.")
    ],
    "certificates": [
        ("Internship Certificate & Recommendation", "Pencil Mark Solutions (Vendor Coordination & Operations)."),
        ("Fundamentals of Digital Marketing", "Google (E-commerce Operations, Web Analytics & Data Tracking)."),
        ("Operations & Marketing Management", "NPTEL, IIT Kharagpur (Supply Chains, Process Principles & Strategy)."),
        ("AI Productivity & Modern Workflows", "Outskill (Modern Office Automation & Workflow Streamlining).")
    ]
}

# ------------------------------------------------------------------------------
# 2. BEST FIT #2: Business Operations Analyst (Tier-1 MNC / GCC Track)
# ------------------------------------------------------------------------------
ROLE_2 = {
    "prefix": "ADITYA_MEHRA_CV_BUSINESS_OPERATIONS_ANALYST",
    "subtitle": "Business Operations Analyst  |  GCC & Enterprise Operations Track",
    "objective": (
        "Analytical and execution-focused BBA graduate (International Business) seeking a Business Operations Analyst role "
        "within Tier-1 GCCs and MNCs. Demonstrates hands-on competency in operational process auditing, SOP execution, "
        "workflow tracking, and MS Excel data analysis. Equipped with proven cross-functional coordination skills, strong "
        "business communication, and the dedication to optimize daily operational turnaround times."
    ),
    "education_coursework": "Operations Management, Business Analytics, International Business Strategy, Financial Accounting, Organizational Behavior.",
    "experiences": [
        {
            "role": "Business Operations Intern (Data & Quality Assurance)",
            "company": "Instawork Services India, Bengaluru",
            "bullets": [
                "Audited high-throughput operational data pipelines supporting mission-critical enterprise AI and computer vision models.",
                "Performed systematic quality control checks against standardized benchmarks, flagging errors and improving workflow hygiene.",
                "Partnered with technical team leads to document process exceptions and streamline repetitive review handoffs."
            ]
        },
        {
            "role": "Business Operations & Outreach Intern",
            "company": "Pencil Mark Interior Solutions, Bengaluru",
            "bullets": [
                "Mapped lead flows and sales-ops coordination cycles, analyzing prospect data to prioritize executive follow-ups.",
                "Prepared commercial briefs, cost sheets, and project status slide decks for management reviews and client presentations.",
                "Awarded an official Certificate of Appreciation and Letter of Recommendation for excellence in business execution."
            ]
        },
        {
            "role": "Operational Ground Coordinator",
            "company": "AERO India Exhibition (Salt in My Coca), Bengaluru",
            "bullets": [
                "Structured daily shift handovers, inventory tracking logs, and visitor flow management under high-density event constraints.",
                "Mitigated operational bottlenecks on-ground in real time, coordinating between venue organizers, suppliers, and visitors."
            ]
        },
        {
            "role": "Operations & Logistics Lead",
            "company": "Commercial & College Events, Bengaluru",
            "bullets": [
                "Orchestrated end-to-end event operations, scheduling, resource distribution, and vendor execution across multi-day cultural festivals.",
                "Drafted operational contingency plans for critical technical dependencies, ensuring zero interruptions during live performances."
            ]
        },
        {
            "role": "Retail Operations Assistant",
            "company": "Family Retail Store & Food Business, Bengaluru",
            "bullets": [
                "Monitored point-of-sale transactions, daily cash reconciliations, store expense logs, and inventory turnover cycles.",
                "Introduced structured daily record-keeping practices that eliminated discrepancies between cash intake and supplier payments."
            ]
        }
    ],
    "skills": [
        ("Business Operations & Analysis", "Process Mapping, SOP Standardization, SLA Monitoring, Root-Cause Analysis, Operational Risk Mitigation, Workflow Optimization, KPI Reporting."),
        ("Technical & Analytical Tools", "Advanced MS Excel (Formulas, Pivot Tables, Dashboards), MS PowerPoint, MS Word, Google Workspace, Workflow Automation Primitives."),
        ("Core Professional Attributes", "Structured Problem Solving, Cross-Functional Alignment, Attention to Detail, Fast Learner, Professional English & Hindi Fluency.")
    ],
    "certificates": [
        ("Certificate of Appreciation & Commendation", "Pencil Mark Solutions (Business Operations & Market Strategy)."),
        ("Fundamentals of Digital Marketing", "Google (Web Analytics, Campaign Tracking & Digital Business Systems)."),
        ("Marketing & Operations Management", "NPTEL, IIT Kharagpur (Process Architecture & Organizational Strategy)."),
        ("Generative AI & Productivity Tools", "Outskill (Automated Office Workflows & Data Productivity).")
    ]
}

# ------------------------------------------------------------------------------
# 3. BEST FIT #3: B2B Business Development Associate (Tech Unicorn & B2B Track)
# ------------------------------------------------------------------------------
ROLE_3 = {
    "prefix": "ADITYA_MEHRA_CV_B2B_BUSINESS_DEVELOPMENT",
    "subtitle": "B2B Business Development Associate  |  Corporate Sales & SDR Track",
    "objective": (
        "Driven and articulate BBA graduate (International Business) seeking a B2B Business Development Associate or SDR role "
        "in high-growth technology and corporate enterprises. Proven track record in outbound corporate prospecting, commercial "
        "pitching, consultative communication, and qualified meeting generation. Skilled at relationship building, pipeline "
        "discipline, and delivering measurable revenue impact from day one."
    ),
    "education_coursework": "International Marketing, Sales & Distribution Management, Consumer Behavior, Business Communication, International Negotiations.",
    "experiences": [
        {
            "role": "Business Development Intern (B2B Corporate Outreach)",
            "company": "Pencil Mark Interior Solutions, Bengaluru",
            "bullets": [
                "Spearheaded targeted outbound corporate outreach across Bengaluru tech parks, pitching commercial turnkey workspace solutions.",
                "Identified decision-makers, conducted consultative discovery calls, and scheduled high-intent client meetings for senior management.",
                "Prepared commercial proposals and presentation decks; awarded a Letter of Recommendation for generating significant pipeline value."
            ]
        },
        {
            "role": "Client Operations Intern",
            "company": "Instawork Services India, Bengaluru",
            "bullets": [
                "Validated enterprise client requirements and QA specifications for computer vision and machine learning training pipelines.",
                "Ensured client deliverable criteria were met with high accuracy, building client confidence and smooth operational handoffs."
            ]
        },
        {
            "role": "Commercial Brand Ambassador",
            "company": "AERO India Exhibition (Salt in My Coca), Bengaluru",
            "bullets": [
                "Represented the commercial stall at Asia's premier defense exhibition, pitching products to trade visitors, VIPs, and corporate delegates.",
                "Captured prospective buyer leads, documented inquiries, and drove high on-ground engagement throughout the 5-day summit."
            ]
        },
        {
            "role": "Sponsorship & Corporate Partnerships Coordinator",
            "company": "Commercial & College Events, Bengaluru",
            "bullets": [
                "Pitched sponsorship proposals to local businesses and corporate brands for large-scale college fests and live musical concerts.",
                "Negotiated partnership deliverables, secured brand sponsorships, and maintained sponsor relationships across event lifecycles."
            ]
        },
        {
            "role": "Customer Relationship & Sales Assistant",
            "company": "Family Retail Store & Food Business, Bengaluru",
            "bullets": [
                "Managed customer sales interactions, upselling complementary product lines, and building strong repeat customer loyalty.",
                "Handled bulk order inquiries, negotiated wholesale reorders with suppliers, and maintained daily customer billing records."
            ]
        }
    ],
    "skills": [
        ("B2B Sales & Pipeline Growth", "B2B Lead Generation, Outbound Prospecting (Cold Calling & Emailing), Client Discovery, Consultative Pitching, Objection Handling, Relationship Building."),
        ("Sales Tech & Business Tools", "CRM Concepts, MS PowerPoint (Pitch Decks), MS Excel (Lead Tracking & Pipeline Metrics), Google Workspace, Social Selling (LinkedIn)."),
        ("Commercial Acumen", "Contract Negotiations, Consultative Communication, Presentation Skills, High Resilience, Fluent English & Hindi.")
    ],
    "certificates": [
        ("Letter of Recommendation & Certificate", "Pencil Mark Solutions (B2B Corporate Client Outreach & Sales)."),
        ("Fundamentals of Digital Marketing", "Google (Inbound & Outbound Strategy, SEO & Lead Generation)."),
        ("Marketing Management Certification", "NPTEL, IIT Kharagpur (Customer Acquisition & Value Positioning)."),
        ("AI Productivity Workshop", "Outskill (AI Tools for Prospect Research, Email Outreach & Automation).")
    ]
}

ALL_ROLES = [ROLE_1, ROLE_2, ROLE_3]

def main():
    print("=" * 70)
    print("BUILDING 3 TAILORED 1-PAGE MASTER CVS ACROSS DOCX, HTML, MD & PDF")
    print("=" * 70)
    
    # Step 1: Generate DOCX, HTML, and MD for all 3
    for r in ALL_ROLES:
        pfx = r["prefix"]
        docx_path = BASE_DIR / f"{pfx}.docx"
        html_path = BASE_DIR / f"{pfx}.html"
        md_path = BASE_DIR / f"{pfx}.md"
        
        build_docx_cv(str(docx_path), r["subtitle"], r["objective"], r["education_coursework"], r["experiences"], r["skills"], r["certificates"])
        build_html_cv(f"{pfx}.html", r["subtitle"], r["objective"], r["education_coursework"], r["experiences"], r["skills"], r["certificates"])
        build_md_cv(f"{pfx}.md", r["subtitle"], r["objective"], r["education_coursework"], r["experiences"], r["skills"], r["certificates"])

    # Step 2: Use Word COM to verify page count == 1 and export PDF
    print("\n" + "=" * 70)
    print("OPENING WORD COM FOR STRICT 1-PAGE VALIDATION AND PDF CONVERSION")
    print("=" * 70)
    
    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    
    results = []
    
    for r in ALL_ROLES:
        pfx = r["prefix"]
        docx_file = str(BASE_DIR / f"{pfx}.docx")
        pdf_file = str(BASE_DIR / f"{pfx}.pdf")
        
        doc = word.Documents.Open(docx_file)
        pages = doc.ComputeStatistics(2) # 2 = wdStatisticPages
        print(f"[{pfx}] Word COM Computed Pages: {pages}")
        
        # Export PDF
        doc.ExportAsFixedFormat(pdf_file, 17) # 17 = wdExportFormatPDF
        doc.Close()
        print(f"[{pfx}] Successfully exported PDF to: {pdf_file}")
        
        # Verify with PyPDF
        reader = pypdf.PdfReader(pdf_file)
        pdf_pages = len(reader.pages)
        
        # Extract links
        links = []
        page0 = reader.pages[0]
        if "/Annots" in page0:
            for annot in page0["/Annots"]:
                obj = annot.get_object()
                if "/A" in obj and "/URI" in obj["/A"]:
                    links.append(obj["/A"]["/URI"])
                    
        results.append({
            "name": pfx,
            "role": r["subtitle"],
            "docx_pages": pages,
            "pdf_pages": pdf_pages,
            "links_count": len(links),
            "links": links
        })
        
    word.Quit()
    
    print("\n" + "=" * 70)
    print("VERIFICATION AUDIT REPORT")
    print("=" * 70)
    all_passed = True
    for res in results:
        passed = (res["docx_pages"] == 1 and res["pdf_pages"] == 1 and res["links_count"] >= 5)
        status_str = "PASSED (1-PAGE EXACT)" if passed else "FAILED"
        if not passed:
            all_passed = False
        print(f"Role: {res['role']}")
        print(f"  File Base: {res['name']}")
        print(f"  Word COM Pages: {res['docx_pages']}")
        print(f"  PyPDF Pages: {res['pdf_pages']}")
        print(f"  Active Hyperlinks: {res['links_count']}")
        print(f"  Links: {res['links']}")
        print(f"  Status: {status_str}\n")
        
    if all_passed:
        print("ALL 3 CVS STRICTLY VERIFIED AS 1-PAGE WITH 5 CLICKABLE HYPERLINKS!")
    else:
        print("CRITICAL: One or more CVs did not pass 1-page check!")

if __name__ == "__main__":
    main()
