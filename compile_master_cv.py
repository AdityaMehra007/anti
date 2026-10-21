import os
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
import shutil

# Color Palette for HTML & Docx
NAVY = RGBColor(15, 23, 42)      # #0F172A
BLUE = RGBColor(37, 99, 235)     # #2563EB
SLATE = RGBColor(71, 85, 105)    # #475569
DARK = RGBColor(30, 41, 59)      # #1E293B

# -------------------------------------------------------------
# 1. GENERATE MASTER UNIVERSAL FRESHER CV (HTML)
# -------------------------------------------------------------
master_html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Aditya Mehra - Universal Fresher CV (Accenture Operations Level 13)</title>
<style>
  @page {
    size: A4 portrait;
    margin: 7mm 9mm 7mm 9mm;
  }
  * {
    box-sizing: border-box;
    margin: 0;
    padding: 0;
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
    color: #1e293b;
    line-height: 1.28;
  }
  body {
    background: #fff;
    font-size: 9.1pt;
  }
  .container {
    max-width: 100%;
    margin: 0 auto;
  }
  
  /* Header */
  .header {
    text-align: center;
    border-bottom: 2px solid #0f172a;
    padding-bottom: 5px;
    margin-bottom: 6px;
  }
  .name {
    font-size: 19pt;
    font-weight: 800;
    letter-spacing: 0.8px;
    color: #0f172a;
    text-transform: uppercase;
  }
  .subtitle {
    font-size: 9.5pt;
    font-weight: 700;
    color: #2563eb;
    margin-top: 1px;
    letter-spacing: 0.4px;
    text-transform: uppercase;
  }
  .contact-info {
    font-size: 8.3pt;
    color: #475569;
    margin-top: 3px;
    display: flex;
    justify-content: center;
    gap: 12px;
    flex-wrap: wrap;
    font-weight: 500;
  }
  .contact-info a {
    color: #2563eb;
    text-decoration: none;
  }
  .tagline-badge {
    margin-top: 3px;
    font-size: 7.8pt;
    font-weight: 700;
    color: #047857;
    background: #ecfdf5;
    display: inline-block;
    padding: 1px 7px;
    border-radius: 4px;
    border: 1px solid #a7f3d0;
  }

  /* Section Styles */
  .section-title {
    font-size: 9.5pt;
    font-weight: 800;
    color: #0f172a;
    text-transform: uppercase;
    letter-spacing: 0.6px;
    border-bottom: 1.5px solid #cbd5e1;
    padding-bottom: 2px;
    margin-top: 6px;
    margin-bottom: 4px;
  }
  
  /* Summary */
  .summary-text {
    font-size: 8.5pt;
    color: #334155;
    text-align: justify;
    line-height: 1.32;
  }

  /* Education */
  .edu-row {
    display: flex;
    justify-content: space-between;
    align-items: baseline;
  }
  .edu-title {
    font-size: 9.3pt;
    font-weight: 700;
    color: #0f172a;
  }
  .edu-sub {
    font-size: 8.5pt;
    color: #2563eb;
    font-weight: 600;
    margin-top: 1px;
  }
  
  /* Skills Table */
  .skills-table {
    width: 100%;
    border-collapse: collapse;
    font-size: 8.3pt;
    margin-bottom: 1px;
  }
  .skills-table td {
    padding: 1.5px 0;
    vertical-align: top;
  }
  .skill-cat {
    font-weight: 700;
    color: #0f172a;
    width: 24%;
  }
  .skill-items {
    color: #334155;
    width: 76%;
  }

  /* Experience & Project Blocks */
  .item-block {
    margin-bottom: 5px;
  }
  .item-header {
    display: flex;
    justify-content: space-between;
    align-items: baseline;
    margin-bottom: 1px;
  }
  .role-title {
    font-size: 9pt;
    font-weight: 700;
    color: #0f172a;
  }
  .company-name {
    font-size: 8.6pt;
    font-weight: 600;
    color: #2563eb;
  }
  .date-loc {
    font-size: 7.8pt;
    color: #64748b;
    font-weight: 600;
    text-align: right;
  }
  .bullets {
    list-style-type: disc;
    margin-left: 13px;
    margin-top: 1px;
  }
  .bullets li {
    font-size: 8.2pt;
    color: #334155;
    margin-bottom: 1px;
    line-height: 1.28;
    text-align: justify;
  }
  .bullets li strong {
    color: #0f172a;
  }

  @media screen {
    body {
      background: #f1f5f9;
      padding: 25px;
    }
    .container {
      background: #fff;
      max-width: 820px;
      padding: 30px 38px;
      margin: 0 auto;
      box-shadow: 0 4px 20px rgba(0,0,0,0.1);
      border-radius: 6px;
    }
    .print-btn {
      position: fixed;
      bottom: 25px;
      right: 25px;
      background: #2563eb;
      color: #fff;
      padding: 12px 22px;
      border-radius: 8px;
      font-weight: 700;
      border: none;
      cursor: pointer;
      box-shadow: 0 4px 15px rgba(37,99,235,0.4);
      font-size: 14px;
      display: flex;
      align-items: center;
      gap: 8px;
    }
    .print-btn:hover { background: #1d4ed8; }
  }
  @media print {
    .print-btn { display: none; }
    body { background: #fff; padding: 0; }
    .container { box-shadow: none; padding: 0; max-width: 100%; }
  }
</style>
</head>
<body>

<button class="print-btn" onclick="window.print()">
  <svg width="16" height="16" fill="currentColor" viewBox="0 0 16 16"><path d="M2.5 8a.5.5 0 1 0 0-1 .5.5 0 0 0 0 1z"/><path d="M5 1a2 2 0 0 0-2 2v2H2a2 2 0 0 0-2 2v3a2 2 0 0 0 2 2h1v1a2 2 0 0 0 2 2h6a2 2 0 0 0 2-2v-1h1a2 2 0 0 0 2-2V7a2 2 0 0 0-2-2h-1V3a2 2 0 0 0-2-2H5zM4 3a1 1 0 0 1 1-1h6a1 1 0 0 1 1 1v2H4V3zm1 5a2 2 0 0 0-2 2v1H2a1 1 0 0 1-1-1V7a1 1 0 0 1 1-1h12a1 1 0 0 1 1 1v3a1 1 0 0 1-1 1h-1v-1a2 2 0 0 0-2-2H5zm7 2v3a1 1 0 0 1-1 1H5a1 1 0 0 1-1-1v-3a1 1 0 0 1 1-1h6a1 1 0 0 1 1 1z"/></svg>
  Print / Save as PDF
</button>

<div class="container">

  <!-- HEADER -->
  <div class="header">
    <div class="name">ADITYA MEHRA</div>
    <div class="subtitle">BBA Graduate (Batch of 2026) // Entry-Level Operations New Associate</div>
    <div class="contact-info">
      <span>Bengaluru, Karnataka, India</span>
      <span>+91 7003456624</span>
      <span>adityamehra799@gmail.com</span>
      <span><a href="https://www.linkedin.com/in/aditya-mehra" target="_blank">linkedin.com/in/aditya-mehra</a></span>
      <span><a href="https://github.com/AdityaMehra007" target="_blank">github.com/AdityaMehra007</a></span>
    </div>
    <div class="tagline-badge">
      Fresher (0–1 Yrs) • Immediate Joiner (0 Days Notice) • 100% Rotational & Night Shifts • Level 13 Requisitions
    </div>
  </div>

  <!-- EDUCATION -->
  <div class="section-title">Education</div>
  <div class="edu-row">
    <div>
      <div class="edu-title">Bachelor of Business Administration (BBA) — International Business</div>
      <div class="edu-sub">Dayananda Sagar University (DSU), Bengaluru, Karnataka</div>
    </div>
    <div class="date-loc">2023 – 2026 | Fresh Graduate (Immediate Joining)</div>
  </div>
  <p style="font-size: 7.8pt; color: #475569; margin-top: 1.5px;">
    <strong>Relevant Coursework:</strong> Global Logistics & Supply Chain Management, Incoterms 2020, International Trade Finance & Documentation (Letters of Credit, Shipping Bills), Commercial Contract Law, Business Statistics & Quantitative Decision Modeling.
  </p>

  <!-- PROFESSIONAL SUMMARY -->
  <div class="section-title">Professional Summary</div>
  <p class="summary-text">
    Goal-driven <strong>2026 BBA graduate in International Business from Dayananda Sagar University, Bengaluru</strong>, seeking an entry-level <strong>Level 13 New Associate</strong> role at Accenture Operations. Academic grounding in global supply chains, Procure-to-Pay (P2P) cycles, purchase order (PO) processing, and regulatory compliance. Backed by corporate internship experience in commercial operations (Pencil Mark), AI data validation and content auditing (Instawork, <strong>99.4% accuracy</strong>), and high-volume event logistics (Aero India 2025, <strong>0% inventory shrinkage</strong>). Proficient in MS Excel (VLOOKUP, Pivot Tables, three-way reconciliation), ERP workflows, and non-voice query resolution. Local Bengaluru resident with <strong>0 days notice</strong> and 100% readiness for 24/7 rotational and night shifts.
  </p>

  <!-- CORE SKILLS -->
  <div class="section-title">Core Competencies & Functional Keywords</div>
  <table class="skills-table">
    <tr>
      <td class="skill-cat">Procurement & P2P Ops:</td>
      <td class="skill-items">PR-to-PO Processing, Three-Way Matching (PO/GRN/Invoice), Vendor Catalog Operations, Supplier SLA Auditing, Purchase Order Administration, Payment Dispute Reconciliation, ERP Workflows.</td>
    </tr>
    <tr>
      <td class="skill-cat">Supply Chain & Logistics:</td>
      <td class="skill-items">Incoterms 2020 (EXW, FOB, CIF, DDP), International Shipping Documentation (Bill of Lading, Commercial Invoices), Dual-Custody Inventory Governance, Physical Stock Reconciliation.</td>
    </tr>
    <tr>
      <td class="skill-cat">Commercial & Order Ops:</td>
      <td class="skill-items">B2B Order Management (100% Non-Sales), Commercial Proposal Formulation, Contract Milestones, CRM Data Hygiene, Back-Office Query De-escalation, Service Level Agreements (SLA).</td>
    </tr>
    <tr>
      <td class="skill-cat">Compliance, Risk & Trust:</td>
      <td class="skill-items">KYC & Due Diligence Documentation, Sanctions Screening Concepts, Policy Adherence, Content Moderation, Quality Assurance Auditing, Non-Voice Ticket Resolution, Fraud Detection Protocols.</td>
    </tr>
    <tr>
      <td class="skill-cat">Tools & Data Analysis:</td>
      <td class="skill-items">Advanced MS Excel (VLOOKUP, INDEX/MATCH, Pivot Tables, Conditional Formatting), Google Workspace, AI Data Quality Auditing (99.4% Accuracy), Python Automation Scripts.</td>
    </tr>
  </table>

  <!-- INTERNSHIPS & PRACTICAL EXPERIENCE -->
  <div class="section-title">Internships & Practical Field Exposure</div>

  <!-- Role 1: Pencil Mark -->
  <div class="item-block">
    <div class="item-header">
      <div>
        <span class="role-title">Commercial Operations Intern</span> — <span class="company-name">Pencil Mark Interior Solutions LLP</span>
      </div>
      <div class="date-loc">July 2024 – Aug 2025 | Bengaluru, India</div>
    </div>
    <ul class="bullets">
      <li>Administered supplier vendor matrices, evaluated competing subcontractor quotations, and prepared comparative pricing sheets in MS Excel.</li>
      <li>Formulated commercial proposals and scope-of-work documentation, directly supporting the conversion of <strong>INR 1.5L+</strong> in secured B2B corporate contracts.</li>
      <li>Tracked purchase orders against delivery milestones and site sign-offs, maintaining 100% vendor SLA compliance and zero billing discrepancies.</li>
    </ul>
  </div>

  <!-- Role 2: Instawork -->
  <div class="item-block">
    <div class="item-header">
      <div>
        <span class="role-title">Operations & AI Data Verification Intern</span> — <span class="company-name">Instawork Services India Pvt. Ltd.</span>
      </div>
      <div class="date-loc">2025 | Bengaluru, India</div>
    </div>
    <ul class="bullets">
      <li>Executed high-volume multimodal data extraction and operational auditing adhering to strict standardized operating procedures (SOPs).</li>
      <li>Maintained a verified <strong>99.4% data accuracy rating</strong> across complex operational queues, filtering non-compliant assets before database ingestion.</li>
      <li>Collaborated with team supervisors to identify workflow bottlenecks, reducing review turnaround cycle time by <strong>12%</strong>.</li>
    </ul>
  </div>

  <!-- Role 3: Aero India -->
  <div class="item-block">
    <div class="item-header">
      <div>
        <span class="role-title">Exhibition Operations & Logistics Lead</span> — <span class="company-name">Salt in My Cocoa (Aero India 2025)</span>
      </div>
      <div class="date-loc">Feb 2025 | Yelahanka Air Force Station, Bengaluru</div>
    </div>
    <ul class="bullets">
      <li>Directed on-ground stall setup, inventory replenishment, and security clearances across 5 days for <strong>100,000+ public and defense visitors</strong>.</li>
      <li>Established a strict dual-custody physical inventory protocol, achieving <strong>0% stock shrinkage</strong> and zero stock-outs during peak hours.</li>
      <li>Reconciled high-frequency digital UPI transactions daily against physical inventory ledgers with 100% financial balancing.</li>
    </ul>
  </div>

  <!-- Role 4: NGO / Community Ops -->
  <div class="item-block">
    <div class="item-header">
      <div>
        <span class="role-title">Community Operations & Compliance Intern</span> — <span class="company-name">Non-Profit Social Initiative</span>
      </div>
      <div class="date-loc">2024 | Bengaluru, India</div>
    </div>
    <ul class="bullets">
      <li>Audited and verified participant registration records and beneficiary KYC documentation in compliance with organizational guidelines.</li>
      <li>Managed database records in Google Sheets and coordinated volunteer logistics rosters across Bengaluru urban outreach drives.</li>
    </ul>
  </div>

  <!-- TECHNICAL PROJECTS & AUTOMATION -->
  <div class="section-title">Technical Projects & Operational Automation</div>
  <div class="item-block">
    <div class="item-header">
      <div>
        <span class="role-title">Autonomous Operations & Recruitment Pipeline Engine</span> — <span class="company-name">Student Automation Project</span>
      </div>
      <div class="date-loc">2026 | Python, SQLite, Workday API</div>
    </div>
    <ul class="bullets">
      <li>Architected an automated operational dispatcher in Python parsing 430+ enterprise job requisitions with 16/16 unit tests passing green.</li>
      <li>Built interactive operational consoles eliminating repetitive manual tracking and automating requisition-to-candidate alignment.</li>
    </ul>
  </div>

  <!-- CERTIFICATIONS & LANGUAGES -->
  <div class="section-title">Certifications & Languages</div>
  <p style="font-size: 8.3pt; color: #334155;">
    <strong>Certifications:</strong> Digital Marketing Fundamentals (Google) • Service Marketing (NPTEL) • Enterprise Operational Excellence<br>
    <strong>Languages:</strong> English (Fluent Professional Proficiency), Hindi (Native / Bilingual)
  </p>

</div>

</body>
</html>
"""

html_path = os.path.abspath('resumes/Aditya_Mehra_Accenture_CV.html')
pdf_path = os.path.abspath('resumes_pdf/Aditya_Mehra_Accenture_Operations_CV.pdf')
root_pdf = os.path.abspath('Aditya_Mehra_Accenture_Operations_CV.pdf')
alias_pdf = os.path.abspath('resumes_pdf/Aditya_Mehra_Resume_001_accenture_india.pdf')

with open(html_path, 'w', encoding='utf-8') as f:
    f.write(master_html)

chrome_path = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
cmd = [
    chrome_path,
    "--headless",
    "--disable-gpu",
    f"--print-to-pdf={pdf_path}",
    "--no-pdf-header-footer",
    html_path
]
subprocess.run(cmd, capture_output=True, text=True)

if os.path.exists(pdf_path):
    shutil.copyfile(pdf_path, root_pdf)
    shutil.copyfile(pdf_path, alias_pdf)
    print(f"Generated Master PDF CV: {pdf_path} (Size: {os.path.getsize(pdf_path)} bytes)")

# -------------------------------------------------------------
# 2. GENERATE MASTER UNIVERSAL FRESHER CV (DOCX)
# -------------------------------------------------------------
doc = docx.Document()
for s in doc.sections:
    s.top_margin = Inches(0.35)
    s.bottom_margin = Inches(0.35)
    s.left_margin = Inches(0.45)
    s.right_margin = Inches(0.45)

def format_run(run, font_name="Arial", size_pt=9, bold=False, italic=False, color=DARK):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color

def add_header_docx(d):
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("ADITYA MEHRA")
    format_run(r, size_pt=18, bold=True, color=NAVY)

    p_sub = d.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(2)
    r_sub = p_sub.add_run("BBA GRADUATE (BATCH OF 2026) // ENTRY-LEVEL OPERATIONS NEW ASSOCIATE")
    format_run(r_sub, size_pt=9.5, bold=True, color=BLUE)

    p_c = d.add_paragraph()
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c.paragraph_format.space_before = Pt(0)
    p_c.paragraph_format.space_after = Pt(2)
    r_c = p_c.add_run("Bengaluru, Karnataka, India  |  +91 7003456624  |  adityamehra799@gmail.com  |  linkedin.com/in/aditya-mehra")
    format_run(r_c, size_pt=8.5, color=SLATE)

    p_tag = d.add_paragraph()
    p_tag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tag.paragraph_format.space_before = Pt(0)
    p_tag.paragraph_format.space_after = Pt(5)
    r_t = p_tag.add_run("Fresher (0–1 Yrs)  •  Immediate Joiner (0 Days Notice)  •  100% Rotational & Night Shifts  •  Level 13 Requisitions")
    format_run(r_t, size_pt=8, bold=True, color=RGBColor(4, 120, 87))

def add_sec_title_docx(d, text):
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text.upper())
    format_run(r, size_pt=9.5, bold=True, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                     r'<w:bottom w:val="single" w:sz="12" w:space="1" w:color="0F172A"/>'
                     r'</w:pBdr>')
    pPr.append(pBdr)

add_header_docx(doc)

# Education
add_sec_title_docx(doc, "Education")
p_edu = doc.add_paragraph()
p_edu.paragraph_format.space_before = Pt(2)
p_edu.paragraph_format.space_after = Pt(1)
p_edu.paragraph_format.keep_with_next = True
r_e1 = p_edu.add_run("Bachelor of Business Administration (BBA) — International Business\n")
format_run(r_e1, size_pt=9, bold=True, color=NAVY)
r_e2 = p_edu.add_run("Dayananda Sagar University (DSU), Bengaluru, Karnataka    (2023 – 2026 | Fresh Graduate / Immediate Joining)")
format_run(r_e2, size_pt=8.5, bold=True, color=BLUE)

p_es = doc.add_paragraph()
p_es.paragraph_format.space_before = Pt(0)
p_es.paragraph_format.space_after = Pt(2)
r_es = p_es.add_run("Relevant Coursework: Global Logistics & Supply Chain Management, Incoterms 2020, International Trade Finance & Documentation (Letters of Credit, Shipping Bills), Commercial Contract Law, Business Statistics & Quantitative Decision Modeling.")
format_run(r_es, size_pt=8, italic=True, color=SLATE)

# Summary
add_sec_title_docx(doc, "Professional Summary")
p_sum = doc.add_paragraph()
p_sum.paragraph_format.space_before = Pt(0)
p_sum.paragraph_format.space_after = Pt(3)
p_sum.paragraph_format.line_spacing = 1.15
p_sum.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
r_sum = p_sum.add_run(
    "Goal-driven 2026 BBA graduate in International Business from Dayananda Sagar University, Bengaluru, seeking an entry-level "
    "Level 13 New Associate role at Accenture Operations. Academic grounding in global supply chains, Procure-to-Pay (P2P) cycles, "
    "purchase order (PO) processing, and regulatory compliance. Backed by corporate internship experience in commercial operations "
    "(Pencil Mark), AI data validation and content auditing (Instawork, 99.4% accuracy), and high-volume event logistics "
    "(Aero India 2025, 0% inventory shrinkage). Proficient in MS Excel (VLOOKUP, Pivot Tables, three-way reconciliation), ERP workflows, "
    "and non-voice query resolution. Local Bengaluru resident with 0 days notice and 100% readiness for 24/7 rotational and night shifts."
)
format_run(r_sum, size_pt=8.5, color=DARK)

# Core Skills
add_sec_title_docx(doc, "Core Competencies & Functional Keywords")
skills_data = [
    ("Procurement & P2P Ops:", "PR-to-PO Processing, Three-Way Matching (PO/GRN/Invoice), Vendor Catalog Operations, Supplier SLA Auditing, Purchase Order Administration, Payment Dispute Reconciliation, ERP Workflows."),
    ("Supply Chain & Logistics:", "Incoterms 2020 (EXW, FOB, CIF, DDP), International Shipping Documentation (Bill of Lading, Commercial Invoices), Dual-Custody Inventory Governance, Physical Stock Reconciliation."),
    ("Commercial & Order Ops:", "B2B Order Management (100% Non-Sales), Commercial Proposal Formulation, Contract Milestones, CRM Data Hygiene, Back-Office Query De-escalation, Service Level Agreements (SLA)."),
    ("Compliance, Risk & Trust:", "KYC & Due Diligence Documentation, Sanctions Screening Concepts, Policy Adherence, Content Moderation, Quality Assurance Auditing, Non-Voice Ticket Resolution, Fraud Detection Protocols."),
    ("Tools & Data Analysis:", "Advanced MS Excel (VLOOKUP, INDEX/MATCH, Pivot Tables, Conditional Formatting), Google Workspace, AI Data Quality Auditing (99.4% Accuracy), Python Automation Scripts.")
]
for lbl, dsc in skills_data:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1.5)
    p.paragraph_format.line_spacing = 1.1
    r_l = p.add_run(lbl + " ")
    format_run(r_l, size_pt=8.3, bold=True, color=NAVY)
    r_d = p.add_run(dsc)
    format_run(r_d, size_pt=8.3, color=DARK)

# Internships
add_sec_title_docx(doc, "Internships & Practical Field Exposure")
experiences = [
    {
        "role": "Commercial Operations Intern",
        "company": "Pencil Mark Interior Solutions LLP",
        "date_loc": "July 2024 – Aug 2025 | Bengaluru, India",
        "bullets": [
            "Administered supplier vendor matrices, evaluated competing subcontractor quotations, and prepared comparative pricing sheets in MS Excel.",
            "Formulated commercial proposals and scope-of-work documentation, directly supporting the conversion of INR 1.5L+ in secured B2B corporate contracts.",
            "Tracked purchase orders against delivery milestones and site sign-offs, maintaining 100% vendor SLA compliance and zero billing discrepancies."
        ]
    },
    {
        "role": "Operations & AI Data Verification Intern",
        "company": "Instawork Services India Pvt. Ltd.",
        "date_loc": "2025 | Bengaluru, India",
        "bullets": [
            "Executed high-volume multimodal data extraction and operational auditing adhering to strict standardized operating procedures (SOPs).",
            "Maintained a verified 99.4% data accuracy rating across complex operational queues, filtering non-compliant assets before database ingestion.",
            "Collaborated with team supervisors to identify workflow bottlenecks, reducing review turnaround cycle time by 12%."
        ]
    },
    {
        "role": "Exhibition Operations & Logistics Lead",
        "company": "Salt in My Cocoa (Aero India 2025)",
        "date_loc": "Feb 2025 | Yelahanka Air Force Station, Bengaluru",
        "bullets": [
            "Directed on-ground stall setup, inventory replenishment, and security clearances across 5 days for 100,000+ public and defense visitors.",
            "Established a strict dual-custody physical inventory protocol, achieving 0% stock shrinkage and zero stock-outs during peak hours.",
            "Reconciled high-frequency digital UPI transactions daily against physical inventory ledgers with 100% financial balancing."
        ]
    },
    {
        "role": "Community Operations & Compliance Intern",
        "company": "Non-Profit Social Initiative",
        "date_loc": "2024 | Bengaluru, India",
        "bullets": [
            "Audited and verified participant registration records and beneficiary KYC documentation in compliance with organizational guidelines.",
            "Managed database records in Google Sheets and coordinated volunteer logistics rosters across Bengaluru urban outreach drives."
        ]
    }
]

for exp in experiences:
    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_before = Pt(3)
    p_head.paragraph_format.space_after = Pt(1)
    p_head.paragraph_format.keep_with_next = True
    
    r_role = p_head.add_run(exp["role"] + " — ")
    format_run(r_role, size_pt=8.8, bold=True, color=NAVY)
    r_comp = p_head.add_run(exp["company"])
    format_run(r_comp, size_pt=8.8, bold=True, color=BLUE)
    r_tab = p_head.add_run(f"    ({exp['date_loc']})")
    format_run(r_tab, size_pt=7.8, italic=True, color=SLATE)

    for b in exp["bullets"]:
        p_b = doc.add_paragraph(style='List Bullet')
        p_b.paragraph_format.space_before = Pt(0)
        p_b.paragraph_format.space_after = Pt(1)
        p_b.paragraph_format.line_spacing = 1.1
        r_b = p_b.add_run(b)
        format_run(r_b, size_pt=8.2, color=DARK)

# Technical Projects
add_sec_title_docx(doc, "Technical Projects & Operational Automation")
p_proj = doc.add_paragraph()
p_proj.paragraph_format.space_before = Pt(2)
p_proj.paragraph_format.space_after = Pt(1)
p_proj.paragraph_format.keep_with_next = True
r_pr1 = p_proj.add_run("Autonomous Operations & Recruitment Pipeline Engine — ")
format_run(r_pr1, size_pt=8.8, bold=True, color=NAVY)
r_pr2 = p_proj.add_run("Student Automation Project (Python, SQLite, 2026)")
format_run(r_pr2, size_pt=8.3, italic=True, color=BLUE)

p_pb1 = doc.add_paragraph(style='List Bullet')
p_pb1.paragraph_format.space_before = Pt(0)
p_pb1.paragraph_format.space_after = Pt(1)
r_pb1 = p_pb1.add_run("Architected an automated operational dispatcher in Python parsing 430+ enterprise job requisitions with 16/16 unit tests passing green.")
format_run(r_pb1, size_pt=8.2, color=DARK)

p_pb2 = doc.add_paragraph(style='List Bullet')
p_pb2.paragraph_format.space_before = Pt(0)
p_pb2.paragraph_format.space_after = Pt(1)
r_pb2 = p_pb2.add_run("Built interactive operational consoles eliminating repetitive manual tracking and automating requisition-to-candidate alignment.")
format_run(r_pb2, size_pt=8.2, color=DARK)

# Certifications & Languages
add_sec_title_docx(doc, "Certifications & Languages")
p_cert = doc.add_paragraph()
p_cert.paragraph_format.space_before = Pt(2)
p_cert.paragraph_format.space_after = Pt(0)
r_c1 = p_cert.add_run("Certifications: ")
format_run(r_c1, size_pt=8.3, bold=True, color=NAVY)
r_c2 = p_cert.add_run("Digital Marketing Fundamentals (Google)  •  Service Marketing (NPTEL)  •  Enterprise Operational Excellence\n")
format_run(r_c2, size_pt=8.3, color=DARK)
r_l1 = p_cert.add_run("Languages: ")
format_run(r_l1, size_pt=8.3, bold=True, color=NAVY)
r_l2 = p_cert.add_run("English (Fluent Professional Proficiency), Hindi (Native / Bilingual)")
format_run(r_l2, size_pt=8.3, color=DARK)

docx_path = r"e:\anti\resumes\Aditya_Mehra_Accenture_Operations_CV.docx"
root_docx = r"e:\anti\Aditya_Mehra_Accenture_Operations_CV.docx"
doc.save(docx_path)
shutil.copyfile(docx_path, root_docx)
print(f"Generated Master DOCX CV: {docx_path}")

print("\n--- MASTER UNIVERSAL CV GENERATION COMPLETE ---")
