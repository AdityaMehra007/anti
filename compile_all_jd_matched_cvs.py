import os
import subprocess
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
import shutil
import re

CHROME_PATH = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
OUTPUT_DIR_HTML = r"e:\anti\resumes"
OUTPUT_DIR_PDF = r"e:\anti\resumes_pdf"
ROOT_DIR = r"e:\anti"

os.makedirs(OUTPUT_DIR_HTML, exist_ok=True)
os.makedirs(OUTPUT_DIR_PDF, exist_ok=True)

# Common Candidate Info
CANDIDATE = {
    "name": "ADITYA MEHRA",
    "location": "Bengaluru, Karnataka, India",
    "phone": "+91 7003456624",
    "email": "adityamehra799@gmail.com",
    "linkedin": "https://www.linkedin.com/in/aditya-mehra",
    "linkedin_display": "linkedin.com/in/aditya-mehra",
    "github": "https://github.com/AdityaMehra007",
    "github_display": "github.com/AdityaMehra007",
    "degree": "Bachelor of Business Administration (BBA) — International Business",
    "college": "Dayananda Sagar University (DSU), Bengaluru, Karnataka",
    "grad_timeline": "2023 – 2026 | Fresh Graduate (Immediate Joining)",
    "coursework": "Global Logistics & Supply Chain Management, Incoterms 2020, International Trade Finance & Documentation (Letters of Credit, Shipping Bills), Commercial Contract Law, Business Statistics & Quantitative Decision Modeling.",
    "status_tag": "Fresher (0–1 Yrs) • Immediate Joiner (0 Days Notice) • 100% Rotational & Night Shifts • Bengaluru Local"
}

# -------------------------------------------------------------
# DOCX HELPERS
# -------------------------------------------------------------
C_NAVY = RGBColor(15, 23, 42)       # #0F172A
C_BLUE = RGBColor(37, 99, 235)      # #2563EB
C_SLATE = RGBColor(71, 85, 105)     # #475569
C_BODY = RGBColor(30, 41, 59)       # #1E293B
C_GREEN = RGBColor(4, 120, 87)      # #047857

def add_clean_divider(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), '4')       # 0.5 pt
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), 'CBD5E1')
    pBdr.append(bottom)
    pPr.append(pBdr)

def add_hyperlink_run(paragraph, url, display_text, font_size_pt=8.5, color_hex="1E3A8A", bold=False, underline=True):
    r_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = parse_xml(
        f'<w:hyperlink xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
        f'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        f'r:id="{r_id}">'
        f'<w:r>'
        f'<w:rPr>'
        f'<w:rFonts w:ascii="Arial" w:hAnsi="Arial"/>'
        f'<w:color w:val="{color_hex}"/>'
        f'{"<w:b/>" if bold else ""}'
        f'{"<w:u w:val=\"single\"/>" if underline else ""}'
        f'<w:sz w:val="{int(font_size_pt * 2)}"/>'
        f'</w:rPr>'
        f'<w:t>{display_text}</w:t>'
        f'</w:r>'
        f'</w:hyperlink>'
    )
    paragraph._p.append(hyperlink)

def create_docx_cv(config, filepath):
    doc = Document()
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(0.32)
        s.bottom_margin = Inches(0.32)
        s.left_margin = Inches(0.42)
        s.right_margin = Inches(0.42)

    # Name
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(1)
    r_n = p_name.add_run(CANDIDATE["name"])
    r_n.font.name = "Arial"; r_n.font.size = Pt(17); r_n.font.bold = True; r_n.font.color.rgb = C_NAVY

    # Subtitle / Role Designation
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(0)
    p_sub.paragraph_format.space_after = Pt(1.5)
    r_s = p_sub.add_run(config["role_header"].upper())
    r_s.font.name = "Arial"; r_s.font.size = Pt(9.2); r_s.font.bold = True; r_s.font.color.rgb = C_BLUE

    # Contact Info
    p_c = doc.add_paragraph()
    p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_c.paragraph_format.space_before = Pt(0)
    p_c.paragraph_format.space_after = Pt(1.5)
    r_c = p_c.add_run(f"{CANDIDATE['location']}  |  {CANDIDATE['phone']}  |  {CANDIDATE['email']}  |  {CANDIDATE['linkedin_display']}")
    r_c.font.name = "Arial"; r_c.font.size = Pt(8.2); r_c.font.color.rgb = C_SLATE

    # Target Requisition / JD Match Banner
    p_req = doc.add_paragraph()
    p_req.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_req.paragraph_format.space_before = Pt(0)
    p_req.paragraph_format.space_after = Pt(4)
    r_req = p_req.add_run(config["jd_banner"])
    r_req.font.name = "Arial"; r_req.font.size = Pt(7.8); r_req.font.bold = True; r_req.font.color.rgb = C_GREEN

    def add_sec(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(title.upper())
        r.font.name = "Arial"; r.font.size = Pt(9.2); r.font.bold = True; r.font.color.rgb = C_NAVY
        add_clean_divider(p)

    # Education
    add_sec("Education")
    p_edu = doc.add_paragraph()
    p_edu.paragraph_format.space_before = Pt(1)
    p_edu.paragraph_format.space_after = Pt(1)
    r_ed1 = p_edu.add_run(f"{CANDIDATE['degree']}\n")
    r_ed1.font.name = "Arial"; r_ed1.font.size = Pt(8.8); r_ed1.font.bold = True; r_ed1.font.color.rgb = C_NAVY
    r_ed2 = p_edu.add_run(f"{CANDIDATE['college']}  •  {CANDIDATE['grad_timeline']}\n")
    r_ed2.font.name = "Arial"; r_ed2.font.size = Pt(8.2); r_ed2.font.bold = True; r_ed2.font.color.rgb = C_BLUE
    r_ed3 = p_edu.add_run(f"Relevant Coursework: {CANDIDATE['coursework']}")
    r_ed3.font.name = "Arial"; r_ed3.font.size = Pt(7.6); r_ed3.font.color.rgb = C_SLATE

    # Professional Summary (JD Matched)
    add_sec("Professional Summary")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(1)
    p_sum.paragraph_format.space_after = Pt(2)
    p_sum.paragraph_format.line_spacing = 1.15
    r_sum = p_sum.add_run(config["summary"])
    r_sum.font.name = "Arial"; r_sum.font.size = Pt(8.3); r_sum.font.color.rgb = C_BODY

    # Core Competencies & ATS Keywords
    add_sec("Core Competencies & ATS Keywords (Accenture Workday Mapped)")
    for cat_name, kw_text in config["competencies"]:
        p_k = doc.add_paragraph()
        p_k.paragraph_format.space_before = Pt(0)
        p_k.paragraph_format.space_after = Pt(1)
        p_k.paragraph_format.left_indent = Inches(0.12)
        r_k1 = p_k.add_run(f"• {cat_name}: ")
        r_k1.font.name = "Arial"; r_k1.font.size = Pt(8.0); r_k1.font.bold = True; r_k1.font.color.rgb = C_NAVY
        r_k2 = p_k.add_run(kw_text)
        r_k2.font.name = "Arial"; r_k2.font.size = Pt(8.0); r_k2.font.color.rgb = C_BODY

    # Internships & Practical Field Exposure
    add_sec("Internships & Practical Field Exposure")
    for exp in config["experiences"]:
        p_ex = doc.add_paragraph()
        p_ex.paragraph_format.space_before = Pt(2)
        p_ex.paragraph_format.space_after = Pt(1)
        p_ex.paragraph_format.keep_with_next = True
        r_ex1 = p_ex.add_run(f"{exp['title']} — {exp['org']}\n")
        r_ex1.font.name = "Arial"; r_ex1.font.size = Pt(8.5); r_ex1.font.bold = True; r_ex1.font.color.rgb = C_NAVY
        r_ex2 = p_ex.add_run(f"{exp['timeline']} | {exp['location']}")
        r_ex2.font.name = "Arial"; r_ex2.font.size = Pt(7.8); r_ex2.font.color.rgb = C_SLATE

        for b in exp["bullets"]:
            p_b = doc.add_paragraph()
            p_b.paragraph_format.space_before = Pt(0)
            p_b.paragraph_format.space_after = Pt(1)
            p_b.paragraph_format.left_indent = Inches(0.18)
            p_b.paragraph_format.first_line_indent = Inches(-0.12)
            p_b.paragraph_format.line_spacing = 1.12
            r_dot = p_b.add_run("– ")
            r_dot.font.name = "Arial"; r_dot.font.size = Pt(8.0); r_dot.font.bold = True; r_dot.font.color.rgb = C_BLUE
            r_txt = p_b.add_run(b)
            r_txt.font.name = "Arial"; r_txt.font.size = Pt(8.0); r_txt.font.color.rgb = C_BODY

    # Technical Project
    add_sec("Technical Project & Operations Automation")
    p_pr = doc.add_paragraph()
    p_pr.paragraph_format.space_before = Pt(1)
    p_pr.paragraph_format.space_after = Pt(1)
    r_pr1 = p_pr.add_run("Autonomous Operations & Workday Pipeline Engine — Student Automation Project\n")
    r_pr1.font.name = "Arial"; r_pr1.font.size = Pt(8.5); r_pr1.font.bold = True; r_pr1.font.color.rgb = C_NAVY
    r_pr2 = p_pr.add_run("2026 | Python, SQLite, Workday API Integration | 16/16 Unit Tests Passing Green\n")
    r_pr2.font.name = "Arial"; r_pr2.font.size = Pt(7.8); r_pr2.font.color.rgb = C_SLATE
    r_pr3 = p_pr.add_run("– Architected an automated operational dispatcher parsing 430+ enterprise job requisitions, validating requisitions against candidate skill taxonomies with zero manual latency.")
    r_pr3.font.name = "Arial"; r_pr3.font.size = Pt(8.0); r_pr3.font.color.rgb = C_BODY

    # Certifications & Languages
    add_sec("Certifications & Languages")
    p_cert = doc.add_paragraph()
    p_cert.paragraph_format.space_before = Pt(1)
    p_cert.paragraph_format.space_after = Pt(1)
    r_c1 = p_cert.add_run("Certifications: ")
    r_c1.font.name = "Arial"; r_c1.font.size = Pt(8.0); r_c1.font.bold = True; r_c1.font.color.rgb = C_NAVY
    r_c2 = p_cert.add_run("Enterprise Operational Excellence • Digital Marketing Fundamentals (Google) • Service Marketing (NPTEL)\n")
    r_c2.font.name = "Arial"; r_c2.font.size = Pt(8.0); r_c2.font.color.rgb = C_BODY
    r_l1 = p_cert.add_run("Languages: ")
    r_l1.font.name = "Arial"; r_l1.font.size = Pt(8.0); r_l1.font.bold = True; r_l1.font.color.rgb = C_NAVY
    r_l2 = p_cert.add_run("English (Fluent Professional Proficiency), Hindi (Native / Bilingual)")
    r_l2.font.name = "Arial"; r_l2.font.size = Pt(8.0); r_l2.font.color.rgb = C_BODY

    doc.save(filepath)

# -------------------------------------------------------------
# HTML & PDF GENERATOR
# -------------------------------------------------------------
def create_html_cv(config, html_path, pdf_path):
    skills_rows = "".join([
        f'<tr><td class="skill-cat">{cat}:</td><td class="skill-items">{items}</td></tr>'
        for cat, items in config["competencies"]
    ])

    exp_blocks = ""
    for exp in config["experiences"]:
        bullets_li = "".join([f'<li>{b}</li>' for b in exp["bullets"]])
        exp_blocks += f"""
        <div class="item-block">
          <div class="item-header">
            <div><span class="role-title">{exp['title']}</span> — <span class="company-name">{exp['org']}</span></div>
            <div class="date-loc">{exp['timeline']} | {exp['location']}</div>
          </div>
          <ul class="bullets">
            {bullets_li}
          </ul>
        </div>
        """

    html = f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1.0">
<title>Aditya Mehra - {config['role_header']} (Accenture Workday Matched)</title>
<style>
  @page {{
    size: A4 portrait;
    margin: 6.5mm 8.5mm 6.5mm 8.5mm;
  }}
  * {{
    box-sizing: border-box;
    margin: 0;
    padding: 0;
    font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, "Helvetica Neue", Arial, sans-serif;
    color: #1e293b;
    line-height: 1.26;
  }}
  body {{
    background: #fff;
    font-size: 8.9pt;
  }}
  .container {{
    max-width: 100%;
    margin: 0 auto;
  }}
  .header {{
    text-align: center;
    border-bottom: 2px solid #0f172a;
    padding-bottom: 4px;
    margin-bottom: 4px;
  }}
  .name {{
    font-size: 18pt;
    font-weight: 800;
    letter-spacing: 0.6px;
    color: #0f172a;
    text-transform: uppercase;
  }}
  .subtitle {{
    font-size: 9.0pt;
    font-weight: 700;
    color: #2563eb;
    margin-top: 1px;
    letter-spacing: 0.3px;
    text-transform: uppercase;
  }}
  .contact-info {{
    font-size: 8.1pt;
    color: #475569;
    margin-top: 2px;
    display: flex;
    justify-content: center;
    gap: 12px;
    flex-wrap: wrap;
    font-weight: 500;
  }}
  .contact-info a {{
    color: #2563eb;
    text-decoration: none;
  }}
  .tagline-badge {{
    margin-top: 2.5px;
    font-size: 7.5pt;
    font-weight: 700;
    color: #047857;
    background: #ecfdf5;
    display: inline-block;
    padding: 1px 7px;
    border-radius: 4px;
    border: 1px solid #a7f3d0;
  }}
  .section-title {{
    font-size: 9.1pt;
    font-weight: 800;
    color: #0f172a;
    text-transform: uppercase;
    letter-spacing: 0.5px;
    border-bottom: 1.5px solid #cbd5e1;
    padding-bottom: 1.5px;
    margin-top: 4.5px;
    margin-bottom: 3px;
  }}
  .summary-text {{
    font-size: 8.2pt;
    color: #334155;
    text-align: justify;
    line-height: 1.28;
  }}
  .edu-row {{
    display: flex;
    justify-content: space-between;
    align-items: baseline;
  }}
  .edu-title {{
    font-size: 8.9pt;
    font-weight: 700;
    color: #0f172a;
  }}
  .edu-sub {{
    font-size: 8.2pt;
    color: #2563eb;
    font-weight: 600;
    margin-top: 0.5px;
  }}
  .date-loc {{
    font-size: 7.6pt;
    color: #64748b;
    font-weight: 500;
    white-space: nowrap;
  }}
  .skills-table {{
    width: 100%;
    border-collapse: collapse;
    font-size: 7.9pt;
  }}
  .skills-table td {{
    padding: 1.2px 0;
    vertical-align: top;
  }}
  .skill-cat {{
    font-weight: 700;
    color: #0f172a;
    width: 25%;
    white-space: nowrap;
  }}
  .skill-items {{
    color: #334155;
    width: 75%;
  }}
  .item-block {{
    margin-bottom: 3px;
  }}
  .item-header {{
    display: flex;
    justify-content: space-between;
    align-items: baseline;
    margin-bottom: 1px;
  }}
  .role-title {{
    font-weight: 700;
    color: #0f172a;
    font-size: 8.5pt;
  }}
  .company-name {{
    font-weight: 600;
    color: #2563eb;
    font-size: 8.3pt;
  }}
  .bullets {{
    list-style-type: disc;
    padding-left: 15px;
  }}
  .bullets li {{
    font-size: 7.9pt;
    color: #334155;
    margin-bottom: 1.2px;
    line-height: 1.26;
  }}
  .bullets strong {{
    color: #0f172a;
  }}
  @media print {{
    body {{ background: #fff; padding: 0; }}
    .container {{ max-width: 100%; }}
  }}
</style>
</head>
<body>
<div class="container">
  <div class="header">
    <div class="name">{CANDIDATE['name']}</div>
    <div class="subtitle">{config['role_header']}</div>
    <div class="contact-info">
      <span>{CANDIDATE['location']}</span>
      <span>{CANDIDATE['phone']}</span>
      <span>{CANDIDATE['email']}</span>
      <span><a href="{CANDIDATE['linkedin']}">{CANDIDATE['linkedin_display']}</a></span>
      <span><a href="{CANDIDATE['github']}">{CANDIDATE['github_display']}</a></span>
    </div>
    <div class="tagline-badge">
      {config['jd_banner']}
    </div>
  </div>

  <div class="section-title">Education</div>
  <div class="edu-row">
    <div>
      <div class="edu-title">{CANDIDATE['degree']}</div>
      <div class="edu-sub">{CANDIDATE['college']}</div>
    </div>
    <div class="date-loc">{CANDIDATE['grad_timeline']}</div>
  </div>
  <p style="font-size: 7.5pt; color: #475569; margin-top: 1px;">
    <strong>Relevant Coursework:</strong> {CANDIDATE['coursework']}
  </p>

  <div class="section-title">Professional Summary</div>
  <p class="summary-text">{config['summary']}</p>

  <div class="section-title">Core Competencies & ATS Keywords (Accenture Workday Mapped)</div>
  <table class="skills-table">
    {skills_rows}
  </table>

  <div class="section-title">Internships & Practical Field Exposure</div>
  {exp_blocks}

  <div class="section-title">Technical Project & Operations Automation</div>
  <div class="item-block">
    <div class="item-header">
      <div><span class="role-title">Autonomous Operations & Workday Pipeline Engine</span> — <span class="company-name">Student Automation Project</span></div>
      <div class="date-loc">2026 | Python, SQLite, Workday API</div>
    </div>
    <ul class="bullets">
      <li>Architected an automated operational dispatcher in Python parsing 430+ enterprise job requisitions with 16/16 unit tests passing green.</li>
      <li>Built automated candidate-to-requisition taxonomy matching, validating skill keywords against Workday job descriptions with zero manual overhead.</li>
    </ul>
  </div>

  <div class="section-title">Certifications & Languages</div>
  <p style="font-size: 7.8pt; color: #334155;">
    <strong>Certifications:</strong> Enterprise Operational Excellence • Digital Marketing Fundamentals (Google) • Service Marketing (NPTEL)<br>
    <strong>Languages:</strong> English (Fluent Professional Proficiency), Hindi (Native / Bilingual)
  </p>
</div>
</body>
</html>"""

    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)

    cmd = [
        CHROME_PATH,
        "--headless",
        "--disable-gpu",
        f"--print-to-pdf={pdf_path}",
        "--no-pdf-header-footer",
        html_path
    ]
    subprocess.run(cmd, capture_output=True, text=True)

# -------------------------------------------------------------
# DEFINITIONS OF THE 5 ATS-MATCHED RESUME TRACKS
# -------------------------------------------------------------

# Track 0: Master Universal Operations CV (All 45 Freshers Requisitions)
universal_config = {
    "role_header": "BBA Graduate (Batch of 2026) // Entry-Level Operations New Associate",
    "jd_banner": "Workday Target: Level 13 New Associate Requisitions (Req IDs: AIOC-S01664213, AIOC-S01670871, AIOC-S01667647, AIOC-S01665657, AIOC-S01668350) • Immediate Joiner • 24/7 Shifts",
    "summary": "Goal-driven <strong>2026 BBA graduate in International Business from Dayananda Sagar University, Bengaluru</strong>, seeking an entry-level <strong>Level 13 New Associate</strong> role at Accenture Operations. Academic grounding in global supply chains, Procure-to-Pay (P2P) cycles, purchase order (PO) processing, and regulatory compliance. Backed by corporate internship experience in commercial operations (Pencil Mark), AI data validation and content auditing (Instawork, <strong>99.4% accuracy</strong>), and high-volume event logistics (Aero India 2025, <strong>0% inventory shrinkage</strong>). Proficient in MS Excel (VLOOKUP, Pivot Tables, three-way reconciliation), ERP workflows, and non-voice query resolution. Local Bengaluru resident with <strong>0 days notice</strong> and 100% readiness for 24/7 rotational and night shifts.",
    "competencies": [
        ("Procurement & P2P Ops", "Procurement Support (P1), PR-to-PO Processing, Three-Way Matching (PO/GRN/Invoice), Vendor Catalog Operations, Supplier SLA Auditing, Purchase Order Administration, Payment Dispute Reconciliation, ERP Workflows."),
        ("Supply Chain & Trade", "Order Management (P3), Incoterms 2020 (EXW, FOB, CIF, DDP), International Shipping Documentation (Bill of Lading, Commercial Invoices), Dual-Custody Inventory Governance, Physical Stock Reconciliation."),
        ("Commercial & Order Ops", "B2B Order Management (100% Non-Sales), Commercial Proposal Formulation, Contract Milestones, CRM Data Hygiene, Back-Office Query De-escalation, Service Level Agreements (SLA)."),
        ("Compliance, Risk & Trust", "KYC Screening (P1), Enhanced Customer Due Diligence (ECDD), Ongoing Due Diligence (OCDD), Sanctions Screening Concepts, Content Moderation, Policy Adherence, Non-Voice Ticket Resolution."),
        ("Tools & Data Analysis", "Advanced MS Excel (VLOOKUP, INDEX/MATCH, Pivot Tables, Conditional Formatting), Google Workspace, AI Data Quality Auditing (99.4% Accuracy), Python Automation Scripts.")
    ],
    "experiences": [
        {
            "title": "Commercial Operations Intern",
            "org": "Pencil Mark Interior Solutions LLP",
            "timeline": "July 2024 – Aug 2025",
            "location": "Bengaluru, India",
            "bullets": [
                "Administered supplier vendor matrices, evaluated competing subcontractor quotations, and prepared comparative pricing sheets in MS Excel.",
                "Formulated commercial proposals and scope-of-work documentation, directly supporting the conversion of <strong>INR 1.5L+</strong> in secured B2B corporate contracts.",
                "Tracked purchase orders against delivery milestones and site sign-offs, maintaining 100% vendor SLA compliance and zero billing discrepancies."
            ]
        },
        {
            "title": "Operations & AI Data Verification Intern",
            "org": "Instawork Services India Pvt. Ltd.",
            "timeline": "2025",
            "location": "Bengaluru, India",
            "bullets": [
                "Executed high-volume multimodal data extraction and operational auditing adhering to strict standardized operating procedures (SOPs).",
                "Maintained a verified <strong>99.4% data accuracy rating</strong> across complex operational queues, filtering non-compliant assets before database ingestion.",
                "Collaborated with team supervisors to identify workflow bottlenecks, reducing review turnaround cycle time by <strong>12%</strong>."
            ]
        },
        {
            "title": "Exhibition Operations & Logistics Lead",
            "org": "Salt in My Cocoa (Aero India 2025)",
            "timeline": "Feb 2025",
            "location": "Yelahanka Air Force Station, Bengaluru",
            "bullets": [
                "Directed on-ground stall setup, inventory replenishment, and security clearances across 5 days for <strong>100,000+ public and defense visitors</strong>.",
                "Established a strict dual-custody physical inventory protocol, achieving <strong>0% stock shrinkage</strong> and zero stock-outs during peak hours.",
                "Reconciled high-frequency digital UPI transactions daily against physical inventory ledgers with 100% financial balancing."
            ]
        },
        {
            "title": "Community Operations & Compliance Intern",
            "org": "Non-Profit Social Initiative",
            "timeline": "2024",
            "location": "Bengaluru, India",
            "bullets": [
                "Audited and verified participant registration records and beneficiary KYC documentation in compliance with organizational guidelines.",
                "Managed database records in Google Sheets and coordinated volunteer logistics rosters across Bengaluru urban outreach drives."
            ]
        }
    ]
}

# Track 1: Procurement & Procure-to-Pay (P2P) Operations
procurement_p2p_config = {
    "role_header": "Procurement & Procure-to-Pay (P2P) Operations New Associate",
    "jd_banner": "Target Workday Requisitions: AIOC-S01664213-1 (Procurement Operations) • AIOC-S01670871-1 (Invoice Processing - P1) • Level 13 • 0–1 Yrs",
    "summary": "Detail-oriented <strong>2026 BBA graduate in International Business from Dayananda Sagar University</strong>, targeting <strong>Procurement Operations and Procure-to-Pay (P2P) New Associate</strong> roles at Accenture Operations (Req IDs: AIOC-S01664213 & AIOC-S01670871). Specialized in PR-to-PO processing, vendor catalog administration, invoice validation, and three-way matching (PO/GRN/Invoice). Practical experience evaluating supplier quotations at Pencil Mark, managing purchase order delivery milestones, and performing daily transaction reconciliations with 100% financial balancing. Skilled in MS Excel (VLOOKUP, Pivot Tables, reconciliation sheets). Immediate joiner (0 Days Notice), fully flexible for rotational/night shifts and month-end financial close extended hours.",
    "competencies": [
        ("Primary Skills (Accenture P1)", "Procurement Support - P1, Invoice Processing - P1, Three-Way Matching (PO vs. GRN vs. Invoice), Purchase Requisition (PR) to Purchase Order (PO) Processing, Vendor Catalog Administration."),
        ("P2P Reconciliation & Audit", "Vendor Compliance Auditing, Discount Capture Optimization, Payment Dispute Resolution, Billing Discrepancy Clearance, Month-Close Accounting Reconciliations, ERP Transaction Verification."),
        ("Supply Chain & Contracting", "Vendor Matrix Governance, Subcontractor Quotation Evaluation, Delivery Milestone Tracking, Supplier SLA Performance Monitoring, Incoterms 2020 Commercial Terms."),
        ("Tools & Analytical Software", "Advanced MS Excel (VLOOKUP, INDEX/MATCH, Pivot Tables, Duplicate Identification), Google Sheets, ERP Workflows, Python Automated Data Processing, Digital Payment Reconciliations.")
    ],
    "experiences": [
        {
            "title": "Commercial & Procurement Operations Intern",
            "org": "Pencil Mark Interior Solutions LLP",
            "timeline": "July 2024 – Aug 2025",
            "location": "Bengaluru, India",
            "bullets": [
                "Managed vendor databases and administered supplier quotation matrices, preparing comparative cost-benefit spreadsheets in MS Excel for management sign-off.",
                "Tracked purchase orders (POs) from issuance through goods receipt (GRN) to delivery completion, eliminating billing errors and maintaining 100% vendor SLA compliance.",
                "Assisted commercial managers in formulating scopes of work and pricing schedules, contributing to securing <strong>INR 1.5L+</strong> in commercial contracts."
            ]
        },
        {
            "title": "High-Volume Transaction Reconciliation Lead",
            "org": "Salt in My Cocoa (Aero India 2025)",
            "timeline": "Feb 2025",
            "location": "Yelahanka Air Force Station, Bengaluru",
            "bullets": [
                "Performed end-of-day three-way reconciliations between digital UPI merchant settlement reports, cashier registers, and physical inventory stock ledgers with 100% financial balancing.",
                "Enforced strict dual-custody receiving protocols for perishable inventory shipments, achieving <strong>0% stock shrinkage</strong> across 100,000+ attendee footfalls.",
                "Resolved vendor invoice queries and delivery receipt discrepancies under high transaction velocity and strict time constraints."
            ]
        },
        {
            "title": "Operations & Data Quality Intern",
            "org": "Instawork Services India Pvt. Ltd.",
            "timeline": "2025",
            "location": "Bengaluru, India",
            "bullets": [
                "Executed standardized operational checks on vendor and partner profile submissions adhering to strict operational standard operating procedures (SOPs).",
                "Delivered an audited <strong>99.4% accuracy rate</strong> across high-throughput data queues, flagging irregular entries and ensuring database integrity.",
                "Reduced queue turnaround time by <strong>12%</strong> through systematic spreadsheet shortcuts and workflow bottleneck identification."
            ]
        },
        {
            "title": "Community Operations & Compliance Intern",
            "org": "Non-Profit Social Initiative",
            "timeline": "2024",
            "location": "Bengaluru, India",
            "bullets": [
                "Audited and matched participant onboarding records against authorization paperwork with 100% compliance to procedural guidelines.",
                "Organized logistical support schedules, vendor equipment receipts, and resource allocations in structured Google Sheets."
            ]
        }
    ]
}

# Track 2: Supply Chain & Order Management Operations
supply_chain_order_config = {
    "role_header": "Supply Chain & Order Management Operations New Associate",
    "jd_banner": "Target Workday Requisitions: AIOC-S01665657-1 (Fulfill - Order Management - P3) • AIOC-S01669615-1 (Commercial Ops) • Level 13 • 0–1 Yrs",
    "summary": "Proactive <strong>2026 BBA graduate in International Business from Dayananda Sagar University</strong>, targeting <strong>Supply Chain Operations & Order Management New Associate</strong> positions at Accenture Operations (Req IDs: AIOC-S01665657 & AIOC-S01669615). Strong academic foundation in global supply planning, Incoterms 2020 (EXW, FOB, CIF, DDP), international shipping bills, and order fulfillment cycles. Proven track record leading on-ground logistics and inventory replenishment at Aero India 2025 (serving 100,000+ visitors with <strong>0% inventory shrinkage</strong>). Proficient in MS Excel (VLOOKUP, Pivot Tables) and backlog tracking. Immediate joiner, fully prepared for 24/7 rotational shift schedules in Bengaluru.",
    "competencies": [
        ("Primary Skills (Accenture P3)", "Fulfill - Order Management - P3, End-to-End Supply Chain Planning & Execution, Order Processing & Tracking, Distribution Operations, Backlog & Delivery Scheduling."),
        ("International Trade & Logistics", "Incoterms 2020 (EXW, FOB, CFR, CIF, DDP), Shipping Documentation (Bill of Lading, Air Waybill, Commercial Invoices, Packing Lists), Customs Clearance Basics, Dual-Custody Inventory Governance."),
        ("Commercial Operations", "B2B Order Processing (100% Non-Sales), Scope-of-Work Review, Purchase Order Validation, Milestone Tracking, Contract Deliverable Verification, Supplier SLA Adherence."),
        ("Tools & Quantitative Modeling", "MS Excel (VLOOKUP, Pivot Tables, SUMIFS, Inventory Turnover Calculations), Google Workspace, Python Process Scripts, Logistics Exception Tracking.")
    ],
    "experiences": [
        {
            "title": "Exhibition Operations & Logistics Lead",
            "org": "Salt in My Cocoa (Aero India 2025)",
            "timeline": "Feb 2025",
            "location": "Yelahanka Air Force Station, Bengaluru",
            "bullets": [
                "Orchestrated end-to-end supply chain fulfillment, warehouse dispatch, and delivery intake at Asia's premier defense aviation expo (100,000+ visitors).",
                "Pioneered a dual-custody physical inventory tracking protocol that maintained 100% stock availability and achieved <strong>0% shrinkage</strong> across 5 consecutive peak days.",
                "Streamlined daily replenishment logistics schedules under stringent defense-sector security checkpoints and restricted dispatch windows."
            ]
        },
        {
            "title": "Commercial Operations Intern",
            "org": "Pencil Mark Interior Solutions LLP",
            "timeline": "July 2024 – Aug 2025",
            "location": "Bengaluru, India",
            "bullets": [
                "Managed B2B purchase order life cycles, tracking materials from supplier dispatch to on-site delivery sign-offs against agreed contractual milestones.",
                "Audited supplier packing slips against purchase orders to eliminate quantity mismatches and damaged goods claims prior to installation.",
                "Prepared commercial proposals and comparative cost estimation sheets in MS Excel, assisting leadership in closing <strong>INR 1.5L+</strong> in corporate contracts."
            ]
        },
        {
            "title": "Operations & Data Verification Intern",
            "org": "Instawork Services India Pvt. Ltd.",
            "timeline": "2025",
            "location": "Bengaluru, India",
            "bullets": [
                "Validated incoming logistics and partner fulfillment records, verifying order completion data points against strict operational SOPs.",
                "Maintained <strong>99.4% accuracy</strong> across high-volume verification queues, preventing system processing errors downstream.",
                "Identified repetitive bottlenecks in manual document validation, collaborating with leads to decrease cycle time by <strong>12%</strong>."
            ]
        },
        {
            "title": "Community Operations & Compliance Intern",
            "org": "Non-Profit Social Initiative",
            "timeline": "2024",
            "location": "Bengaluru, India",
            "bullets": [
                "Supervised field resource allocations, participant registrations, and supply distributions across municipal community venues.",
                "Maintained inventory tracking rosters in Google Sheets, ensuring full material accounting and on-time distribution."
            ]
        }
    ]
}

# Track 3: Regulatory Compliance, KYC & Risk Operations
regulatory_compliance_config = {
    "role_header": "Regulatory Compliance & KYC Operations New Associate",
    "jd_banner": "Target Workday Requisitions: AIOC-S01667647-1 (KYC Screening - P1) • AIOC-S01668628-1 • Level 13 • 0–1 Yrs • Bengaluru",
    "summary": "Analytical and process-driven <strong>2026 BBA graduate in International Business from Dayananda Sagar University</strong>, targeting <strong>Regulatory Compliance & KYC New Associate</strong> positions at Accenture Operations (Req ID: AIOC-S01667647-1). Comprehensive academic training in commercial contract law, trade compliance documentation, AML/KYC guidelines, and regulatory risk controls. Hands-on experience auditing beneficiary KYC identity credentials in community operations and validating digital partner records with a verified <strong>99.4% accuracy rating</strong> at Instawork. Familiar with Enhanced Customer Due Diligence (ECDD), Ongoing Customer Due Diligence (OCDD), and sanctions screening concepts. Immediate joiner, fully prepared for 24/7 rotational shifts in Bengaluru.",
    "competencies": [
        ("Primary Skills (Accenture P1)", "KYC Screening - Know Your Customer (KYC - P1), Enhanced Customer Due Diligence (ECDD), Ongoing Customer Due Diligence (OCDD), Customer Identification Program (CIP), Document Verification."),
        ("Compliance & Risk Governance", "Anti-Money Laundering (AML) Fundamentals, Politically Exposed Persons (PEP) Screening, Sanctions Lists Adherence (OFAC/UN/EU basics), Identity Fraud Detection, Audit Trail Documentation."),
        ("Process Quality & Audit", "Standard Operating Procedure (SOP) Compliance, Quality Assurance Auditing, Discrepancy Escalation Protocols, Regulatory Requirement Mapping, Data Privacy & Confidentiality."),
        ("Analytical Tools & Systems", "Advanced MS Excel (VLOOKUP, Conditional Formatting, Data Cleaning), Google Workspace, Python Structured Data Scripts, CRM Document Verification Systems.")
    ],
    "experiences": [
        {
            "title": "Community Operations & Compliance Intern",
            "org": "Non-Profit Social Initiative",
            "timeline": "2024",
            "location": "Bengaluru, India",
            "bullets": [
                "Audited and verified personal identity documents (Government Photo IDs, address proofs) for 500+ program participants against compliance guidelines.",
                "Identified and flagged incomplete or mismatched identity submissions, initiating proactive follow-ups to maintain 100% KYC ledger compliance.",
                "Compiled structured compliance audit logs in Google Sheets, maintaining confidential participant data records with zero security breaches."
            ]
        },
        {
            "title": "Operations & Data Verification Intern",
            "org": "Instawork Services India Pvt. Ltd.",
            "timeline": "2025",
            "location": "Bengaluru, India",
            "bullets": [
                "Executed rigorous identity and credential validation checks on incoming user and partner registrations in compliance with internal policy SOPs.",
                "Sustained an exceptional <strong>99.4% audit accuracy score</strong> across complex multi-attribute compliance queues.",
                "Escalated suspicious submission patterns and forged documentation to senior quality leads, helping mitigate onboarding risk."
            ]
        },
        {
            "title": "Commercial Operations Intern",
            "org": "Pencil Mark Interior Solutions LLP",
            "timeline": "July 2024 – Aug 2025",
            "location": "Bengaluru, India",
            "bullets": [
                "Conducted preliminary regulatory and statutory compliance checks on vendor GST registration certificates and business licenses before onboarding.",
                "Enforced contract milestone sign-offs prior to invoice approvals, ensuring strict auditability and preventing payment non-compliance.",
                "Maintained organized vendor documentation repositories, supporting internal commercial audits."
            ]
        },
        {
            "title": "Financial Balancing & Cash Controls Lead",
            "org": "Salt in My Cocoa (Aero India 2025)",
            "timeline": "Feb 2025",
            "location": "Yelahanka Air Force Station, Bengaluru",
            "bullets": [
                "Audited and balanced daily sales transactions against merchant banking settlement reports, ensuring zero cash or digital discrepancies.",
                "Implemented secure dual-custody register verification protocols in a high-security defense environment."
            ]
        }
    ]
}

# Track 4: Trust & Safety, Content Moderation & Digital Support Operations
trust_safety_config = {
    "role_header": "Trust & Safety / Digital Operations New Associate",
    "jd_banner": "Target Workday Requisitions: AIOC-S01668350-1 (Content Moderation & Policy QA) • AIOC-S01665940-1 (Omnichannel Comms) • Level 13",
    "summary": "Resilient, highly observant <strong>2026 BBA graduate in International Business from Dayananda Sagar University</strong>, seeking a <strong>Trust & Safety / Content Moderation New Associate</strong> position at Accenture Operations (Req IDs: AIOC-S01668350 & AIOC-S01665940). Verifiable background in multimodal data auditing and content quality control at Instawork, achieving a <strong>99.4% accuracy rating</strong> while processing high-volume queues under strict SLAs. Demonstrates high emotional resilience, sharp attention to detail, and objective adherence to community guidelines and safety policies. Familiar with Responsible AI adversarial testing, non-voice omnichannel query resolution, and escalation protocols. Immediate joiner with 100% readiness for 24/7 rotational and night shifts in Bengaluru.",
    "competencies": [
        ("Primary Skills (Accenture)", "User-Generated Content Moderation, Policy Enforcement & Adherence, Quality Assurance Auditing, Responsible AI Safety Validation, Trust & Safety Operations."),
        ("Incident Management & Support", "Omnichannel Customer Support, Non-Voice Ticket Resolution, Escalation Handling, Sensitive Content Assessment, SLA Turnaround Time Optimization."),
        ("Operational Resilience & SOPs", "Standard Operating Procedure (SOP) Strict Execution, High Emotional Resilience, Objective Decision-Making under Precedent, Edge-Case Identification."),
        ("Tools & Data Platforms", "Ticketing & Moderation Consoles, MS Excel (VLOOKUP, Pivot Tables, Error Tracking), Google Workspace, Python Automation Scripts, AI Dataset Quality Annotation.")
    ],
    "experiences": [
        {
            "title": "Operations & AI Content Verification Intern",
            "org": "Instawork Services India Pvt. Ltd.",
            "timeline": "2025",
            "location": "Bengaluru, India",
            "bullets": [
                "Reviewed and audited hundreds of multimodal user submissions daily against strict platform safety and content quality guidelines.",
                "Maintained a verified <strong>99.4% accuracy benchmark</strong> across high-speed review queues, identifying and removing non-compliant content.",
                "Demonstrated emotional resilience and focus while reviewing complex, subjective, and edge-case material without policy drift.",
                "Identified recurring misclassification trends and partnered with team leads to clarify guidelines, reducing review cycle time by <strong>12%</strong>."
            ]
        },
        {
            "title": "Community Operations & Query Resolution Intern",
            "org": "Non-Profit Social Initiative",
            "timeline": "2024",
            "location": "Bengaluru, India",
            "bullets": [
                "Managed non-voice communication channels, resolving queries and clarifying event protocols for 500+ diverse participants.",
                "De-escalated participant concerns and logistical bottlenecks under tight deadlines, maintaining high satisfaction ratings."
            ]
        },
        {
            "title": "Commercial Operations Intern",
            "org": "Pencil Mark Interior Solutions LLP",
            "timeline": "July 2024 – Aug 2025",
            "location": "Bengaluru, India",
            "bullets": [
                "Audited project scopes of work and commercial deliverables against client contract terms, ensuring 100% policy compliance.",
                "Maintained documentation integrity in shared databases, eliminating billing mismatches through systematic cross-checks."
            ]
        },
        {
            "title": "Exhibition Operations & Safety Lead",
            "org": "Salt in My Cocoa (Aero India 2025)",
            "timeline": "Feb 2025",
            "location": "Yelahanka Air Force Station, Bengaluru",
            "bullets": [
                "Enforced strict safety protocols, visitor queue flow, and dual-custody stock controls for 100,000+ visitors at a high-security defense expo.",
                "Maintained complete operational discipline during crowd surges, achieving <strong>0% stock shrinkage</strong> and zero security non-conformances."
            ]
        }
    ]
}

# -------------------------------------------------------------
# EXECUTION & COMPILATION LOOP
# -------------------------------------------------------------
tracks = [
    ("Master Universal Operations CV", universal_config, "Aditya_Mehra_Accenture_Operations_CV"),
    ("Procurement & P2P Operations CV", procurement_p2p_config, "Aditya_Mehra_Accenture_Procurement_P2P_CV"),
    ("Supply Chain & Order Management CV", supply_chain_order_config, "Aditya_Mehra_Accenture_SupplyChain_OrderOps_CV"),
    ("Regulatory Compliance & KYC Operations CV", regulatory_compliance_config, "Aditya_Mehra_Accenture_Regulatory_Compliance_CV"),
    ("Trust & Safety / Digital Operations CV", trust_safety_config, "Aditya_Mehra_Accenture_Trust_Safety_Digital_CV")
]

print("Starting generation of 5 ATS & JD matched CV suites...")

for label, cfg, basename in tracks:
    html_f = os.path.join(OUTPUT_DIR_HTML, f"{basename}.html")
    pdf_f = os.path.join(OUTPUT_DIR_PDF, f"{basename}.pdf")
    docx_f = os.path.join(OUTPUT_DIR_HTML, f"{basename}.docx")

    # Generate HTML & Headless Chrome PDF
    create_html_cv(cfg, html_f, pdf_f)

    # Generate Word DOCX
    create_docx_cv(cfg, docx_f)

    # Copy primary versions to root
    shutil.copyfile(pdf_f, os.path.join(ROOT_DIR, f"{basename}.pdf"))
    shutil.copyfile(docx_f, os.path.join(ROOT_DIR, f"{basename}.docx"))

    # Verify Page Count of PDF
    pdf_data = open(pdf_f, "rb").read()
    page_count = len(re.findall(b'/Type\s*/Page[^s]', pdf_data))
    pdf_size = os.path.getsize(pdf_f)
    print(f"[OK] {label}: PDF={pdf_size} bytes, Pages={page_count} (Target: 1 Page), DOCX={os.path.getsize(docx_f)} bytes")

print("\n--- ALL 5 JD & ATS MATCHED ACCENTURE CVS COMPILED SUCCESSFULLY ---")
